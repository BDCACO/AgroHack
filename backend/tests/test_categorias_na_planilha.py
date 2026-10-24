"""A categoria na planilha: a coluna `categoria` da importação (só ao cadastrar insumo
novo), o modelo com a lista das cadastradas e o relatório, que a leva na cópia de
segurança."""

import base64
from datetime import timedelta
from io import BytesIO
from typing import Any

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.core.relogio import hoje
from app.models import Base


def _br(dias_atras: int) -> str:
    return (hoje() - timedelta(days=dias_atras)).strftime("%d/%m/%Y")


def _importar(client: TestClient, conteudo: str, *, confirmar: bool = False) -> dict[str, Any]:
    resposta = client.post(
        "/api/v1/importacoes", json={"conteudo": conteudo, "confirmar": confirmar}
    )
    assert resposta.status_code == 200, resposta.json()
    return resposta.json()


def _post(client: TestClient, rota: str, **corpo: Any) -> dict[str, Any]:
    resposta = client.post(f"/api/v1/{rota}", json=corpo)
    assert resposta.status_code == 201, resposta.json()
    return resposta.json()


def _categorias(client: TestClient) -> dict[str, str | None]:
    return {i["nome"]: i["categoria"] for i in client.get("/api/v1/insumos").json()}


def test_categoria_da_planilha_vale_para_o_insumo_novo(client: TestClient) -> None:
    _post(client, "categorias", nome="Fertilizante")
    adubo = _post(client, "categorias", nome="Corretivo de solo")
    _post(client, "insumos", nome="Calcário", unidade="t", categoria_id=adubo["id"])
    planilha = (
        "data;insumo;unidade;tipo;quantidade;categoria\n"
        # Casada sem diferenciar maiúsculas.
        f"{_br(5)};Ureia;kg;entrada;100;fertilizante\n"
        f"{_br(4)};Ureia;;saída;10;\n"
        # Categoria que não existe é cadastrada, como o fornecedor novo.
        f"{_br(4)};Ivermectina 1%;mL;entrada;500;Medicamento veterinário\n"
        f"{_br(3)};Ivermectina 1%;;saída;20;medicamento VETERINÁRIO\n"
        # Insumo já cadastrado: a categoria dele é a da tela, e a da planilha não a muda.
        f"{_br(3)};Calcário;t;entrada;2;Fertilizante\n"
        # Insumo novo sem categoria: fica sem, como sempre.
        f"{_br(2)};Sal mineral;kg;entrada;30;\n"
    )

    previa = _importar(client, planilha)

    assert (previa["erros"], previa["categorias_novas"]) == ([], ["Medicamento veterinário"])
    assert len(client.get("/api/v1/categorias").json()) == 2  # a prévia não grava
    assert _importar(client, planilha, confirmar=True)["importado"] is True
    assert _categorias(client) == {
        "Calcário": "Corretivo de solo",
        "Ivermectina 1%": "Medicamento veterinário",
        "Sal mineral": None,
        "Ureia": "Fertilizante",
    }
    nomes = [c["nome"] for c in client.get("/api/v1/categorias").json()]
    assert nomes == ["Corretivo de solo", "Fertilizante", "Medicamento veterinário"]


def test_categoria_errada_na_planilha_e_erro_na_linha(client: TestClient) -> None:
    _post(client, "categorias", nome="Fertilizante")
    planilha = (
        "data;insumo;unidade;tipo;quantidade;categoria\n"
        f"{_br(5)};Ureia;kg;entrada;100;Fertilizante\n"
        # O mesmo insumo novo em duas categorias: escolher uma seria chute.
        f"{_br(4)};Ureia;;entrada;10;Defensivo agrícola\n"
        f"{_br(4)};Glifosato;L;entrada;10;Defensivo\x0bagrícola\n"
        f"{_br(4)};Calcário;t;entrada;10;{'C' * 61}\n"
    )

    resultado = _importar(client, planilha, confirmar=True)

    assert resultado["erros"] == [
        'Linha 3: "Ureia" já está na categoria Fertilizante em outra linha, não em '
        "Defensivo agrícola.",
        "Linha 4: o nome da categoria tem um caractere inválido.",
        "Linha 5: nome da categoria com mais de 60 caracteres.",
    ]
    assert client.get("/api/v1/insumos").json() == []


def test_modelo_traz_a_coluna_e_as_categorias(client: TestClient) -> None:
    _post(client, "categorias", nome="Fertilizante")
    _post(client, "categorias", nome="Defensivo agrícola")

    livro = load_workbook(BytesIO(client.get("/api/v1/importacoes/modelo.xlsx").content))

    aba = livro["Movimentações"]
    assert aba["J1"].value == "categoria"
    assert "insumo novo" in aba["J1"].comment.text
    validacoes = {str(v.sqref): v for v in aba.data_validations.dataValidation}
    lista = validacoes["J2:J50001"]
    assert lista.formula1 == "'Categorias cadastradas'!$A$2:$A$3"
    assert not lista.showErrorMessage  # categoria nova vale, como o fornecedor novo
    assert list(livro["Categorias cadastradas"].iter_rows(values_only=True)) == [
        ("categoria",),
        ("Defensivo agrícola",),
        ("Fertilizante",),
    ]
    # O exemplo do modelo, com categorias, importa sem erro.
    exemplo = livro["Exemplo"]
    assert exemplo["J2"].value == "Fertilizante"


def test_copia_de_seguranca_devolve_a_categoria(client: TestClient, db: Session) -> None:
    adubo = _post(client, "categorias", nome="Fertilizante")
    ureia = _post(client, "insumos", nome="Ureia", unidade="kg", categoria_id=adubo["id"])
    _post(client, "insumos", nome="Sal mineral", unidade="kg")
    for insumo_id in (ureia["id"], ureia["id"] + 1):
        _post(
            client,
            "movimentacoes",
            insumo_id=insumo_id,
            tipo="entrada",
            data=str(hoje()),
            quantidade=10,
        )
    relatorio = client.get("/api/v1/relatorio.xlsx").content
    aba = load_workbook(BytesIO(relatorio))["Movimentações"]
    assert aba["J1"].value == "categoria"
    assert {linha[1]: linha[9] for linha in aba.iter_rows(min_row=2, values_only=True)} == {
        "Ureia": "Fertilizante",
        "Sal mineral": None,
    }

    # Num sistema vazio, a cópia de segurança devolve cada insumo à sua categoria.
    tabelas = ", ".join(t.name for t in reversed(Base.metadata.sorted_tables))
    for tabela in tabelas.split(", "):
        db.execute(text(f"DELETE FROM {tabela}"))  # noqa: S608 — nomes dos modelos
    db.commit()
    restaurado = client.post(
        "/api/v1/importacoes",
        json={"xlsx": base64.b64encode(relatorio).decode(), "confirmar": True},
    ).json()

    assert (restaurado["erros"], restaurado["importado"]) == ([], True)
    assert _categorias(client) == {"Sal mineral": None, "Ureia": "Fertilizante"}
