"""GET /relatorios/periodo.xlsx e .pdf: o mesmo relatório da tela, nos mesmos filtros, em
arquivo para baixar. As somas são testadas em tests/test_relatorio_por_periodo*.py; aqui,
o que cada arquivo leva e as recusas, que são as da tela."""

from datetime import date, timedelta
from io import BytesIO
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.domain.relatorio_por_periodo import RelatorioPorPeriodo, Totais
from app.services import relatorio_por_periodo_em_arquivo as arquivo
from tests import test_relatorio_por_periodo_rota as rota

# A safra dos testes da rota, reaproveitada: as mesmas compras, aplicações e perdas.
safra = rota.safra
dia = rota.dia

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def baixar(client: TestClient, extensao: str, consulta: str) -> bytes:
    resposta = client.get(f"/api/v1/relatorios/periodo.{extensao}{consulta}")
    assert resposta.status_code == 200, resposta.text
    return resposta.content


def test_excel_com_o_resumo_os_periodos_e_cada_recorte(
    client: TestClient, safra: dict[str, Any]
) -> None:
    resposta = client.get(
        f"/api/v1/relatorios/periodo.xlsx?desde={dia(20)}&ate={dia(0)}&agrupar=dia"
    )

    assert resposta.headers["content-type"] == XLSX
    assert resposta.headers["content-disposition"] == (
        f'attachment; filename="relatorio-campocerto-{dia(20)}-a-{dia(0)}.xlsx"'
    )
    livro = load_workbook(BytesIO(resposta.content))
    assert livro.sheetnames == [
        "Resumo",
        "Por período",
        "Consumo por insumo",
        "Por categoria",
        "Por fornecedor",
    ]
    resumo = {linha[0]: linha[1] for linha in livro["Resumo"].iter_rows(values_only=True)}
    # Os mesmos números de GET /relatorios/periodo (test_relatorio_por_dia_com_os_totais).
    assert resumo["Gasto em compras"] == 4550.0
    assert resumo["Consumo"] == 909.0
    assert resumo["Perdas (descarte e falta)"] == 181.8
    assert resumo["Compras sem preço (fora do gasto)"] == 1

    periodos = list(livro["Por período"].iter_rows(min_row=2, values_only=True))
    assert len(periodos) == 21
    assert livro["Por período"]._charts  # 21 dias cabem no gráfico
    # O que mais custou primeiro; o glifosato, sem preço, com o valor vazio, e não R$ 0,00.
    consumo = list(livro["Consumo por insumo"].iter_rows(min_row=2, values_only=True))
    assert consumo == [("Ureia", 300, "kg", 909.0), ("Glifosato", 5, "L", None)]
    fornecedores = list(livro["Por fornecedor"].iter_rows(min_row=2, values_only=True))
    assert [linha[0] for linha in fornecedores] == ["Agro Norte", "Casa do Adubo"]


def test_o_filtro_aparece_pelo_nome_no_cabecalho(client: TestClient, safra: dict[str, Any]) -> None:
    consulta = f"?desde={dia(20)}&ate={dia(0)}&categoria_id={safra['adubo']['id']}"

    livro = load_workbook(BytesIO(baixar(client, "xlsx", consulta)))
    textos = [linha[0] for linha in livro["Resumo"].iter_rows(values_only=True)]
    assert "Categoria: Fertilizante" in textos
    # Recortado como a tela: só a ureia, que é fertilizante.
    consumo = list(livro["Consumo por insumo"].iter_rows(min_row=2, values_only=True))
    assert [linha[0] for linha in consumo] == ["Ureia"]


def test_pdf_para_imprimir(client: TestClient, safra: dict[str, Any]) -> None:
    resposta = client.get(f"/api/v1/relatorios/periodo.pdf?desde={dia(20)}&ate={dia(0)}")

    assert resposta.headers["content-type"] == "application/pdf"
    assert resposta.headers["content-disposition"].endswith(f'-a-{dia(0)}.pdf"')
    assert resposta.content.startswith(b"%PDF-")
    assert b"/Title" in resposta.content


def test_os_arquivos_recusam_o_que_a_tela_recusa(client: TestClient) -> None:
    for extensao in ("xlsx", "pdf"):
        trocado = client.get(f"/api/v1/relatorios/periodo.{extensao}?desde={dia(0)}&ate={dia(5)}")
        assert trocado.status_code == 422
        assert trocado.json()["detail"] == "O período começa depois de terminar: troque as datas."
        demais = client.get(f"/api/v1/relatorios/periodo.{extensao}?desde=2000-01-01&agrupar=dia")
        assert demais.status_code == 422


def test_sem_movimento_os_arquivos_saem_com_o_aviso(client: TestClient) -> None:
    livro = load_workbook(BytesIO(baixar(client, "xlsx", "")))
    assert livro["Consumo por insumo"]["A2"].value == "Nenhum consumo no período."
    assert baixar(client, "pdf", "").startswith(b"%PDF-")


def _vazio() -> Totais:
    return Totais(
        consumo=[],
        gasto=0,
        compras_sem_preco=0,
        perdas=0,
        perdas_sem_preco=0,
        por_categoria=[],
        por_fornecedor=[],
    )


@pytest.mark.parametrize("nome", ["Café ☕ do sítio", "Adubo 🌱", "Ração \x07"])
def test_nome_com_o_que_a_fonte_nao_tem_nao_derruba_o_pdf_nem_o_excel(nome: str) -> None:
    relatorio = RelatorioPorPeriodo(
        desde=date(2026, 10, 1),
        ate=date(2026, 10, 1) + timedelta(days=2),
        agrupar="dia",
        periodos=[],
        total=_vazio(),
        limitacoes=[],
    )

    assert arquivo.pdf(relatorio, [f"Insumo: {nome}"], date(2026, 10, 3)).startswith(b"%PDF-")
    assert arquivo.planilha(relatorio, [f"Insumo: {nome}"], date(2026, 10, 3)).startswith(b"PK")


def test_numeros_como_se_escreve_no_brasil() -> None:
    assert arquivo._reais(1234.5) == "R$ 1.234,50"
    assert arquivo._numero(1234.5) == "1.234,5"
    assert arquivo._numero(300) == "300"
    assert arquivo._reais_curto(12500) == "R$ 12,5 mil"
