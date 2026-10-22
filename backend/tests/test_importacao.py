"""Importação de movimentações por planilha CSV (POST /importacoes)."""

import base64
import re
import time
from collections.abc import Callable
from datetime import datetime, timedelta
from io import BytesIO
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.core.relogio import hoje
from app.models import Insumo
from app.services import importacao

HOJE = hoje()


def br(dias_atras: int) -> str:
    """Data como o Excel brasileiro escreve: DD/MM/AAAA."""
    return f"{HOJE - timedelta(days=dias_atras):%d/%m/%Y}"


def importar(client: TestClient, conteudo: str, *, confirmar: bool = False) -> dict[str, Any]:
    resposta = client.post(
        "/api/v1/importacoes", json={"conteudo": conteudo, "confirmar": confirmar}
    )
    assert resposta.status_code == 200, resposta.json()
    return resposta.json()


# Como sai do Excel em português: ponto e vírgula, vírgula decimal, BOM no início.
PLANILHA = (
    "\ufeffData;Insumo;Unidade;Tipo;Quantidade;Validade\r\n"
    f"{br(30)};Ureia agrícola;kg;entrada;1.000;\r\n"
    f"{br(20)};ureia agrícola;;saída;25,5;\r\n"
    f"{br(10)};Ivermectina 1%;L;compra;5;{br(-200)}\r\n"
    ";;;;;\r\n"  # linha em branco que o Excel deixa no fim
)


def test_previa_resume_e_nao_grava(client: TestClient) -> None:
    previa = importar(client, PLANILHA)

    assert previa == {
        "linhas": 3,
        "entradas": 2,
        "saidas": 1,
        # A grafia da primeira linha vale; "ureia agrícola" é o mesmo insumo. O total de
        # cada um é onde o produtor percebe um erro de escala antes de gravar.
        "insumos": [
            {
                "nome": "Ureia agrícola",
                "unidade": "kg",
                "novo": True,
                "entradas": 1000,
                "saidas": 25.5,
            },
            {"nome": "Ivermectina 1%", "unidade": "L", "novo": True, "entradas": 5, "saidas": 0},
        ],
        "primeira_data": str(HOJE - timedelta(days=30)),
        "ultima_data": str(HOJE - timedelta(days=10)),
        "ja_lancadas": 0,
        "erros": [],
        "importado": False,
    }
    assert client.get("/api/v1/insumos").json() == []


def test_confirmar_grava_insumos_e_movimentacoes(client: TestClient) -> None:
    resultado = importar(client, PLANILHA, confirmar=True)

    assert resultado["importado"] is True
    estoques = {i["nome"]: i["estoque_atual"] for i in client.get("/api/v1/insumos").json()}
    assert estoques == {"Ivermectina 1%": 5, "Ureia agrícola": 974.5}
    validades = [m["validade"] for m in client.get("/api/v1/movimentacoes").json()]
    assert str(HOJE + timedelta(days=200)) in validades


def test_reimportar_avisa_que_ja_foi_lancada(client: TestClient) -> None:
    importar(client, PLANILHA, confirmar=True)

    previa = importar(client, PLANILHA)

    assert previa["ja_lancadas"] == 3
    assert [i["novo"] for i in previa["insumos"]] == [False, False]


def test_descarte_entra_como_saida_que_nao_e_consumo(client: TestClient) -> None:
    planilha = (
        "data;insumo;unidade;tipo;quantidade;validade\n"
        f"{br(30)};Ivermectina;L;entrada;10;\n"
        f"{br(20)};Ivermectina;;saída;2;\n"
        f"{br(10)};Ivermectina;;Descarte;3;\n"
    )

    previa = importar(client, planilha)
    importar(client, planilha, confirmar=True)

    assert (previa["entradas"], previa["saidas"]) == (1, 2)
    assert previa["insumos"][0]["saidas"] == 5
    movimentacoes = client.get("/api/v1/movimentacoes").json()
    assert [(m["tipo"], m["motivo"]) for m in movimentacoes] == [
        ("saida", "descarte"),
        ("saida", None),
        ("entrada", None),
    ]
    [insumo] = client.get("/api/v1/insumos").json()
    assert insumo["estoque_atual"] == 5


def test_reimportar_nao_confunde_descarte_com_retirada(client: TestClient) -> None:
    planilha = f"data;insumo;unidade;tipo;quantidade\n{br(9)};Sal;kg;entrada;10\n"
    importar(client, planilha + f"{br(5)};Sal;kg;descarte;3\n", confirmar=True)

    previa = importar(client, planilha + f"{br(5)};Sal;kg;saída;3\n")

    assert previa["ja_lancadas"] == 1  # só a entrada


def test_insumo_existente_com_outra_unidade_e_erro(client: TestClient) -> None:
    client.post("/api/v1/insumos", json={"nome": "Calcário", "unidade": "t"})
    planilha = f"data,insumo,unidade,tipo,quantidade\n{br(1)},CALCÁRIO,kg,entrada,2.5\n"

    previa = importar(client, planilha)

    assert previa["erros"] == ['Linha 2: "CALCÁRIO" já está cadastrado em t, não em kg.']


def test_com_um_erro_nada_e_gravado(client: TestClient) -> None:
    planilha = (
        "data;insumo;unidade;tipo;quantidade;validade\n"
        f"{br(5)};Ureia;kg;entrada;100;\n"
        f"31/02/2026;Ureia;kg;saida;10;\n"
        f"{br(-1)};Ureia;kg;saida;10;\n"
        f"{br(3)};Ureia;kg;devolução;10;\n"
        f"{br(3)};Ureia;kg;saida;0;\n"
        f"{br(3)};Ureia;kg;saida;dez;\n"
        f"{br(3)};Ureia;kg;saida;5;{br(-30)}\n"
        f"{br(3)};Glifosato;;entrada;5;\n"
        f"{br(3)};;kg;entrada;5;\n"
    )

    resultado = importar(client, planilha, confirmar=True)

    assert resultado["importado"] is False
    assert resultado["erros"] == [
        'Linha 3: data "31/02/2026" inválida; use DD/MM/AAAA.',
        "Linha 4: data no futuro.",
        'Linha 5: tipo "devolução" desconhecido; use entrada, saída ou descarte.',
        "Linha 6: quantidade deve ser maior que zero.",
        'Linha 7: quantidade "dez" não é um número.',
        "Linha 8: validade só se informa numa entrada.",
        'Linha 9: "Glifosato" não está cadastrado; informe a unidade para cadastrá-lo.',
        "Linha 10: falta o nome do insumo.",
    ]
    assert client.get("/api/v1/insumos").json() == []


def test_coluna_valor_traz_o_preco_da_compra(client: TestClient) -> None:
    planilha = (
        "data;insumo;unidade;tipo;quantidade;validade;valor\n"
        f"{br(30)};Ureia;kg;entrada;1.000;;3,20\n"
        f"{br(20)};Ureia;kg;compra;500;;R$ 3,80\n"
        f"{br(15)};Ureia;kg;entrada;100;;\n"  # sem preço: vale, fica fora do custo
        f"{br(10)};Ureia;kg;saída;25,5;;\n"
    )

    resultado = importar(client, planilha, confirmar=True)

    assert (resultado["erros"], resultado["importado"]) == ([], True)
    precos = [m["valor_unitario"] for m in client.get("/api/v1/movimentacoes").json()]
    assert precos == [None, None, 3.8, 3.2]
    # (1000 × 3,20 + 500 × 3,80) / 1500 = 3,40
    assert client.get("/api/v1/insumos").json()[0]["custo_medio"] == pytest.approx(3.4)


def test_valor_errado_ou_fora_de_uma_compra_e_erro_na_linha(client: TestClient) -> None:
    planilha = (
        "data;insumo;unidade;tipo;quantidade;valor\n"
        f"{br(9)};Sal;kg;entrada;100;2,50\n"
        f"{br(8)};Sal;kg;saída;10;2,50\n"
        f"{br(7)};Sal;kg;descarte;1;2,50\n"
        f"{br(6)};Sal;kg;falta;1;2,50\n"
        f"{br(5)};Sal;kg;sobra;1;2,50\n"
        f"{br(4)};Sal;kg;entrada;1;caro\n"
        f"{br(3)};Sal;kg;entrada;1;0\n"
        f"{br(2)};Sal;kg;entrada;1;-2,5\n"
    )

    assert importar(client, planilha)["erros"] == [
        "Linha 3: valor só se informa numa entrada (compra): é o preço pago.",
        "Linha 4: valor só se informa numa entrada (compra): é o preço pago.",
        "Linha 5: valor só se informa numa entrada (compra): é o preço pago.",
        "Linha 6: valor só se informa numa entrada (compra): é o preço pago.",
        'Linha 7: valor "caro" não é um número.',
        "Linha 8: valor deve ser maior que zero.",
        "Linha 9: valor deve ser maior que zero.",
    ]


def test_valor_tambem_decide_o_separador_decimal(client: TestClient) -> None:
    # Separada por vírgula, "1.000" sozinho seria um (o costume internacional); o valor
    # "3,20" mostra que a vírgula é o decimal, e o ponto, o milhar.
    planilha = f'data,insumo,unidade,tipo,quantidade,valor\n{br(1)},Ureia,kg,entrada,1.000,"3,20"\n'

    previa = importar(client, planilha)

    assert previa["erros"] == []
    assert previa["insumos"][0]["entradas"] == 1000
    misturado = (
        "data;insumo;unidade;tipo;quantidade;valor\n"
        f"{br(2)};A;kg;entrada;25,5;\n{br(1)};A;kg;entrada;2;3.20\n"
    )
    assert importar(client, misturado)["erros"] == [
        "Números com vírgula decimal (linha 2: 25,5) e com ponto decimal (linha 3, valor: "
        "3.20) na mesma planilha: use um formato só."
    ]


def test_reimportar_com_outro_preco_nao_e_a_mesma_compra(client: TestClient) -> None:
    planilha = "data;insumo;unidade;tipo;quantidade;valor\n"
    importar(client, planilha + f"{br(5)};Sal;kg;entrada;10;2,50\n", confirmar=True)

    assert importar(client, planilha + f"{br(5)};Sal;kg;entrada;10;2,50\n")["ja_lancadas"] == 1
    assert importar(client, planilha + f"{br(5)};Sal;kg;entrada;10;2,60\n")["ja_lancadas"] == 0


def test_planilha_sem_as_colunas_do_modelo(client: TestClient) -> None:
    previa = importar(client, "produto;qtd\nUreia;10\n")

    assert previa["erros"] == [
        "Faltam as colunas: data, insumo, tipo, quantidade. A primeira linha precisa ter os "
        "títulos. Use a planilha modelo."
    ]
    assert importar(client, "")["erros"] == ["O arquivo está vazio."]
    so_cabecalho = importar(client, "data;insumo;tipo;quantidade\n")
    assert so_cabecalho["erros"] == ["Nenhuma movimentação na planilha."]


def _quantidades(client: TestClient, planilha: str) -> dict[str, tuple[float, float]]:
    previa = importar(client, planilha)
    assert previa["erros"] == [], previa["erros"]
    return {i["nome"]: (i["entradas"], i["saidas"]) for i in previa["insumos"]}


def test_formato_dos_numeros_e_decidido_pela_planilha_inteira(client: TestClient) -> None:
    cabecalho = "data;insumo;unidade;tipo;quantidade\n"
    # Ponto e vírgula, sem pista: o padrão do Excel em português — ponto é milhar.
    assert _quantidades(client, cabecalho + f"{br(1)};A;kg;entrada;1.000\n") == {"A": (1000, 0)}
    # "0.250" e "2.5" só podem ter ponto decimal; então "1.000" na mesma planilha é 1.
    planilha = cabecalho + f"{br(2)};A;kg;entrada;0.250\n{br(1)};B;kg;entrada;1.000\n"
    assert _quantidades(client, planilha) == {"A": (0.25, 0), "B": (1, 0)}
    # Separada por vírgula, sem pista: o padrão internacional — ponto é decimal.
    planilha = f"data,insumo,unidade,tipo,quantidade\n{br(1)},A,kg,entrada,25.500\n"
    assert _quantidades(client, planilha) == {"A": (25.5, 0)}
    planilha = f'data,insumo,unidade,tipo,quantidade\n{br(1)},A,kg,entrada,"1,234.5"\n'
    assert _quantidades(client, planilha) == {"A": (1234.5, 0)}


def test_planilha_que_mistura_virgula_e_ponto_decimal(client: TestClient) -> None:
    planilha = (
        f"data;insumo;unidade;tipo;quantidade\n{br(2)};A;kg;entrada;25,5\n{br(1)};A;kg;saida;2.5\n"
    )

    assert importar(client, planilha)["erros"] == [
        "Quantidades com vírgula decimal (linha 2: 25,5) e com ponto decimal (linha 3: 2.5) "
        "na mesma planilha: use um formato só."
    ]


def test_linha_com_celulas_a_mais_ou_aspa_sem_fechar(client: TestClient) -> None:
    planilha = (
        "data,insumo,unidade,tipo,quantidade\n"
        f"{br(3)},Ureia,kg,entrada,25,5\n"  # vírgula decimal num arquivo separado por vírgula
        f'{br(2)},"Ureia,kg,entrada,10\n'
        f"{br(1)},Ureia,kg,saida,1\n"
    )

    assert importar(client, planilha)["erros"] == [
        "Linha 2: 6 colunas, mas o cabeçalho tem 5. Se a quantidade tem vírgula decimal, "
        "separe as colunas por ponto e vírgula ou ponha o número entre aspas.",
        'Linha 3: aspa (") sem fechar; o resto do arquivo virou uma célula só.',
    ]


def test_aspa_sem_fechar_num_arquivo_grande_nao_derruba_a_api(client: TestClient) -> None:
    planilha = "data;insumo;unidade;tipo;quantidade\n" + f'{br(1)};"Ureia' + "x" * 200_000

    [erro] = importar(client, planilha)["erros"]

    assert erro.startswith("Linha 2: ")


def test_nomes_que_so_diferem_em_maiusculas_no_cadastro(client: TestClient) -> None:
    client.post("/api/v1/insumos", json={"nome": "Ureia", "unidade": "kg"})
    client.post("/api/v1/insumos", json={"nome": "UREIA", "unidade": "kg"})

    erros = importar(client, f"data;insumo;tipo;quantidade\n{br(1)};ureia;entrada;5\n")["erros"]

    assert erros == [
        'Linha 2: há mais de um insumo cadastrado como "ureia" (UREIA, Ureia); '
        "renomeie um deles na tela de Insumos."
    ]


# --- Planilha do Excel (.xlsx) ---------------------------------------------------------

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def baixar_modelo(client: TestClient) -> Workbook:
    resposta = client.get("/api/v1/importacoes/modelo.xlsx")
    assert resposta.status_code == 200
    assert resposta.headers["content-type"] == XLSX
    assert 'filename="modelo-campocerto.xlsx"' in resposta.headers["content-disposition"]
    return load_workbook(BytesIO(resposta.content))


def importar_xlsx(client: TestClient, livro: Workbook, *, confirmar: bool = False) -> Any:
    arquivo = BytesIO()
    livro.save(arquivo)
    xlsx = base64.b64encode(arquivo.getvalue()).decode()
    resposta = client.post("/api/v1/importacoes", json={"xlsx": xlsx, "confirmar": confirmar})
    assert resposta.status_code == 200, resposta.json()
    return resposta.json()


def preencher(aba: Any, linhas: list[tuple]) -> None:
    for numero, linha in enumerate(linhas, start=2):
        for coluna, valor in enumerate(linha, start=1):
            aba.cell(numero, coluna, valor)


def em_datetime(dias_atras: int) -> datetime:
    return datetime.combine(HOJE - timedelta(days=dias_atras), datetime.min.time())


def test_planilha_modelo_vem_formatada_e_com_os_insumos(client: TestClient) -> None:
    client.post("/api/v1/insumos", json={"nome": "Ureia", "unidade": "kg"})

    livro = baixar_modelo(client)

    assert livro.sheetnames == [
        "Movimentações",
        "Como preencher",
        "Exemplo",
        "Insumos cadastrados",
        "Talhões cadastrados",
    ]
    aba = livro["Movimentações"]
    assert [c.value for c in aba[1]] == [
        "data",
        "insumo",
        "unidade",
        "tipo",
        "quantidade",
        "validade",
        "valor",
        "talhao",
    ]
    assert aba["A1"].font.bold
    assert aba["A1"].comment.text.startswith("Dia em que a compra chegou")
    assert aba.freeze_panes == "A2"
    # Formato da coluna inteira: vale também para a linha 30.000.
    assert aba.column_dimensions["A"].number_format == "DD/MM/YYYY"
    assert aba.column_dimensions["F"].number_format == "DD/MM/YYYY"
    assert aba.column_dimensions["G"].number_format == '"R$" #,##0.00'
    assert aba["G1"].comment.text.startswith("Preço pago por unidade")
    # Validação até a última linha que a importação aceita. Data, tipo e quantidade
    # barram valor errado na hora; insumo e unidade só sugerem (insumo novo vale).
    validacoes = {str(v.sqref): v for v in aba.data_validations.dataValidation}
    assert validacoes["B2:B50001"].formula1 == "'Insumos cadastrados'!$A$2:$A$2"
    assert validacoes["D2:D50001"].formula1 == '"entrada,saída,descarte"'
    barram = {faixa[0] for faixa, v in validacoes.items() if v.showErrorMessage}
    assert barram == {"A", "D", "E", "F", "G"}
    assert (validacoes["G2:G50001"].type, validacoes["G2:G50001"].formula1) == ("decimal", "0")
    assert livro["Insumos cadastrados"]["A2"].value == "Ureia"


def test_modelo_nao_vira_formula_nem_quebra_com_nome_estranho(
    client: TestClient, db: Session
) -> None:
    # Gravado direto no banco: a API já não aceita caractere de controle no nome.
    db.add_all([Insumo(nome="=1+1", unidade="kg"), Insumo(nome="Sal\x0bmineral", unidade="kg")])
    db.commit()

    insumos = baixar_modelo(client)["Insumos cadastrados"]

    assert [(c.value, c.data_type) for c in insumos["A"][1:]] == [
        ("=1+1", "s"),
        ("Salmineral", "s"),
    ]
    # Nome antigo, de antes da regra, não derruba nenhuma tela: a regra vale na entrada.
    estranho = next(i for i in client.get("/api/v1/insumos").json() if "\x0b" in i["nome"])
    assert client.get(f"/api/v1/insumos/{estranho['id']}").status_code == 200
    assert client.get("/api/v1/relatorio.xlsx").status_code == 200


def test_nome_com_caractere_de_controle_e_recusado(client: TestClient) -> None:
    resposta = client.post("/api/v1/insumos", json={"nome": "Sal\x0bmineral", "unidade": "kg"})
    assert resposta.status_code == 422

    planilha = f"data;insumo;unidade;tipo;quantidade\n{br(1)};Sal\x0bmineral;kg;entrada;5\n"
    assert importar(client, planilha)["erros"] == [
        "Linha 2: o nome do insumo tem um caractere inválido."
    ]


def test_modelo_preenchido_volta_pela_importacao(client: TestClient) -> None:
    client.post("/api/v1/insumos", json={"nome": "Ureia", "unidade": "kg"})
    livro = baixar_modelo(client)
    preencher(
        livro["Movimentações"],
        [
            # Data e número com tipo, como o Excel grava a célula formatada...
            (em_datetime(10), "Ureia", None, "entrada", 1000, None, 3.2),
            (em_datetime(5), "ureia", "kg", "saída", 25.5, None),
            # ...e digitados como texto, que também valem.
            (br(3), "Vacina contra aftosa", "dose", "entrada", "12,5", br(-100)),
        ],
    )

    previa = importar_xlsx(client, livro)

    assert previa["erros"] == []
    assert previa["insumos"] == [
        {"nome": "Ureia", "unidade": "kg", "novo": False, "entradas": 1000, "saidas": 25.5},
        {
            "nome": "Vacina contra aftosa",
            "unidade": "dose",
            "novo": True,
            "entradas": 12.5,
            "saidas": 0,
        },
    ]
    assert importar_xlsx(client, livro, confirmar=True)["importado"] is True
    assert client.get("/api/v1/insumos").json()[0]["custo_medio"] == 3.2


def test_o_exemplo_do_modelo_importa_sem_erro(client: TestClient) -> None:
    exemplo = Workbook()
    for linha in baixar_modelo(client)["Exemplo"].iter_rows(values_only=True):
        exemplo.active.append(linha)

    previa = importar_xlsx(client, exemplo)

    assert (previa["erros"], previa["linhas"]) == ([], 3)


def test_modelo_vazio_nao_tem_o_que_importar(client: TestClient) -> None:
    assert importar_xlsx(client, baixar_modelo(client))["erros"] == [
        "Nenhuma movimentação na planilha."
    ]


def test_erros_do_xlsx_apontam_a_linha_do_excel(client: TestClient) -> None:
    livro = Workbook()
    aba = livro.active
    aba.append(["Data", "Insumo", "Unidade", "Tipo", "Quantidade"])
    preencher(
        aba,
        [
            (em_datetime(5), "Ureia", "kg", "entrada", 10),
            (),  # linha 3 em branco: pulada, mas a numeração segue a do Excel
            (em_datetime(4), "Ureia", "kg", "devolução", 5),
            (46000, "Ureia", "kg", "entrada", 5),  # número de série numa célula sem formato
            (em_datetime(2), "Ureia", "kg", "entrada", 1, None, None, "anotação"),
        ],
    )

    assert importar_xlsx(client, livro)["erros"] == [
        'Linha 4: tipo "devolução" desconhecido; use entrada, saída ou descarte.',
        'Linha 5: data "46000" inválida; use DD/MM/AAAA.',
        "Linha 6: há valores na coluna H, que não tem título na linha 1; apague-os ou dê um "
        "título à coluna.",
    ]


def test_arquivo_que_nao_e_xlsx(client: TestClient) -> None:
    resposta = client.post(
        "/api/v1/importacoes", json={"xlsx": base64.b64encode(b"isto nao e uma planilha").decode()}
    )

    [erro] = resposta.json()["erros"]
    assert erro.startswith("Não deu para abrir a planilha do Excel.")


def test_pedido_com_os_dois_formatos_ou_nenhum(client: TestClient) -> None:
    xlsx = base64.b64encode(b"x").decode()
    for corpo in ({}, {"conteudo": "a", "xlsx": xlsx}):
        assert client.post("/api/v1/importacoes", json=corpo).status_code == 422


# --- Planilhas adulteradas ou fora do padrão -------------------------------------------


def com_xml_trocado(livro: Workbook, trocar: Callable[[str], str]) -> bytes:
    """O .xlsx com o XML da primeira aba reescrito: é como se monta um arquivo que o
    Excel não gera, mas que pode chegar num upload."""
    original = BytesIO()
    livro.save(original)
    adulterado = BytesIO()
    with ZipFile(original) as entrada, ZipFile(adulterado, "w", ZIP_DEFLATED) as saida:
        for item in entrada.infolist():
            dados = entrada.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                dados = trocar(dados.decode()).encode()
            saida.writestr(item.filename, dados)
    return adulterado.getvalue()


def importar_bytes(client: TestClient, dados: bytes) -> Any:
    resposta = client.post("/api/v1/importacoes", json={"xlsx": base64.b64encode(dados).decode()})
    assert resposta.status_code == 200, resposta.json()
    return resposta.json()


def planilha_com(linhas: int) -> Workbook:
    livro = Workbook()
    livro.active.append(["data", "insumo", "unidade", "tipo", "quantidade"])
    for dias_atras in range(1, linhas + 1):
        livro.active.append([em_datetime(dias_atras), "Ureia", "kg", "entrada", 1])
    return livro


def test_dimensao_errada_gravada_no_arquivo_nao_corta_a_importacao(client: TestClient) -> None:
    # Alguns exportadores gravam a dimensão errada; lida ao pé da letra, sumiam linhas.
    dados = com_xml_trocado(
        planilha_com(10),
        lambda xml: re.sub(r'<dimension ref="[^"]*"', '<dimension ref="A1:C3"', xml),
    )

    assert importar_bytes(client, dados)["linhas"] == 10


def test_planilha_inflada_e_recusada_rapido(client: TestClient) -> None:
    # Dimensão do tamanho máximo do Excel e uma linha além do limite dele: o leitor
    # preencheria milhões de linhas e colunas vazias.
    def inflar(xml: str) -> str:
        xml = re.sub(r'<dimension ref="[^"]*"', '<dimension ref="A1:XFD1048576"', xml)
        xml = xml.replace('<row r="2"', '<row r="2000000"')
        return re.sub(r'r="([A-Z]+)2"', r'r="\g<1>2000000"', xml)

    inicio = time.perf_counter()
    [erro] = importar_bytes(client, com_xml_trocado(planilha_com(1), inflar))["erros"]

    assert erro.startswith("A planilha é grande demais")
    assert time.perf_counter() - inicio < 5


def test_planilha_que_descompacta_demais_e_recusada(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    # O teto real é de dezenas de MB; aqui baixa para não gerar um arquivo enorme.
    monkeypatch.setattr(importacao, "DESCOMPACTADO_MAXIMO", 1_000)
    arquivo = BytesIO()
    planilha_com(1).save(arquivo)

    [erro] = importar_bytes(client, arquivo.getvalue())["erros"]

    assert erro.startswith("A planilha é grande demais")


def test_valor_em_coluna_sem_titulo_diz_a_coluna(client: TestClient) -> None:
    planilha = f"data;insumo;;tipo;quantidade\n{br(1)};Ureia;kg;entrada;5\n"

    assert importar(client, planilha)["erros"] == [
        "Linha 2: há valores na coluna C, que não tem título na linha 1; apague-os ou dê um "
        "título à coluna."
    ]


def test_pedido_grande_demais_e_recusado_antes_de_ler(client: TestClient) -> None:
    resposta = client.post(
        "/api/v1/importacoes",
        content=b"{}",
        headers={"content-type": "application/json", "content-length": "50000000"},
    )

    assert resposta.status_code == 413
