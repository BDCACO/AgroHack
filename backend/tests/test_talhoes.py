"""Talhões (/talhoes): o cadastro de onde os insumos são aplicados, as aplicações, a
carência e o custo de cada um."""

import base64
from datetime import datetime, timedelta
from io import BytesIO
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core.relogio import hoje

HOJE = hoje()


def dia(dias_atras: int) -> str:
    return str(HOJE - timedelta(days=dias_atras))


def criar_talhao(client: TestClient, **campos: Any) -> dict[str, Any]:
    corpo = {"nome": "Talhão Norte", "area_ha": 120, "cultura": "Soja"} | campos
    resposta = client.post("/api/v1/talhoes", json=corpo)
    assert resposta.status_code == 201, resposta.json()
    return resposta.json()


def test_cadastro_de_talhao_e_nome_duplicado(client: TestClient) -> None:
    talhao = criar_talhao(client, nome="  Talhão Norte  ", area_ha=120.5)

    cadastro = {k: talhao[k] for k in ("nome", "area_ha", "cultura", "demonstracao")}
    assert cadastro == {
        "nome": "Talhão Norte",
        "area_ha": 120.5,
        "cultura": "Soja",
        "demonstracao": False,
    }
    repetido = client.post("/api/v1/talhoes", json={"nome": "Talhão Norte", "area_ha": 3})
    assert repetido.status_code == 409
    assert repetido.json()["detail"] == "Já existe um talhão com esse nome."


def test_cultura_e_opcional(client: TestClient) -> None:
    talhao = criar_talhao(client, nome="Pasto 1", area_ha=45, cultura=None)
    sem = client.post("/api/v1/talhoes", json={"nome": "Pasto 2", "area_ha": 10})
    em_branco = criar_talhao(client, nome="Pasto 3", cultura="  ")

    assert talhao["cultura"] is None
    assert sem.json()["cultura"] is None
    assert em_branco["cultura"] is None


@pytest.mark.parametrize(
    "corpo",
    [
        {"nome": "", "area_ha": 10},
        {"nome": "Norte", "area_ha": 0},  # área é o divisor do custo por hectare
        {"nome": "Norte", "area_ha": -5},
        {"nome": "Norte", "area_ha": 100_000_000},  # passa do NUMERIC(10, 2)
        {"nome": "Nor\x0bte", "area_ha": 10},
        {"nome": "Norte", "area_ha": 10, "cultura": "So\x00ja"},
        {"nome": "N" * 121, "area_ha": 10},
        {"area_ha": 10},
    ],
)
def test_talhao_invalido_e_recusado(client: TestClient, corpo: dict[str, Any]) -> None:
    assert client.post("/api/v1/talhoes", json=corpo).status_code == 422
    assert client.get("/api/v1/talhoes").json() == []


def test_listar_alterar_detalhar_e_remover_talhao(client: TestClient) -> None:
    sul = criar_talhao(client, nome="Talhão Sul", area_ha=80, cultura="Milho")
    norte = criar_talhao(client)

    assert [t["nome"] for t in client.get("/api/v1/talhoes").json()] == [
        "Talhão Norte",
        "Talhão Sul",
    ]
    alterado = client.put(
        f"/api/v1/talhoes/{sul['id']}", json={"nome": "Talhão Sul", "area_ha": 82.25}
    )
    assert alterado.status_code == 200
    assert (alterado.json()["area_ha"], alterado.json()["cultura"]) == (82.25, None)
    duplicado = client.put(
        f"/api/v1/talhoes/{sul['id']}", json={"nome": "Talhão Norte", "area_ha": 1}
    )
    assert duplicado.status_code == 409
    assert client.get(f"/api/v1/talhoes/{norte['id']}").json()["nome"] == "Talhão Norte"

    assert client.delete(f"/api/v1/talhoes/{sul['id']}").status_code == 204
    assert client.get(f"/api/v1/talhoes/{sul['id']}").status_code == 404
    assert client.delete(f"/api/v1/talhoes/{sul['id']}").status_code == 404
    assert (
        client.put(f"/api/v1/talhoes/{sul['id']}", json={"nome": "X", "area_ha": 1}).status_code
        == 404
    )


def criar_insumo(client: TestClient, nome: str = "Glifosato 480", **campos: Any) -> int:
    resposta = client.post("/api/v1/insumos", json={"nome": nome, "unidade": "L"} | campos)
    assert resposta.status_code == 201, resposta.json()
    return resposta.json()["id"]


def lancar(client: TestClient, **corpo: Any) -> Any:
    return client.post("/api/v1/movimentacoes", json=corpo)


def test_retirada_para_uso_vai_para_um_talhao(client: TestClient) -> None:
    talhao = criar_talhao(client)
    glifosato = criar_insumo(client)
    lancar(client, insumo_id=glifosato, tipo="entrada", data=dia(10), quantidade=100)

    resposta = lancar(
        client,
        insumo_id=glifosato,
        tipo="saida",
        data=dia(2),
        quantidade=15,
        talhao_id=talhao["id"],
    )
    sem_talhao = lancar(client, insumo_id=glifosato, tipo="saida", data=dia(1), quantidade=1)

    assert resposta.status_code == 201, resposta.json()
    assert resposta.json()["talhao_id"] == talhao["id"]
    assert sem_talhao.json()["talhao_id"] is None
    listadas = client.get("/api/v1/movimentacoes").json()
    assert [m["talhao_id"] for m in listadas] == [None, talhao["id"], None]
    # Para o estoque, a aplicação é uma retirada como outra qualquer.
    assert client.get(f"/api/v1/insumos/{glifosato}").json()["estoque_atual"] == 84


@pytest.mark.parametrize(
    "corpo",
    [
        {"tipo": "entrada"},
        {"tipo": "saida", "motivo": "descarte"},
    ],
)
def test_talhao_so_numa_retirada_para_uso(client: TestClient, corpo: dict[str, Any]) -> None:
    talhao = criar_talhao(client)
    glifosato = criar_insumo(client)

    resposta = lancar(
        client, insumo_id=glifosato, data=dia(1), quantidade=1, talhao_id=talhao["id"], **corpo
    )

    assert resposta.status_code == 422
    assert "talhão só se informa numa retirada para uso" in str(resposta.json())
    assert client.get("/api/v1/movimentacoes").json() == []


def test_talhao_inexistente_e_422(client: TestClient) -> None:
    glifosato = criar_insumo(client)

    resposta = lancar(
        client, insumo_id=glifosato, tipo="saida", data=dia(1), quantidade=1, talhao_id=99
    )

    assert resposta.status_code == 422
    assert resposta.json()["detail"] == "Talhão não encontrado."


def test_talhao_com_aplicacao_nao_se_remove(client: TestClient) -> None:
    talhao = criar_talhao(client)
    glifosato = criar_insumo(client)
    aplicacao = lancar(
        client, insumo_id=glifosato, tipo="saida", data=dia(1), quantidade=1, talhao_id=talhao["id"]
    ).json()

    recusa = client.delete(f"/api/v1/talhoes/{talhao['id']}")

    assert recusa.status_code == 409
    assert recusa.json()["detail"] == (
        "Talhão Norte tem 1 aplicação registrada: removê-lo apagaria para onde foi cada "
        "insumo. Remova as aplicações antes, se foram lançadas por engano."
    )
    assert client.get(f"/api/v1/talhoes/{talhao['id']}").status_code == 200
    client.delete(f"/api/v1/movimentacoes/{aplicacao['id']}")
    assert client.delete(f"/api/v1/talhoes/{talhao['id']}").status_code == 204


def test_fila_do_app_lanca_no_talhao(client: TestClient) -> None:
    talhao = criar_talhao(client)
    glifosato = criar_insumo(client)
    base = {"insumo_id": glifosato, "data": dia(1), "quantidade": 1}

    resultado = client.post(
        "/api/v1/movimentacoes/lote",
        json={
            "movimentacoes": [
                base
                | {
                    "uuid": "6f1c2a4e-8b3d-4f5a-9c7e-1d2b3a4c5d6e",
                    "tipo": "saida",
                    "talhao_id": talhao["id"],
                },
                base
                | {
                    "uuid": "7f1c2a4e-8b3d-4f5a-9c7e-1d2b3a4c5d6e",
                    "tipo": "saida",
                    "talhao_id": 99,
                },
                base
                | {
                    "uuid": "8f1c2a4e-8b3d-4f5a-9c7e-1d2b3a4c5d6e",
                    "tipo": "entrada",
                    "talhao_id": talhao["id"],
                },
            ]
        },
    ).json()

    gravada, inexistente, em_entrada = resultado["resultados"]
    assert gravada["situacao"] == "gravada"
    assert gravada["movimentacao"]["talhao_id"] == talhao["id"]
    assert (inexistente["situacao"], inexistente["erros"]) == (
        "recusada",
        ["O talhão 99 não existe: pode ter sido removido no site."],
    )
    assert em_entrada["situacao"] == "recusada"
    assert em_entrada["erros"] == [
        "Talhão só se informa numa retirada para uso: é onde o insumo foi aplicado."
    ]


def test_reenvio_com_outro_talhao_nao_e_o_mesmo_lancamento(client: TestClient) -> None:
    norte = criar_talhao(client)
    sul = criar_talhao(client, nome="Talhão Sul")
    glifosato = criar_insumo(client)
    corpo = {
        "uuid": "6f1c2a4e-8b3d-4f5a-9c7e-1d2b3a4c5d6e",
        "insumo_id": glifosato,
        "tipo": "saida",
        "data": dia(1),
        "quantidade": 1,
        "talhao_id": norte["id"],
    }
    assert lancar(client, **corpo).status_code == 201

    assert lancar(client, **corpo).status_code == 200
    outro = lancar(client, **corpo | {"talhao_id": sul["id"]})
    assert outro.status_code == 409
    assert "talhao_id" in outro.json()["detail"]


def test_carencia_no_cadastro_do_insumo(client: TestClient) -> None:
    fungicida = client.post(
        "/api/v1/insumos", json={"nome": "Fungicida", "unidade": "L", "carencia_dias": 30}
    ).json()
    ureia = client.post("/api/v1/insumos", json={"nome": "Ureia", "unidade": "kg"}).json()

    assert (fungicida["carencia_dias"], ureia["carencia_dias"]) == (30, None)
    assert client.get(f"/api/v1/insumos/{fungicida['id']}").json()["carencia_dias"] == 30
    for carencia in (-1, 366, 1.5):
        corpo = {"nome": "Outro", "unidade": "L", "carencia_dias": carencia}
        assert client.post("/api/v1/insumos", json=corpo).status_code == 422, carencia


def test_alterar_insumo_sem_a_carencia_mantem_a_gravada(client: TestClient) -> None:
    # O site e o app de antes não conhecem a carência: um PUT deles não pode apagá-la.
    fungicida = criar_insumo(client, "Fungicida", carencia_dias=30)

    sem = client.put(f"/api/v1/insumos/{fungicida}", json={"nome": "Fungicida", "unidade": "L"})
    zerando = client.put(
        f"/api/v1/insumos/{fungicida}",
        json={"nome": "Fungicida", "unidade": "L", "carencia_dias": 0},
    )
    limpando = client.put(
        f"/api/v1/insumos/{fungicida}",
        json={"nome": "Fungicida", "unidade": "L", "carencia_dias": None},
    )

    assert sem.json()["carencia_dias"] == 30
    assert zerando.json()["carencia_dias"] == 0
    assert limpando.json()["carencia_dias"] is None


def aplicar(client: TestClient, insumo_id: int, talhao_id: int, dias_atras: int, quantidade: float):
    resposta = lancar(
        client,
        insumo_id=insumo_id,
        tipo="saida",
        data=dia(dias_atras),
        quantidade=quantidade,
        talhao_id=talhao_id,
    )
    assert resposta.status_code == 201, resposta.json()
    return resposta.json()


def test_talhao_em_carencia_com_a_conta_e_o_custo(client: TestClient) -> None:
    norte = criar_talhao(client, area_ha=40)
    sul = criar_talhao(client, nome="Talhão Sul", area_ha=10, cultura="Milho")
    fungicida = criar_insumo(client, "Fungicida X", carencia_dias=15)
    ureia = criar_insumo(client, "Ureia", unidade="kg", carencia_dias=0)
    calcario = criar_insumo(client, "Calcário", unidade="t")  # sem carência nem preço
    lancar(
        client, insumo_id=fungicida, tipo="entrada", data=dia(30), quantidade=10, valor_unitario=100
    )
    lancar(client, insumo_id=ureia, tipo="entrada", data=dia(30), quantidade=500, valor_unitario=3)
    lancar(client, insumo_id=calcario, tipo="entrada", data=dia(30), quantidade=5)
    aplicar(client, fungicida, norte["id"], 20, 3)  # liberada há 5 dias
    aplicar(client, fungicida, norte["id"], 6, 2)  # prende até daqui a 9 dias
    aplicar(client, ureia, norte["id"], 3, 100)
    aplicar(client, calcario, sul["id"], 2, 1)

    talhoes = {t["nome"]: t for t in client.get("/api/v1/talhoes").json()}

    norte = talhoes["Talhão Norte"]
    assert norte["colheita_liberada_em"] == str(HOJE + timedelta(days=9))
    assert norte["em_carencia"] is True
    [alerta] = norte["alertas"]
    assert (alerta["tipo"], alerta["gravidade"]) == ("carencia", "alta")
    assert alerta["justificativa"] == (
        f"Talhão Norte: não colher nem abater antes de {HOJE + timedelta(days=9):%d/%m/%Y} "
        f"(faltam 9 dias). Fungicida X aplicado em {HOJE - timedelta(days=6):%d/%m/%Y}, "
        "carência de 15 dias."
    )
    # 5 L × R$ 100,00 + 100 kg × R$ 3,00 = R$ 800,00, em 40 ha.
    assert (norte["custo_total"], norte["custo_por_ha"]) == (800, 20)
    assert any("custo médio de hoje" in t for t in norte["limitacoes"])

    sul = talhoes["Talhão Sul"]
    assert (sul["colheita_liberada_em"], sul["em_carencia"], sul["alertas"]) == (None, False, [])
    assert (sul["custo_total"], sul["custo_por_ha"]) == (None, None)
    assert any("sem carência cadastrada (Calcário)" in t for t in sul["limitacoes"])
    assert any("1 aplicação sem preço (Calcário)" in t for t in sul["limitacoes"])


def test_detalhe_do_talhao_traz_aplicacoes_e_consumo(client: TestClient) -> None:
    norte = criar_talhao(client, area_ha=40)
    fungicida = criar_insumo(client, "Fungicida X", carencia_dias=15)
    lancar(
        client, insumo_id=fungicida, tipo="entrada", data=dia(30), quantidade=10, valor_unitario=100
    )
    antiga = aplicar(client, fungicida, norte["id"], 20, 3)
    recente = aplicar(client, fungicida, norte["id"], 6, 2)

    detalhe = client.get(f"/api/v1/talhoes/{norte['id']}").json()

    assert [a["id"] for a in detalhe["aplicacoes"]] == [recente["id"], antiga["id"]]
    assert detalhe["consumo"] == [
        {
            "insumo_id": fungicida,
            "insumo": "Fungicida X",
            "unidade": "L",
            "quantidade": 5,
            "custo": 500,
        }
    ]
    assert detalhe["em_carencia"] is True


def test_custo_de_uma_safra_pelo_periodo(client: TestClient) -> None:
    norte = criar_talhao(client, area_ha=10)
    fungicida = criar_insumo(client, "Fungicida X", carencia_dias=15)
    lancar(
        client,
        insumo_id=fungicida,
        tipo="entrada",
        data=dia(300),
        quantidade=100,
        valor_unitario=10,
    )
    aplicar(client, fungicida, norte["id"], 250, 50)  # safra passada
    aplicar(client, fungicida, norte["id"], 60, 3)
    aplicar(client, fungicida, norte["id"], 6, 2)
    safra = {"desde": dia(90), "ate": dia(0)}

    [talhao] = client.get("/api/v1/talhoes", params=safra).json()
    detalhe = client.get(f"/api/v1/talhoes/{norte['id']}", params=safra).json()
    tudo = client.get(f"/api/v1/talhoes/{norte['id']}").json()

    # 5 L × R$ 10,00 na safra; 55 L no total.
    assert (talhao["custo_total"], talhao["custo_por_ha"]) == (50, 5)
    assert len(detalhe["aplicacoes"]) == 2
    assert tudo["custo_total"] == 550
    # A carência não depende do período: colher hoje é proibido em qualquer recorte.
    so_a_safra_passada = client.get(
        "/api/v1/talhoes", params={"desde": dia(300), "ate": dia(200)}
    ).json()[0]
    assert so_a_safra_passada["em_carencia"] is True


def test_periodo_invertido_e_422(client: TestClient) -> None:
    norte = criar_talhao(client)

    for url in ("/api/v1/talhoes", f"/api/v1/talhoes/{norte['id']}"):
        resposta = client.get(url, params={"desde": dia(0), "ate": dia(10)})
        assert resposta.status_code == 422
        assert resposta.json()["detail"] == "O período começa depois de terminar: troque as datas."


def test_talhao_sem_aplicacao(client: TestClient) -> None:
    criar_talhao(client)

    [talhao] = client.get("/api/v1/talhoes").json()

    assert talhao | {"id": 0} == {
        "id": 0,
        "nome": "Talhão Norte",
        "area_ha": 120,
        "cultura": "Soja",
        "demonstracao": False,
        "colheita_liberada_em": None,
        "em_carencia": False,
        "alertas": [],
        "custo_total": 0,
        "custo_por_ha": 0,
        "limitacoes": [],
    }


def br(dias_atras: int) -> str:
    return f"{HOJE - timedelta(days=dias_atras):%d/%m/%Y}"


def importar(client: TestClient, conteudo: str, *, confirmar: bool = False) -> dict[str, Any]:
    resposta = client.post(
        "/api/v1/importacoes", json={"conteudo": conteudo, "confirmar": confirmar}
    )
    assert resposta.status_code == 200, resposta.json()
    return resposta.json()


def test_planilha_lanca_a_aplicacao_no_talhao(client: TestClient) -> None:
    norte = criar_talhao(client)
    planilha = (
        "data;insumo;unidade;tipo;quantidade;talhão\n"
        f"{br(10)};Glifosato 480;L;entrada;100;\n"
        f"{br(5)};glifosato 480;;saída;15;talhão norte\n"
        f"{br(4)};Glifosato 480;;saída;2;\n"
    )

    resultado = importar(client, planilha, confirmar=True)

    assert (resultado["erros"], resultado["importado"]) == ([], True)
    movimentacoes = client.get("/api/v1/movimentacoes").json()
    assert [m["talhao_id"] for m in movimentacoes] == [None, norte["id"], None]
    # Reimportar: a mesma aplicação, no mesmo talhão, já está lá.
    assert importar(client, planilha)["ja_lancadas"] == 3


def test_talhao_na_planilha_precisa_estar_cadastrado_e_ser_de_uma_saida(
    client: TestClient,
) -> None:
    criar_talhao(client)
    planilha = (
        "data;insumo;unidade;tipo;quantidade;talhao\n"
        f"{br(10)};Glifosato 480;L;entrada;100;Talhão Norte\n"
        f"{br(5)};Glifosato 480;L;descarte;1;Talhão Norte\n"
        f"{br(4)};Glifosato 480;L;saída;2;Talhão Leste\n"
        f"{br(3)};Glifosato 480;L;saída;2;Talhão Norte\n"
    )

    resultado = importar(client, planilha, confirmar=True)

    assert resultado["erros"] == [
        "Linha 2: talhão só se informa numa saída (retirada para uso): é onde o insumo foi "
        "aplicado.",
        "Linha 3: talhão só se informa numa saída (retirada para uso): é onde o insumo foi "
        "aplicado.",
        'Linha 4: talhão "Talhão Leste" não cadastrado; cadastre na tela de Talhões.',
    ]
    assert client.get("/api/v1/movimentacoes").json() == []


def test_planilha_modelo_traz_a_coluna_e_os_talhoes(client: TestClient) -> None:
    criar_talhao(client)
    criar_talhao(client, nome="Pasto 1", area_ha=45.5, cultura="Pasto")

    resposta = client.get("/api/v1/importacoes/modelo.xlsx")
    livro = load_workbook(BytesIO(resposta.content))

    aba = livro["Movimentações"]
    assert aba["H1"].value == "talhao"
    assert aba["H1"].comment.text.startswith("Onde o insumo foi aplicado")
    validacoes = {str(v.sqref): v for v in aba.data_validations.dataValidation}
    lista = validacoes["H2:H50001"]
    assert lista.formula1 == "'Talhões cadastrados'!$A$2:$A$3"
    assert not lista.showErrorMessage  # só sugere: a importação diz o que está errado
    talhoes = list(livro["Talhões cadastrados"].iter_rows(values_only=True))
    assert talhoes == [
        ("talhão", "área (ha)", "cultura"),
        ("Pasto 1", 45.5, "Pasto"),
        ("Talhão Norte", 120, "Soja"),
    ]
    instrucoes = [c.value for c in livro["Como preencher"]["C"] if c.value]
    assert any("Talhões" in texto for texto in instrucoes)


def test_planilha_modelo_sem_talhao_cadastrado(client: TestClient) -> None:
    livro = load_workbook(BytesIO(client.get("/api/v1/importacoes/modelo.xlsx").content))

    validadas = {str(v.sqref)[0] for v in livro["Movimentações"].data_validations.dataValidation}
    assert "H" not in validadas
    assert livro["Talhões cadastrados"]["A2"].value.startswith("Nenhum talhão cadastrado")


def test_relatorio_traz_os_talhoes_e_a_copia_leva_o_talhao(client: TestClient) -> None:
    norte = criar_talhao(client, area_ha=40)
    criar_talhao(client, nome="Pasto 1", area_ha=10, cultura=None)
    fungicida = criar_insumo(client, "Fungicida X", carencia_dias=15)
    lancar(
        client, insumo_id=fungicida, tipo="entrada", data=dia(30), quantidade=10, valor_unitario=100
    )
    aplicar(client, fungicida, norte["id"], 6, 2)
    lancar(client, insumo_id=fungicida, tipo="saida", data=dia(1), quantidade=1)

    resposta = client.get("/api/v1/relatorio.xlsx")
    livro = load_workbook(BytesIO(resposta.content))

    assert livro.sheetnames == [
        "Resumo",
        "Alertas",
        "Reposição",
        "Estoque",
        "Talhões",
        "Movimentações",
    ]
    talhoes = list(livro["Talhões"].iter_rows(values_only=True))
    assert talhoes[0] == (
        "talhão",
        "área (ha)",
        "cultura",
        "situação",
        "colheita liberada em",
        "custo total (R$)",
        "custo por ha (R$)",
        "carência",
        "limitações",
    )
    pasto, norte_no_relatorio = talhoes[1:]
    assert pasto[:8] == ("Pasto 1", 10, None, "Sem alertas", None, 0, 0, None)
    assert norte_no_relatorio[:7] == (
        "Talhão Norte",
        40,
        "Soja",
        "Urgente",
        datetime.combine(HOJE + timedelta(days=9), datetime.min.time()),
        200,
        5,
    )
    assert norte_no_relatorio[7].startswith("Talhão Norte: não colher nem abater antes de")
    assert livro["Talhões"]["E3"].number_format == "DD/MM/YYYY"
    assert livro["Talhões"]["F3"].number_format == '"R$" #,##0.00'
    resumo = {
        linha[0]: linha[1] for linha in livro["Resumo"].iter_rows(values_only=True) if linha[0]
    }
    assert resumo["Talhões em carência"] == 1

    movimentacoes = list(livro["Movimentações"].iter_rows(values_only=True))
    assert movimentacoes[0][7] == "talhao"
    assert [linha[7] for linha in movimentacoes[1:]] == [None, "Talhão Norte", None]
    previa = client.post(
        "/api/v1/importacoes", json={"xlsx": base64.b64encode(resposta.content).decode()}
    ).json()
    # A cópia de segurança volta igual: a aplicação, no mesmo talhão, já está lá.
    assert (previa["erros"], previa["linhas"], previa["ja_lancadas"]) == ([], 3, 3)
