"""/fornecedores: o cadastro das revendas de quem se compra, e o que se pagou a cada uma."""

from datetime import timedelta
from io import BytesIO
from typing import Any
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core.relogio import hoje

REVENDA = {
    "nome": "  Revenda Boa Safra ",
    "contato": "Seu Antônio",
    "telefone": "(69) 99999-0000",
    "cidade": "Ariquemes",
    "observacao": "Entrega na quinta.\nFrete grátis acima de 20 sacos.",
}


def cadastrar(client: TestClient, **campos: Any) -> dict[str, Any]:
    resposta = client.post("/api/v1/fornecedores", json=REVENDA | campos)
    assert resposta.status_code == 201, resposta.json()
    return resposta.json()


def test_cadastra_lista_e_detalha(client: TestClient) -> None:
    revenda = cadastrar(client)
    cadastrar(client, nome="Agropecuária Central", contato="", telefone=None, observacao=" ")

    assert revenda["nome"] == "Revenda Boa Safra"
    assert (revenda["ativo"], revenda["demonstracao"]) == (True, False)
    assert revenda["observacao"] == "Entrega na quinta.\nFrete grátis acima de 20 sacos."
    lista = client.get("/api/v1/fornecedores").json()
    assert [f["nome"] for f in lista] == ["Agropecuária Central", "Revenda Boa Safra"]
    # Campo apagado na tela chega como "": é fornecedor sem contato, não contato "".
    assert (lista[0]["contato"], lista[0]["telefone"], lista[0]["observacao"]) == (None,) * 3
    detalhe = client.get(f"/api/v1/fornecedores/{revenda['id']}").json()
    assert detalhe["telefone"] == "(69) 99999-0000"
    assert client.get("/api/v1/fornecedores/999").status_code == 404


def test_nome_repetido_e_409_e_dado_invalido_422(client: TestClient) -> None:
    cadastrar(client)

    repetido = client.post("/api/v1/fornecedores", json={"nome": "Revenda Boa Safra"})
    assert repetido.status_code == 409
    assert repetido.json()["detail"] == "Já existe um fornecedor com esse nome."
    for corpo in (
        {"nome": ""},
        {"nome": "Revenda\x00"},
        {"nome": "Ok", "telefone": "9" * 31},
        {"nome": "Ok", "cidade": "Porto\x07Velho"},
        {"nome": "Ok", "observacao": "nulo\x00aqui"},
    ):
        assert client.post("/api/v1/fornecedores", json=corpo).status_code == 422, corpo


def test_alterar_muda_so_o_que_veio(client: TestClient) -> None:
    revenda = cadastrar(client)
    outra = cadastrar(client, nome="Cooperativa Vale Verde")

    resposta = client.put(
        f"/api/v1/fornecedores/{revenda['id']}", json={"nome": "Revenda Boa Safra", "ativo": False}
    )

    assert resposta.status_code == 200
    # O telefone, que não veio, fica: a tela que só desativa não apaga o resto.
    assert (resposta.json()["ativo"], resposta.json()["telefone"]) == (False, "(69) 99999-0000")
    limpo = client.put(f"/api/v1/fornecedores/{revenda['id']}", json={"nome": "R", "telefone": ""})
    assert (limpo.json()["telefone"], limpo.json()["ativo"]) == (None, False)
    nome_de_outro = {"nome": "Cooperativa Vale Verde"}
    assert (
        client.put(f"/api/v1/fornecedores/{revenda['id']}", json=nome_de_outro).status_code == 409
    )
    assert client.put("/api/v1/fornecedores/999", json=nome_de_outro).status_code == 404
    assert client.get(f"/api/v1/fornecedores/{outra['id']}").json()["ativo"] is True


def test_remove_o_fornecedor_sem_historia(client: TestClient) -> None:
    revenda = cadastrar(client)

    assert client.delete(f"/api/v1/fornecedores/{revenda['id']}").status_code == 204
    assert client.get("/api/v1/fornecedores").json() == []
    assert client.delete(f"/api/v1/fornecedores/{revenda['id']}").status_code == 404


def _insumo(client: TestClient, nome: str = "Ureia", unidade: str = "kg") -> int:
    resposta = client.post("/api/v1/insumos", json={"nome": nome, "unidade": unidade})
    return resposta.json()["id"]


def _compra(client: TestClient, insumo_id: int, **campos: Any):
    corpo = {"insumo_id": insumo_id, "tipo": "entrada", "data": str(hoje())} | campos
    return client.post("/api/v1/movimentacoes", json=corpo)


def test_compra_diz_de_quem_foi(client: TestClient) -> None:
    revenda = cadastrar(client)
    ureia = _insumo(client)

    resposta = _compra(client, ureia, quantidade=500, fornecedor_id=revenda["id"])

    assert resposta.status_code == 201
    compra = resposta.json()
    assert (compra["fornecedor_id"], compra["comprado_de"]) == (revenda["id"], "Revenda Boa Safra")
    assert compra["item_do_pedido_id"] is None
    listada = client.get("/api/v1/movimentacoes").json()[0]
    assert listada["comprado_de"] == "Revenda Boa Safra"
    # Sem fornecedor, como sempre.
    assert _compra(client, ureia, quantidade=1).json()["comprado_de"] is None


def test_fornecedor_so_na_compra_e_cadastrado(client: TestClient) -> None:
    revenda = cadastrar(client)
    ureia = _insumo(client)
    base = {"insumo_id": ureia, "data": str(hoje()), "quantidade": 1}

    for corpo in (
        {"tipo": "saida"},
        {"tipo": "saida", "motivo": "descarte"},
    ):
        resposta = client.post(
            "/api/v1/movimentacoes", json=base | corpo | {"fornecedor_id": revenda["id"]}
        )
        assert resposta.status_code == 422, corpo
        assert "fornecedor só se informa numa compra" in str(resposta.json())
    inexistente = _compra(client, ureia, quantidade=1, fornecedor_id=999)
    assert inexistente.status_code == 422
    assert inexistente.json()["detail"] == "Fornecedor não encontrado."
    assert client.get("/api/v1/movimentacoes").json() == []


def test_fornecedor_na_fila_do_app(client: TestClient) -> None:
    revenda = cadastrar(client)
    ureia = _insumo(client)
    base = {"insumo_id": ureia, "data": str(hoje()), "quantidade": 1}
    fila = [
        base | {"tipo": "entrada", "uuid": str(uuid4()), "fornecedor_id": revenda["id"]},
        base | {"tipo": "entrada", "uuid": str(uuid4()), "fornecedor_id": 999},
        base | {"tipo": "saida", "uuid": str(uuid4()), "fornecedor_id": revenda["id"]},
    ]

    resultados = client.post("/api/v1/movimentacoes/lote", json={"movimentacoes": fila}).json()

    gravada, inexistente, saida = resultados["resultados"]
    assert gravada["movimentacao"]["comprado_de"] == "Revenda Boa Safra"
    assert inexistente["situacao"] == "recusada"
    assert inexistente["erros"] == ["O fornecedor 999 não existe: pode ter sido removido no site."]
    assert saida["erros"] == [
        "Fornecedor só se informa numa compra (entrada): é de quem se comprou."
    ]
    # O reenvio do mesmo uuid com outro fornecedor não é o mesmo lançamento.
    outro = cadastrar(client, nome="Cooperativa Vale Verde")
    reenvio = client.post("/api/v1/movimentacoes", json=fila[0] | {"fornecedor_id": outro["id"]})
    assert reenvio.status_code == 409
    assert "fornecedor_id" in reenvio.json()["detail"]


def test_preco_pago_a_cada_fornecedor(client: TestClient) -> None:
    revenda = cadastrar(client)
    cooperativa = cadastrar(client, nome="Cooperativa Vale Verde")
    ureia = _insumo(client)
    sal = _insumo(client, "Sal mineral")

    def dia(atras: int) -> str:
        return str(hoje() - timedelta(days=atras))

    # Na revenda: 1000 kg a R$ 3,00 e 500 kg a R$ 3,60 (média ponderada R$ 3,20).
    _compra(
        client, ureia, data=dia(60), quantidade=1000, valor_unitario=3, fornecedor_id=revenda["id"]
    )
    _compra(
        client, ureia, data=dia(10), quantidade=500, valor_unitario=3.6, fornecedor_id=revenda["id"]
    )
    # Na cooperativa, mais barato; e uma compra sem preço, que fica fora da média.
    _compra(
        client,
        ureia,
        data=dia(30),
        quantidade=800,
        valor_unitario=2.9,
        fornecedor_id=cooperativa["id"],
    )
    _compra(client, ureia, data=dia(20), quantidade=100, fornecedor_id=cooperativa["id"])
    _compra(
        client, sal, data=dia(5), quantidade=200, valor_unitario=4.5, fornecedor_id=revenda["id"]
    )
    _compra(client, ureia, data=dia(1), quantidade=50, valor_unitario=9)  # sem fornecedor

    detalhe = client.get(f"/api/v1/fornecedores/{revenda['id']}").json()

    assert [m["quantidade"] for m in detalhe["compras"]] == [200, 500, 1000]
    assert detalhe["precos"] == [
        {
            "insumo_id": sal,
            "insumo": "Sal mineral",
            "unidade": "kg",
            "fornecedor_id": revenda["id"],
            "fornecedor": "Revenda Boa Safra",
            "compras": 1,
            "quantidade": 200,
            "preco_medio": 4.5,
            "ultimo_preco": 4.5,
            "ultima_compra": dia(5),
        },
        {
            "insumo_id": ureia,
            "insumo": "Ureia",
            "unidade": "kg",
            "fornecedor_id": revenda["id"],
            "fornecedor": "Revenda Boa Safra",
            "compras": 2,
            "quantidade": 1500,
            "preco_medio": 3.2,
            "ultimo_preco": 3.6,
            "ultima_compra": dia(10),
        },
    ]
    # Na página do insumo, as revendas lado a lado, a mais barata primeiro.
    precos = client.get(f"/api/v1/insumos/{ureia}").json()["precos_por_fornecedor"]
    assert [(p["fornecedor"], p["preco_medio"], p["compras"]) for p in precos] == [
        ("Cooperativa Vale Verde", 2.9, 2),
        ("Revenda Boa Safra", 3.2, 2),
    ]


def test_fornecedor_com_compra_nao_se_remove_mas_desativa(client: TestClient) -> None:
    revenda = cadastrar(client)
    _compra(client, _insumo(client), quantidade=10, fornecedor_id=revenda["id"])

    resposta = client.delete(f"/api/v1/fornecedores/{revenda['id']}")

    assert resposta.status_code == 409
    assert resposta.json()["detail"] == (
        "Revenda Boa Safra tem 1 compra no histórico: removê-lo apagaria de quem foi cada "
        "compra e o preço pago a ele. Desative-o (ativo: false): ele sai das escolhas, e a "
        "história fica."
    )
    desativar = {"nome": "Revenda Boa Safra", "ativo": False}
    assert client.put(f"/api/v1/fornecedores/{revenda['id']}", json=desativar).status_code == 200
    # A compra antiga continua dizendo de quem foi.
    assert client.get("/api/v1/movimentacoes").json()[0]["comprado_de"] == "Revenda Boa Safra"


def _br(dias_atras: int) -> str:
    return (hoje() - timedelta(days=dias_atras)).strftime("%d/%m/%Y")


def _importar(client: TestClient, conteudo: str, *, confirmar: bool = False) -> dict[str, Any]:
    resposta = client.post(
        "/api/v1/importacoes", json={"conteudo": conteudo, "confirmar": confirmar}
    )
    assert resposta.status_code == 200, resposta.json()
    return resposta.json()


def test_planilha_com_a_coluna_fornecedor(client: TestClient) -> None:
    cadastrar(client)
    planilha = (
        "data;insumo;unidade;tipo;quantidade;valor;fornecedor\n"
        f"{_br(5)};Ureia;kg;entrada;100;3,20;revenda boa safra\n"
        f"{_br(4)};Ureia;;entrada;50;;Cooperativa Nova\n"
        f"{_br(3)};Ureia;;entrada;20;;cooperativa nova\n"
        f"{_br(2)};Ureia;;saída;10;;\n"
    )

    previa = _importar(client, planilha)

    # O fornecedor novo nasce da planilha, como o insumo novo: só precisa do nome.
    assert (previa["erros"], previa["fornecedores_novos"]) == ([], ["Cooperativa Nova"])
    assert len(client.get("/api/v1/fornecedores").json()) == 1  # a prévia não grava
    assert _importar(client, planilha, confirmar=True)["importado"] is True
    fornecedores = [f["nome"] for f in client.get("/api/v1/fornecedores").json()]
    assert fornecedores == ["Cooperativa Nova", "Revenda Boa Safra"]
    compras = client.get("/api/v1/movimentacoes").json()
    assert [m["comprado_de"] for m in reversed(compras)] == [
        "Revenda Boa Safra",
        "Cooperativa Nova",
        "Cooperativa Nova",
        None,
    ]
    # Reimportada, tudo já está lá, com os mesmos fornecedores.
    assert _importar(client, planilha)["ja_lancadas"] == 4


def test_fornecedor_errado_na_planilha_e_erro_na_linha(client: TestClient) -> None:
    planilha = (
        "data;insumo;unidade;tipo;quantidade;fornecedor\n"
        f"{_br(2)};Ureia;kg;saída;10;Revenda\n"
        f"{_br(2)};Ureia;kg;descarte;1;Revenda\n"
        f"{_br(2)};Ureia;kg;entrada;1;Revenda\x0bBoa\n"
        f"{_br(2)};Ureia;kg;entrada;1;{'R' * 121}\n"
    )

    resultado = _importar(client, planilha, confirmar=True)

    assert resultado["erros"] == [
        "Linha 2: fornecedor só se informa numa entrada (compra): é de quem se comprou.",
        "Linha 3: fornecedor só se informa numa entrada (compra): é de quem se comprou.",
        "Linha 4: o nome do fornecedor tem um caractere inválido.",
        "Linha 5: nome do fornecedor com mais de 120 caracteres.",
    ]
    assert client.get("/api/v1/fornecedores").json() == []


def test_ja_lancadas_trata_o_fornecedor_como_o_preco(client: TestClient) -> None:
    revenda = cadastrar(client)
    cadastrar(client, nome="Cooperativa Vale Verde")
    ureia = _insumo(client)
    _compra(client, ureia, data=str(hoje() - timedelta(days=5)), quantidade=100)
    _compra(
        client,
        ureia,
        data=str(hoje() - timedelta(days=4)),
        quantidade=50,
        fornecedor_id=revenda["id"],
    )

    def previa(linhas: str) -> int:
        cabecalho = "data;insumo;tipo;quantidade;fornecedor\n"
        return _importar(client, cabecalho + linhas)["ja_lancadas"]

    # A compra lançada sem fornecedor, reimportada com ele: completa, não é outra.
    assert previa(f"{_br(5)};Ureia;entrada;100;Revenda Boa Safra\n") == 1
    assert previa(f"{_br(5)};Ureia;entrada;100;Fornecedor que não existe\n") == 1
    # Fornecedor diferente do lançado é outra compra.
    assert previa(f"{_br(4)};Ureia;entrada;50;Cooperativa Vale Verde\n") == 0
    assert previa(f"{_br(4)};Ureia;entrada;50;Fornecedor que não existe\n") == 0
    assert previa(f"{_br(4)};Ureia;entrada;50;revenda boa safra\n") == 1
    assert previa(f"{_br(4)};Ureia;entrada;50;\n") == 1


def test_modelo_traz_a_coluna_e_os_fornecedores(client: TestClient) -> None:
    cadastrar(client)
    cadastrar(client, nome="Cooperativa Vale Verde", cidade="Ariquemes")

    livro = load_workbook(BytesIO(client.get("/api/v1/importacoes/modelo.xlsx").content))

    aba = livro["Movimentações"]
    assert aba["I1"].value == "fornecedor"
    assert aba["I1"].comment.text.startswith("De quem se comprou")
    validacoes = {str(v.sqref): v for v in aba.data_validations.dataValidation}
    lista = validacoes["I2:I50001"]
    assert lista.formula1 == "'Fornecedores cadastrados'!$A$2:$A$3"
    assert not lista.showErrorMessage  # fornecedor novo vale, como o insumo novo
    assert list(livro["Fornecedores cadastrados"].iter_rows(values_only=True)) == [
        ("fornecedor", "cidade"),
        ("Cooperativa Vale Verde", "Ariquemes"),
        ("Revenda Boa Safra", "Ariquemes"),
    ]
