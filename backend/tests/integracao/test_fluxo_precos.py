"""Fluxo real de uso, por HTTP: o produtor que passou a anotar o preço das compras. Cadastra
a ureia com o saco de 50 kg e o pedido mínimo da revenda, lança compras com preços
diferentes e as retiradas, vê o estoque em reais, os alertas com o dinheiro em jogo, a
compra sugerida em sacos fechados, as perdas no resumo, e restaura tudo da cópia de
segurança num CampoCerto vazio, como faria pelo site ou pelo app."""

import base64
from datetime import timedelta
from io import BytesIO
from typing import Any

import httpx
import pytest
from openpyxl import load_workbook

from app.core.relogio import hoje

HOJE = hoje()


def dia(dias_atras: int) -> str:
    return str(HOJE - timedelta(days=dias_atras))


def ok(resposta: httpx.Response, status: int = 200) -> Any:
    assert resposta.status_code == status, resposta.text
    return resposta.json() if resposta.content else None


def lancar(api: httpx.Client, **corpo: Any) -> dict[str, Any]:
    return ok(api.post("/movimentacoes", json=corpo), 201)


def alertas(insumo: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {a["tipo"]: a for a in insumo["analise"]["alertas"]}


CADASTROS = {
    "Ureia agrícola": {
        "unidade": "kg",
        "estoque_minimo": 300,
        "prazo_reposicao_dias": 10,
        "embalagem": 50,
        "pedido_minimo": 500,
    },
    "Ivermectina 1%": {"unidade": "L", "estoque_minimo": 2, "prazo_reposicao_dias": 10},
    "Calcário dolomítico": {"unidade": "t", "prazo_reposicao_dias": 20},
}


def fazenda(api: httpx.Client) -> dict[str, int]:
    ids = {
        nome: ok(api.post("/insumos", json={"nome": nome} | cadastro), 201)["id"]
        for nome, cadastro in CADASTROS.items()
    }
    ureia, ivermectina, calcario = ids.values()
    # Duas compras de ureia, a segunda mais cara: custo médio de R$ 3,30/kg.
    lancar(api, insumo_id=ureia, tipo="entrada", data=dia(100), quantidade=1000, valor_unitario=3)
    lancar(api, insumo_id=ureia, tipo="entrada", data=dia(40), quantidade=1000, valor_unitario=3.6)
    # 30 retiradas de 61 kg nos últimos 90 dias: 20,33 kg/dia.
    for retirada in range(30):
        lancar(api, insumo_id=ureia, tipo="saida", data=dia(89 - 3 * retirada), quantidade=61)
    # Ivermectina: o lote de R$ 90 venceu há 52 dias; o de R$ 100 está em dia.
    lancar(
        api,
        insumo_id=ivermectina,
        tipo="entrada",
        data=dia(120),
        quantidade=10,
        validade=dia(52),
        valor_unitario=90,
    )
    lancar(
        api,
        insumo_id=ivermectina,
        tipo="entrada",
        data=dia(60),
        quantidade=10,
        validade=dia(-300),
        valor_unitario=100,
    )
    for semana in range(8):
        lancar(api, insumo_id=ivermectina, tipo="saida", data=dia(50 - 7 * semana), quantidade=1)
    # Calcário comprado para anos: 100 t a R$ 200, e quase nada saiu.
    lancar(
        api, insumo_id=calcario, tipo="entrada", data=dia(100), quantidade=100, valor_unitario=200
    )
    lancar(api, insumo_id=calcario, tipo="saida", data=dia(80), quantidade=1)
    lancar(api, insumo_id=calcario, tipo="saida", data=dia(40), quantidade=1)
    return ids


def test_precos_embalagem_e_perdas_de_ponta_a_ponta(api: httpx.Client) -> None:
    ids = fazenda(api)
    ureia = ok(api.get(f"/insumos/{ids['Ureia agrícola']}"))

    # (1000 × 3,00 + 1000 × 3,60) / 2000; restam 2000 − 30 × 61 = 170 kg.
    assert (ureia["custo_medio"], ureia["estoque_atual"]) == (3.3, 170)
    assert ureia["valor_em_estoque"] == pytest.approx(561)
    compras = [m["valor_unitario"] for m in ureia["movimentacoes"] if m["tipo"] == "entrada"]
    assert compras == [3.6, 3]
    # A conta dá 300 + 30 × 20,33 = 910 kg: acima do pedido mínimo (500), e em sacos de
    # 50 kg são 19 sacos, 950 kg, a R$ 3,30.
    reposicao = ureia["analise"]["reposicao"]
    assert reposicao["pedir_hoje"] is True
    assert reposicao["quantidade"] == 950
    assert reposicao["custo_estimado"] == pytest.approx(3135)
    assert reposicao["justificativa"].startswith("Pedir hoje 19 embalagens de 50 kg (950 kg)")
    assert "A conta dá 910 kg" in reposicao["justificativa"]
    assert not any("não considera" in t for t in ureia["analise"]["limitacoes"])

    ivermectina = ok(api.get(f"/insumos/{ids['Ivermectina 1%']}"))
    vencido = alertas(ivermectina)["vencido"]
    # O lote vencido inteiro (10 L) pelo custo médio de R$ 95,00/L.
    assert vencido["valor"] == pytest.approx(950)
    assert vencido["justificativa"].startswith("10 L vencidos")
    assert vencido["justificativa"].endswith(
        "Pelo custo médio de R$ 95,00/L, são cerca de R$ 950,00."
    )

    calcario = ok(api.get(f"/insumos/{ids['Calcário dolomítico']}"))
    excesso = alertas(calcario)["excesso"]
    # 2 t em 90 dias: 180 dias de consumo são 4 t; as outras 94 t a R$ 200 estão paradas.
    assert excesso["valor"] == pytest.approx(18_800)
    assert "R$ 18.800,00" in excesso["justificativa"]

    # Devolve o lote vencido e conta o galpão: meio litro sumiu. Os dois são perda.
    lancar(
        api,
        insumo_id=ids["Ivermectina 1%"],
        tipo="saida",
        motivo="descarte",
        data=dia(0),
        quantidade=10,
    )
    ajuste = ok(
        api.post(f"/insumos/{ids['Ivermectina 1%']}/inventario", json={"quantidade_contada": 1.5})
    )
    assert ajuste["diferenca"] == -0.5
    insumos = ok(api.get("/insumos"))
    resumo = ok(api.get("/resumo"))

    # 10 L + 0,5 L a R$ 95,00.
    assert resumo["perdas"] == {"ultimos_30_dias": 997.5, "total": 997.5}
    assert resumo["valor_em_estoque"] == pytest.approx(sum(i["valor_em_estoque"] for i in insumos))
    assert resumo["valor_em_estoque"] == pytest.approx(561 + 1.5 * 95 + 98 * 200)
    reposicoes = [i["analise"]["reposicao"] for i in insumos if i["analise"]["reposicao"]]
    proximas = [r for r in reposicoes if r["pedir_em"] <= str(HOJE + timedelta(days=30))]
    assert resumo["compras_sugeridas"] == {
        "valor": pytest.approx(sum(r["custo_estimado"] for r in proximas)),
        "pedidos": len(proximas),
        "pedir_hoje": sum(r["pedir_hoje"] for r in proximas),
    }
    assert resumo["compras_sugeridas"]["pedir_hoje"] == 2  # ureia e ivermectina
    assert resumo["insumos_sem_preco"] == 0
    assert resumo["alertas"]["alta"] >= 1

    # O relatório leva os preços: custo médio no Estoque e o valor de cada compra.
    relatorio = api.get("/relatorio.xlsx")
    assert relatorio.status_code == 200
    livro = load_workbook(BytesIO(relatorio.content))
    estoque = {linha[0]: linha for linha in livro["Estoque"].iter_rows(min_row=2, values_only=True)}
    assert estoque["Ureia agrícola"][7:9] == (3.3, 561)
    precos = [linha[6] for linha in livro["Movimentações"].iter_rows(min_row=2, values_only=True)]
    assert sorted(p for p in precos if p) == [3, 3.6, 90, 100, 200]
    resumo_do_relatorio = {
        linha[0]: linha[1] for linha in livro["Resumo"].iter_rows(values_only=True) if linha[0]
    }
    assert resumo_do_relatorio["Perdas no total"] == 997.5


def test_restauracao_da_copia_de_seguranca_traz_os_precos(api: httpx.Client) -> None:
    """Perdeu tudo: o relatório importado num CampoCerto vazio refaz preços, custo médio e
    estoque; relançado o cadastro (que a cópia não leva), a análise volta igual."""
    fazenda(api)
    originais = {i["nome"]: i for i in ok(api.get("/insumos"))}
    movimentacoes = ok(api.get("/movimentacoes", params={"limite": 1000}))
    relatorio = api.get("/relatorio.xlsx").content

    for insumo in originais.values():
        ok(api.delete(f"/insumos/{insumo['id']}"), 204)
    assert ok(api.get("/insumos")) == []
    resultado = ok(
        api.post(
            "/importacoes",
            json={"xlsx": base64.b64encode(relatorio).decode(), "confirmar": True},
        )
    )

    assert (resultado["erros"], resultado["importado"]) == ([], True)
    restaurados = {i["nome"]: i for i in ok(api.get("/insumos"))}

    def sem_id(lista: list[dict[str, Any]]) -> list[tuple]:
        return sorted(
            (m["data"], m["tipo"], m["quantidade"], m["validade"], m["valor_unitario"])
            for m in lista
        )

    assert sem_id(ok(api.get("/movimentacoes", params={"limite": 1000}))) == sem_id(movimentacoes)
    for nome, original in originais.items():
        restaurado = restaurados[nome]
        assert restaurado["custo_medio"] == original["custo_medio"]
        assert restaurado["valor_em_estoque"] == original["valor_em_estoque"]
        # O cadastro não vai na cópia: relançado, a análise (em reais) é a mesma.
        ok(api.put(f"/insumos/{restaurado['id']}", json={"nome": nome} | CADASTROS[nome]))
        assert ok(api.get(f"/insumos/{restaurado['id']}"))["analise"] == original["analise"]
    # Importar de novo a mesma cópia: tudo já está lá, com os mesmos preços.
    previa = ok(api.post("/importacoes", json={"xlsx": base64.b64encode(relatorio).decode()}))
    assert previa["ja_lancadas"] == previa["linhas"] == len(movimentacoes)


def test_erros_de_preco_e_embalagem_respondem_422(api: httpx.Client) -> None:
    insumo = ok(api.post("/insumos", json={"nome": "Sal mineral", "unidade": "kg"}), 201)
    base = {"insumo_id": insumo["id"], "data": dia(0), "quantidade": 1}

    for corpo in (
        {"tipo": "saida", "valor_unitario": 4.5},  # preço é da compra
        {"tipo": "saida", "motivo": "descarte", "valor_unitario": 4.5},
        {"tipo": "entrada", "valor_unitario": -4.5},
        {"tipo": "entrada", "valor_unitario": 0},
    ):
        assert api.post("/movimentacoes", json=base | corpo).status_code == 422, corpo
    for cadastro in ({"embalagem": 0}, {"pedido_minimo": -1}):
        resposta = api.post("/insumos", json={"nome": "Ureia", "unidade": "kg"} | cadastro)
        assert resposta.status_code == 422
    assert ok(api.get("/movimentacoes")) == []
    assert [i["nome"] for i in ok(api.get("/insumos"))] == ["Sal mineral"]
    # Planilha com preço numa saída: erro na linha, e nada é gravado.
    planilha = f"data;insumo;tipo;quantidade;valor\n{HOJE:%d/%m/%Y};Sal mineral;saída;1;4,50\n"
    resultado = ok(api.post("/importacoes", json={"conteudo": planilha, "confirmar": True}))
    assert resultado["erros"] == [
        "Linha 2: valor só se informa numa entrada (compra): é o preço pago."
    ]
    assert ok(api.get("/movimentacoes")) == []
