"""Fluxo real de uso, por HTTP: o produtor que lança no campo, sem sinal. O app guarda a fila
no celular e a manda em lote quando a rede volta; a resposta se perde e o lote vai de novo;
um item ruim no meio não trava os outros; o celular reenvia ao mesmo tempo que o primeiro
envio ainda grava. No fim, o estoque, a análise, o relatório e a restauração da cópia de
segurança têm de estar como se cada lançamento tivesse sido feito uma vez só."""

import base64
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from io import BytesIO
from typing import Any
from uuid import uuid4

import httpx
import pytest
from openpyxl import load_workbook

from app.core.relogio import hoje

HOJE = hoje()


def dia(dias_atras: int) -> str:
    return str(HOJE - timedelta(days=dias_atras))


def ok(resposta: httpx.Response, status: int = 200) -> Any:
    assert resposta.status_code == status, resposta.text
    return resposta.json() if resposta.content else None


def lancamento(insumo_id: int, **campos: Any) -> dict[str, Any]:
    """Como o app monta um lançamento: o uuid nasce no celular, na hora de lançar."""
    return {"uuid": str(uuid4()), "insumo_id": insumo_id, "tipo": "saida", "data": dia(0)} | campos


def sincronizar(api: httpx.Client, fila: list[Any]) -> dict[str, Any]:
    return ok(api.post("/movimentacoes/lote", json={"movimentacoes": fila}))


def fila_do_campo(ureia: int, ivermectina: int) -> list[dict[str, Any]]:
    # A ureia chegou há 60 dias; 20 retiradas de 20 kg, uma a cada 3 dias.
    fila = [lancamento(ureia, tipo="entrada", data=dia(60), quantidade=1000, valor_unitario=3.25)]
    fila += [lancamento(ureia, data=dia(59 - 3 * i), quantidade=20) for i in range(20)]
    fila += [
        lancamento(ivermectina, tipo="entrada", data=dia(20), quantidade=2, validade=dia(-200)),
        lancamento(ivermectina, data=dia(10), quantidade=0.5),
    ]
    return fila


@pytest.fixture
def insumos(api: httpx.Client) -> tuple[int, int]:
    ureia = ok(
        api.post(
            "/insumos",
            json={"nome": "Ureia", "unidade": "kg", "estoque_minimo": 100, "embalagem": 50},
        ),
        201,
    )
    ivermectina = ok(api.post("/insumos", json={"nome": "Ivermectina 1%", "unidade": "L"}), 201)
    return ureia["id"], ivermectina["id"]


def test_fila_offline_resposta_perdida_e_item_ruim(
    api: httpx.Client, insumos: tuple[int, int]
) -> None:
    ureia, ivermectina = insumos
    fila = fila_do_campo(ureia, ivermectina)

    primeiro = sincronizar(api, fila)

    assert primeiro["resumo"] == {"gravadas": 23, "ja_existiam": 0, "recusadas": 0}
    assert [r["uuid"] for r in primeiro["resultados"]] == [i["uuid"] for i in fila]
    # Nada a conferir, a não ser a data dos lançamentos de mais de 30 dias atrás.
    atrasados = [i["data"] < dia(30) for i in fila]
    assert [bool(r["avisos"]) for r in primeiro["resultados"]] == atrasados
    assert all(
        len(r["avisos"]) == 1 and r["avisos"][0].startswith("Lançamento atrasado")
        for r in primeiro["resultados"]
        if r["avisos"]
    )
    antes = ok(api.get("/insumos"))

    # A resposta se perdeu no caminho: o app manda a mesma fila de novo.
    segundo = sincronizar(api, fila)

    assert segundo["resumo"] == {"gravadas": 0, "ja_existiam": 23, "recusadas": 0}
    assert [r["movimentacao"] for r in segundo["resultados"]] == [
        r["movimentacao"] for r in primeiro["resultados"]
    ]
    assert ok(api.get("/insumos")) == antes
    assert len(ok(api.get("/movimentacoes", params={"limite": 1000}))) == 23

    # Mais lançamentos no campo, um deles com a quantidade digitada errada.
    seguintes = [
        lancamento(ureia, data=dia(1), quantidade=20),
        lancamento(ureia, data=dia(1), quantidade="vinte"),
        lancamento(ureia, data=dia(0), quantidade=20),
    ]
    terceiro = sincronizar(api, seguintes)

    assert [r["situacao"] for r in terceiro["resultados"]] == ["gravada", "recusada", "gravada"]
    assert terceiro["resultados"][1]["erros"] == ["quantidade tem de ser um número."]
    assert terceiro["resultados"][1]["uuid"] == seguintes[1]["uuid"]

    # Estoque e análise como se cada lançamento tivesse sido feito uma vez só.
    detalhe = ok(api.get(f"/insumos/{ureia}"))
    assert detalhe["estoque_atual"] == 1000 - 22 * 20
    # 440 kg em 61 dias (de 60 dias atrás até hoje, inclusive).
    assert detalhe["analise"]["consumo_diario"] == pytest.approx(440 / 61)
    assert detalhe["custo_medio"] == 3.25
    assert {m["uuid"] for m in detalhe["movimentacoes"]} >= {i["uuid"] for i in fila[:21]}
    assert ok(api.get(f"/insumos/{ivermectina}"))["estoque_atual"] == 1.5


def test_avisos_de_estoque_negativo_e_de_data_atrasada(
    api: httpx.Client, insumos: tuple[int, int]
) -> None:
    ureia, ivermectina = insumos
    fila = [
        lancamento(ivermectina, tipo="entrada", data=dia(3), quantidade=1),
        # A entrada que trouxe o resto não foi lançada: o estoque fica negativo.
        lancamento(ivermectina, data=dia(2), quantidade=4),
        # O mês digitado errado: 40 dias atrás.
        lancamento(ureia, tipo="entrada", data=dia(40), quantidade=50),
    ]

    resultado = sincronizar(api, fila)

    assert resultado["resumo"]["gravadas"] == 3
    assert [r["avisos"] for r in resultado["resultados"]] == [
        [],
        ["O estoque de Ivermectina 1% fica em -3 L com este envio: falta lançar uma entrada?"],
        [
            f"Lançamento atrasado: a data é {HOJE - timedelta(days=40):%d/%m/%Y}, 40 dias "
            "atrás. Confira a data."
        ],
    ]
    assert ok(api.get(f"/insumos/{ivermectina}"))["estoque_atual"] == -3


def test_lancamento_digitado_no_site_e_no_app(api: httpx.Client, insumos: tuple[int, int]) -> None:
    ureia, _ = insumos
    no_site = {"insumo_id": ureia, "tipo": "entrada", "data": dia(0), "quantidade": 0.1}
    ok(api.post("/movimentacoes", json=no_site), 201)

    resultado = sincronizar(api, [lancamento(ureia) | no_site])

    [item] = resultado["resultados"]
    assert item["situacao"] == "gravada"
    assert item["avisos"] == [
        "Já existe um lançamento igual neste dia, feito sem o app (no site ou pela planilha). "
        "Se for o mesmo, remova um dos dois."
    ]


def test_post_simples_com_uuid_repetido_e_com_outro_conteudo(
    api: httpx.Client, insumos: tuple[int, int]
) -> None:
    ureia, _ = insumos
    corpo = lancamento(ureia, tipo="entrada", quantidade=12.3456, valor_unitario=3.123456)

    primeiro = ok(api.post("/movimentacoes", json=corpo), 201)
    segundo = ok(api.post("/movimentacoes", json=corpo), 200)
    conflito = api.post("/movimentacoes", json=corpo | {"data": dia(1)})

    # O PostgreSQL guardou 3 casas na quantidade e 4 no preço; o reenvio ainda é o mesmo.
    assert primeiro == segundo
    assert (primeiro["quantidade"], primeiro["valor_unitario"]) == (12.346, 3.1235)
    assert conflito.status_code == 409
    assert conflito.json()["detail"] == (
        "Este uuid já identifica outro lançamento, diferente deste em: data. Cada lançamento "
        "precisa de um uuid novo."
    )
    assert len(ok(api.get("/movimentacoes"))) == 1


def test_envios_simultaneos_do_mesmo_lancamento(
    api: httpx.Client, insumos: tuple[int, int]
) -> None:
    """O app reenvia enquanto o primeiro envio ainda grava: as requisições chegam juntas ao
    uvicorn, e só o UNIQUE do banco separa uma da outra."""
    ureia, ivermectina = insumos
    base = str(api.base_url)

    def enviar(caminho: str, corpo: dict[str, Any]) -> httpx.Response:
        with httpx.Client(base_url=base, timeout=30) as cliente:
            return cliente.post(caminho, json=corpo)

    lancamentos = [lancamento(ureia, quantidade=i + 1) for i in range(10)]
    with ThreadPoolExecutor(max_workers=8) as threads:
        simples = [
            list(threads.map(lambda corpo: enviar("/movimentacoes", corpo), [corpo] * 8))
            for corpo in lancamentos
        ]
        fila = fila_do_campo(ureia, ivermectina)
        lotes = list(
            threads.map(lambda _: enviar("/movimentacoes/lote", {"movimentacoes": fila}), range(6))
        )

    for respostas in simples:
        # Um gravou; os outros receberam o mesmo lançamento, já gravado. Nenhum 500.
        assert sorted(r.status_code for r in respostas) == [200] * 7 + [201]
        assert len({r.json()["id"] for r in respostas}) == 1
    assert {r.status_code for r in lotes} == {200}
    for posicao in range(len(fila)):
        situacoes = Counter(r.json()["resultados"][posicao]["situacao"] for r in lotes)
        assert situacoes == {"gravada": 1, "ja_existia": 5}
    gravados = Counter(m["uuid"] for m in ok(api.get("/movimentacoes", params={"limite": 1000})))
    assert len(gravados) == len(lancamentos) + len(fila)
    assert set(gravados.values()) == {1}
    assert ok(api.get(f"/insumos/{ureia}"))["estoque_atual"] == 1000 - 20 * 20 - 55


def test_relatorio_e_restauracao_com_lancamentos_do_app(
    api: httpx.Client, insumos: tuple[int, int]
) -> None:
    """O uuid é detalhe técnico: não vai para a planilha, e a cópia de segurança restaurada
    refaz o estoque como antes."""
    ureia, ivermectina = insumos
    sincronizar(api, fila_do_campo(ureia, ivermectina))
    originais = {i["nome"]: i["estoque_atual"] for i in ok(api.get("/insumos"))}

    relatorio = api.get("/relatorio.xlsx")

    assert relatorio.status_code == 200
    aba = load_workbook(BytesIO(relatorio.content))["Movimentações"]
    cabecalho = [str(c.value).lower() for c in next(aba.iter_rows(max_row=1))]
    assert not any("uuid" in titulo for titulo in cabecalho)
    assert aba.max_row == 1 + 23

    for insumo_id in insumos:
        ok(api.delete(f"/insumos/{insumo_id}"), 204)
    xlsx = base64.b64encode(relatorio.content).decode()
    resultado = ok(api.post("/importacoes", json={"xlsx": xlsx, "confirmar": True}))

    assert (resultado["erros"], resultado["importado"]) == ([], True)
    assert {i["nome"]: i["estoque_atual"] for i in ok(api.get("/insumos"))} == originais
    restauradas = ok(api.get("/movimentacoes", params={"limite": 1000}))
    assert len(restauradas) == 23
    assert {m["uuid"] for m in restauradas} == {None}
