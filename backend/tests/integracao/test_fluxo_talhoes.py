"""Fluxo real de uso, por HTTP: o produtor que passou a dizer onde aplicou cada insumo.
Cadastra os talhões e os defensivos com a carência da bula e o preço, lança as aplicações
pelo site e pela fila do app, vê até quando não colher (com a conta), quanto custou cada
talhão por hectare na safra, e restaura tudo da cópia de segurança num CampoCerto vazio."""

import base64
from datetime import timedelta
from io import BytesIO
from typing import Any

import httpx
from openpyxl import load_workbook
from sqlalchemy import Engine, text

from app.core.relogio import hoje
from app.models import Base

HOJE = hoje()


def dia(dias_atras: int) -> str:
    return str(HOJE - timedelta(days=dias_atras))


def br(dias: int) -> str:
    return f"{HOJE + timedelta(days=dias):%d/%m/%Y}"


def ok(resposta: httpx.Response, status: int = 200) -> Any:
    assert resposta.status_code == status, resposta.text
    return resposta.json() if resposta.content else None


def lancar(api: httpx.Client, **corpo: Any) -> dict[str, Any]:
    return ok(api.post("/movimentacoes", json=corpo), 201)


TALHOES = {
    "Talhão Norte": {"area_ha": 40, "cultura": "Soja"},
    "Talhão Sul": {"area_ha": 12.5, "cultura": "Milho"},
}
INSUMOS = {
    "Fungicida X": {"unidade": "L", "carencia_dias": 30},
    "Inseticida Y": {"unidade": "L", "carencia_dias": 14},
    "Ureia": {"unidade": "kg", "carencia_dias": 0},
    "Adubo foliar": {"unidade": "L"},  # sem carência cadastrada nem preço
}


def fazenda(api: httpx.Client) -> tuple[dict[str, int], dict[str, int]]:
    talhoes = {
        nome: ok(api.post("/talhoes", json={"nome": nome} | cadastro), 201)["id"]
        for nome, cadastro in TALHOES.items()
    }
    insumos = {
        nome: ok(api.post("/insumos", json={"nome": nome} | cadastro), 201)["id"]
        for nome, cadastro in INSUMOS.items()
    }
    for nome, quantidade, preco in (
        ("Fungicida X", 20, 120),
        ("Inseticida Y", 10, 80),
        ("Ureia", 1000, 3),
    ):
        lancar(
            api,
            insumo_id=insumos[nome],
            tipo="entrada",
            data=dia(200),
            quantidade=quantidade,
            valor_unitario=preco,
        )
    lancar(api, insumo_id=insumos["Adubo foliar"], tipo="entrada", data=dia(200), quantidade=50)
    return talhoes, insumos


def aplicar(api: httpx.Client, insumo: int, talhao: int, dias_atras: int, quantidade: float):
    return lancar(
        api,
        insumo_id=insumo,
        tipo="saida",
        data=dia(dias_atras),
        quantidade=quantidade,
        talhao_id=talhao,
    )


def test_carencia_e_custo_por_talhao_de_ponta_a_ponta(api: httpx.Client) -> None:
    talhoes, insumos = fazenda(api)
    norte, sul = talhoes["Talhão Norte"], talhoes["Talhão Sul"]
    # Safra passada no Norte: entra no custo de sempre, não no da safra.
    aplicar(api, insumos["Ureia"], norte, 150, 200)
    # Pelo site: fungicida há 10 dias (30 de carência, libera em 20) e ureia.
    aplicar(api, insumos["Fungicida X"], norte, 10, 4)
    aplicar(api, insumos["Ureia"], norte, 8, 100)
    # Pelo app, sem sinal: inseticida ontem (14 de carência, libera em 13) e adubo foliar.
    fila = {
        "movimentacoes": [
            {
                "uuid": "1a2b3c4d-0000-4000-8000-000000000001",
                "insumo_id": insumos["Inseticida Y"],
                "tipo": "saida",
                "data": dia(1),
                "quantidade": 2,
                "talhao_id": norte,
            },
            {
                "uuid": "1a2b3c4d-0000-4000-8000-000000000002",
                "insumo_id": insumos["Adubo foliar"],
                "tipo": "saida",
                "data": dia(1),
                "quantidade": 5,
                "talhao_id": sul,
            },
        ]
    }
    lote = ok(api.post("/movimentacoes/lote", json=fila))
    assert lote["resumo"] == {"gravadas": 2, "ja_existiam": 0, "recusadas": 0}
    # O sinal caiu antes da resposta: o app manda de novo, e nada duplica.
    assert ok(api.post("/movimentacoes/lote", json=fila))["resumo"]["ja_existiam"] == 2

    lista = {t["nome"]: t for t in ok(api.get("/talhoes"))}

    # Vale a carência que acaba por último (o fungicida), e não a da aplicação mais recente.
    norte_hoje = lista["Talhão Norte"]
    assert norte_hoje["em_carencia"] is True
    assert norte_hoje["colheita_liberada_em"] == str(HOJE + timedelta(days=20))
    [alerta] = norte_hoje["alertas"]
    assert alerta["gravidade"] == "alta"
    assert alerta["justificativa"] == (
        f"Talhão Norte: não colher nem abater antes de {br(20)} (faltam 20 dias). "
        f"Fungicida X aplicado em {br(-10)}, carência de 30 dias."
    )
    # 4 L × R$ 120 + 2 L × R$ 80 + 300 kg × R$ 3 = R$ 1.540,00, em 40 ha.
    assert (norte_hoje["custo_total"], norte_hoje["custo_por_ha"]) == (1540, 38.5)
    sul_hoje = lista["Talhão Sul"]
    assert (sul_hoje["em_carencia"], sul_hoje["custo_total"]) == (False, None)
    assert any("(Adubo foliar)" in t and "carência" in t for t in sul_hoje["limitacoes"])
    assert any("1 aplicação sem preço" in t for t in sul_hoje["limitacoes"])
    assert ok(api.get("/resumo"))["talhoes_em_carencia"] == 1

    # A safra atual: só os últimos 90 dias. A carência não muda com o recorte.
    safra = ok(api.get(f"/talhoes/{norte}", params={"desde": dia(90), "ate": dia(0)}))
    # 4 × 120 + 2 × 80 + 100 × 3 = R$ 940,00 / 40 ha.
    assert (safra["custo_total"], safra["custo_por_ha"]) == (940, 23.5)
    assert [a["data"] for a in safra["aplicacoes"]] == [dia(1), dia(8), dia(10)]
    assert [(c["insumo"], c["quantidade"], c["custo"]) for c in safra["consumo"]] == [
        ("Fungicida X", 4, 480),
        ("Ureia", 100, 300),
        ("Inseticida Y", 2, 160),
    ]
    assert safra["colheita_liberada_em"] == norte_hoje["colheita_liberada_em"]
    invertido = api.get("/talhoes", params={"desde": dia(0), "ate": dia(5)})
    assert invertido.status_code == 422

    # Remover o talhão com aplicações perderia a rastreabilidade.
    recusa = api.delete(f"/talhoes/{norte}")
    assert recusa.status_code == 409
    assert recusa.json()["detail"].startswith("Talhão Norte tem 4 aplicações registradas")
    assert ok(api.get(f"/talhoes/{norte}"))["nome"] == "Talhão Norte"

    # O relatório leva os talhões, com a carência em destaque.
    livro = load_workbook(BytesIO(api.get("/relatorio.xlsx").content))
    linhas = {linha[0]: linha for linha in livro["Talhões"].iter_rows(min_row=2, values_only=True)}
    assert linhas["Talhão Norte"][3] == "Urgente"
    assert linhas["Talhão Norte"][5:7] == (1540, 38.5)
    assert linhas["Talhão Norte"][7] == alerta["justificativa"]
    assert linhas["Talhão Sul"][3:7] == ("Sem alertas", None, "sem preço", "sem preço")


def test_restauracao_da_copia_de_seguranca_traz_as_aplicacoes(
    api: httpx.Client, banco: Engine
) -> None:
    """Perdeu tudo: com os talhões cadastrados de novo (a cópia diz quais faltam), o
    relatório importado num CampoCerto vazio devolve cada aplicação ao seu talhão."""
    talhoes, insumos = fazenda(api)
    aplicar(api, insumos["Fungicida X"], talhoes["Talhão Norte"], 10, 4)
    aplicar(api, insumos["Ureia"], talhoes["Talhão Sul"], 8, 100)
    lancar(api, insumo_id=insumos["Ureia"], tipo="saida", data=dia(3), quantidade=7)
    originais = {t["nome"]: t for t in ok(api.get("/talhoes"))}
    relatorio = base64.b64encode(api.get("/relatorio.xlsx").content).decode()

    def aplicacoes() -> list[tuple]:
        nomes = {t["id"]: t["nome"] for t in ok(api.get("/talhoes"))}
        return sorted(
            (m["data"], m["quantidade"], nomes.get(m["talhao_id"]))
            for m in ok(api.get("/movimentacoes", params={"limite": 1000}))
            if m["tipo"] == "saida"
        )

    antes = aplicacoes()
    # Perdeu tudo (ou é um CampoCerto novo): o banco vazio. Pela API não dá para chegar
    # aqui, de propósito: insumo e talhão com aplicação não se removem.
    tabelas = ", ".join(t.name for t in Base.metadata.sorted_tables)
    with banco.begin() as conexao:
        conexao.execute(text(f"TRUNCATE {tabelas} RESTART IDENTITY CASCADE"))
    assert ok(api.get("/insumos")) == ok(api.get("/talhoes")) == []

    sem_talhoes = ok(api.post("/importacoes", json={"xlsx": relatorio, "confirmar": True}))
    assert sem_talhoes["importado"] is False
    assert sorted(e.split(": ", 1)[1] for e in sem_talhoes["erros"]) == [
        'talhão "Talhão Norte" não cadastrado; cadastre na tela de Talhões.',
        'talhão "Talhão Sul" não cadastrado; cadastre na tela de Talhões.',
    ]
    for nome, cadastro in TALHOES.items():
        ok(api.post("/talhoes", json={"nome": nome} | cadastro), 201)
    resultado = ok(api.post("/importacoes", json={"xlsx": relatorio, "confirmar": True}))

    assert (resultado["erros"], resultado["importado"]) == ([], True)
    assert aplicacoes() == antes
    # Com a carência relançada no cadastro (a cópia não leva), a análise do talhão volta igual.
    for insumo in ok(api.get("/insumos")):
        cadastro = {"nome": insumo["nome"]} | INSUMOS[insumo["nome"]]
        ok(api.put(f"/insumos/{insumo['id']}", json=cadastro))
    for restaurado in ok(api.get("/talhoes")):
        original = originais[restaurado["nome"]]
        assert restaurado | {"id": 0} == original | {"id": 0}


def test_erros_de_talhao_respondem_422(api: httpx.Client) -> None:
    talhoes, insumos = fazenda(api)
    base = {"insumo_id": insumos["Ureia"], "data": dia(0), "quantidade": 1}

    for corpo in (
        {"tipo": "entrada", "talhao_id": talhoes["Talhão Norte"]},
        {"tipo": "saida", "motivo": "descarte", "talhao_id": talhoes["Talhão Norte"]},
        {"tipo": "saida", "talhao_id": 999},
    ):
        resposta = api.post("/movimentacoes", json=base | corpo)
        assert resposta.status_code == 422, corpo
    inexistente = api.post("/movimentacoes", json=base | {"tipo": "saida", "talhao_id": 999})
    assert inexistente.json()["detail"] == "Talhão não encontrado."
    for cadastro in ({"area_ha": 0}, {"area_ha": -1}, {}):
        resposta = api.post("/talhoes", json={"nome": "Talhão Leste"} | cadastro)
        assert resposta.status_code == 422, cadastro
    for carencia in (-1, 366):
        resposta = api.post(
            "/insumos", json={"nome": "Outro", "unidade": "L", "carencia_dias": carencia}
        )
        assert resposta.status_code == 422
    assert ok(api.get(f"/talhoes/{talhoes['Talhão Norte']}"))["aplicacoes"] == []
    assert api.get("/talhoes/999").status_code == 404
