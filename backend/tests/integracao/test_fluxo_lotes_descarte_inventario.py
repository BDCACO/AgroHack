"""Fluxo real de uso, por HTTP: o produtor de gado com um lote de ivermectina que venceu
na prateleira. Lança compras e retiradas, vê o alerta certo, descarta o vencido, conta o
galpão, baixa o relatório e o importa de volta, como faria pelo site ou pelo app."""

import base64
from datetime import timedelta
from io import BytesIO
from typing import Any

import httpx
from openpyxl import load_workbook

from app.core.relogio import hoje

HOJE = hoje()


def dia(dias_atras: int) -> str:
    return str(HOJE - timedelta(days=dias_atras))


def ok(resposta: httpx.Response, status: int = 200) -> Any:
    assert resposta.status_code == status, resposta.text
    return resposta.json() if resposta.content else None


def lancar(api: httpx.Client, **corpo: Any) -> dict[str, Any]:
    return ok(api.post("/movimentacoes", json=corpo), 201)


def alertas(insumo: dict[str, Any]) -> dict[str, str]:
    return {a["tipo"]: a["justificativa"] for a in insumo["analise"]["alertas"]}


def test_lote_vencido_descarte_e_inventario_de_ponta_a_ponta(api: httpx.Client) -> None:
    insumo = ok(
        api.post(
            "/insumos",
            json={
                "nome": "Ivermectina 1%",
                "unidade": "L",
                "estoque_minimo": 2,
                "prazo_reposicao_dias": 10,
            },
        ),
        201,
    )
    id_ = insumo["id"]
    # Lote A vence há 52 dias; lote B, novo, chegou antes disso.
    lancar(api, insumo_id=id_, tipo="entrada", data=dia(120), quantidade=10, validade=dia(52))
    lancar(api, insumo_id=id_, tipo="entrada", data=dia(60), quantidade=10, validade=dia(-300))
    # Oito retiradas semanais, todas depois do vencimento do lote A.
    for semana in range(8):
        lancar(api, insumo_id=id_, tipo="saida", data=dia(50 - 7 * semana), quantidade=1)

    detalhe = ok(api.get(f"/insumos/{id_}"))
    assert detalhe["estoque_atual"] == 12
    # O vencido está inteiro na prateleira: as retiradas saíram do lote bom.
    assert alertas(detalhe)["vencido"].startswith("10 L vencidos")
    consumo_antes = detalhe["analise"]["consumo_diario"]

    # O produtor devolve o lote vencido ao fornecedor.
    descarte = lancar(
        api, insumo_id=id_, tipo="saida", motivo="descarte", data=dia(0), quantidade=10
    )
    assert descarte["motivo"] == "descarte"
    detalhe = ok(api.get(f"/insumos/{id_}"))
    assert detalhe["estoque_atual"] == 2
    assert "vencido" not in alertas(detalhe)
    assert detalhe["analise"]["consumo_diario"] == consumo_antes  # descarte não é consumo
    # Sobram 2 L bons, no mínimo: o produtor precisa comprar, e a análise diz.
    assert detalhe["analise"]["reposicao"] is not None

    # Contou o galpão e achou 1,5 L: meio litro saiu sem registro.
    ajuste = ok(api.post(f"/insumos/{id_}/inventario", json={"quantidade_contada": 1.5}))
    assert (ajuste["saldo_registrado"], ajuste["diferenca"]) == (2, -0.5)
    assert ajuste["movimentacao"]["tipo"] == "saida"
    detalhe = ok(api.get(f"/insumos/{id_}"))
    assert detalhe["estoque_atual"] == 1.5
    assert any("ajuste de inventário" in t for t in detalhe["analise"]["limitacoes"])
    assert detalhe["analise"]["consumo_diario"] == consumo_antes
    # Recontar o mesmo número não lança nada.
    assert (
        ok(api.post(f"/insumos/{id_}/inventario", json={"quantidade_contada": 1.5}))["movimentacao"]
        is None
    )
    assert len(detalhe["movimentacoes"]) == 12

    # A cópia de segurança leva descarte e falta e volta igual, sem nada a lançar.
    relatorio = api.get("/relatorio.xlsx")
    assert relatorio.status_code == 200
    aba = load_workbook(BytesIO(relatorio.content))["Movimentações"]
    tipos = [linha[3] for linha in aba.iter_rows(min_row=2, values_only=True)]
    assert sorted(set(tipos)) == ["descarte", "entrada", "falta", "saída"]
    previa = ok(
        api.post("/importacoes", json={"xlsx": base64.b64encode(relatorio.content).decode()})
    )
    assert (previa["erros"], previa["linhas"], previa["ja_lancadas"]) == ([], 12, 12)


def test_restauracao_da_copia_de_seguranca_num_banco_vazio(api: httpx.Client) -> None:
    """Perdeu tudo: o relatório importado num CampoCerto vazio refaz estoque e alertas."""
    planilha = (
        "data;insumo;unidade;tipo;quantidade;validade\n"
        f"{HOJE - timedelta(days=90):%d/%m/%Y};Sal mineral;kg;entrada;500;\n"
        f"{HOJE - timedelta(days=60):%d/%m/%Y};Sal mineral;kg;saída;100;\n"
        f"{HOJE - timedelta(days=30):%d/%m/%Y};Sal mineral;kg;descarte;20;\n"
        f"{HOJE - timedelta(days=10):%d/%m/%Y};Sal mineral;kg;sobra;5;\n"
        f"{HOJE - timedelta(days=5):%d/%m/%Y};Sal mineral;kg;falta;3;\n"
    )
    ok(api.post("/importacoes", json={"conteudo": planilha, "confirmar": True}))
    original = ok(api.get("/insumos"))[0]
    relatorio = api.get("/relatorio.xlsx").content

    ok(api.delete(f"/insumos/{original['id']}"), 204)
    assert ok(api.get("/insumos")) == []
    resultado = ok(
        api.post(
            "/importacoes",
            json={"xlsx": base64.b64encode(relatorio).decode(), "confirmar": True},
        )
    )

    assert resultado["importado"] is True
    [restaurado] = ok(api.get("/insumos"))
    assert restaurado["estoque_atual"] == original["estoque_atual"] == 382
    assert restaurado["analise"] == original["analise"]
    motivos = sorted((m["tipo"], m["motivo"] or "") for m in ok(api.get("/movimentacoes")))
    assert motivos == [
        ("entrada", ""),
        ("entrada", "inventario"),
        ("saida", ""),
        ("saida", "descarte"),
        ("saida", "inventario"),
    ]


def test_erros_de_uso_respondem_com_o_status_certo(api: httpx.Client) -> None:
    insumo = ok(api.post("/insumos", json={"nome": "Ureia", "unidade": "kg"}), 201)
    base = {"insumo_id": insumo["id"], "data": dia(0), "quantidade": 1}

    assert (
        api.post(
            "/movimentacoes", json=base | {"tipo": "entrada", "motivo": "descarte"}
        ).status_code
        == 422
    )
    assert (
        api.post(
            "/movimentacoes", json=base | {"tipo": "saida", "motivo": "inventario"}
        ).status_code
        == 422
    )
    assert api.post("/insumos/999/inventario", json={"quantidade_contada": 1}).status_code == 404
    assert (
        api.post(
            f"/insumos/{insumo['id']}/inventario", json={"quantidade_contada": 1, "data": dia(-1)}
        ).status_code
        == 422
    )
    assert ok(api.get("/movimentacoes")) == []
