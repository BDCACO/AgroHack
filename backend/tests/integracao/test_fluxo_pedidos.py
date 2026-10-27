"""Fluxo real de uso, por HTTP: o ciclo da compra, sugerido → pedido → recebido. O dono
cadastra as revendas, vê a sugestão de pedido, faz o pedido; a análise passa a descontar o
que está a caminho (sem mandar comprar de novo); o funcionário confere a carga no galpão e
dá a entrada em duas vezes; o preço pago a cada revenda fica lado a lado. E o relatório,
com a aba Pedidos e a cópia de segurança que devolve de quem foi cada compra."""

import base64
import subprocess
import sys
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta
from io import BytesIO
from typing import Any

import httpx
from openpyxl import load_workbook
from sqlalchemy import Engine, text

from app.core.relogio import hoje
from app.core.senha import gerar_hash
from tests.integracao.conftest import BACKEND, ambiente_do_entrypoint

HOJE = hoje()
SENHA = "senha-do-fluxo-1"


def dia(dias: int) -> str:
    """Dias a partir de hoje: negativo é passado."""
    return str(HOJE + timedelta(days=dias))


def ok(resposta: httpx.Response, status: int = 200) -> Any:
    assert resposta.status_code == status, resposta.text
    return resposta.json() if resposta.content else None


def risco(insumo: dict[str, Any]) -> dict[str, Any]:
    return next(a for a in insumo["analise"]["alertas"] if a["tipo"] == "risco_de_falta")


def ureia_que_vai_faltar(api: httpx.Client, cabecalho: dict[str, str], revendas: dict) -> int:
    """Saco de 50 kg, mínimo 300, entrega em 10 dias. Comprada na cooperativa (R$ 3,00) e,
    por último, na revenda (R$ 3,60); 30 retiradas de 61 kg nos últimos 90 dias: restam
    170 kg a 20,33 kg/dia, que duram 8 dias. A conta da reposição dá 910 kg: 19 sacos."""
    ureia = ok(
        api.post(
            "/insumos",
            headers=cabecalho,
            json={
                "nome": "Ureia agrícola",
                "unidade": "kg",
                "estoque_minimo": 300,
                "prazo_reposicao_dias": 10,
                "embalagem": 50,
            },
        ),
        201,
    )["id"]
    compra = {"insumo_id": ureia, "tipo": "entrada", "quantidade": 1000}
    for data, preco, de in (
        (dia(-100), 3, "Cooperativa Vale Verde"),
        (dia(-40), 3.6, "Revenda Boa Safra"),
    ):
        corpo = compra | {"data": data, "valor_unitario": preco, "fornecedor_id": revendas[de]}
        ok(api.post("/movimentacoes", headers=cabecalho, json=corpo), 201)
    for retirada in range(30):
        saida = {"insumo_id": ureia, "tipo": "saida", "data": dia(-89 + 3 * retirada)}
        ok(api.post("/movimentacoes", headers=cabecalho, json=saida | {"quantidade": 61}), 201)
    return ureia


def cadastrar_revendas(api: httpx.Client, cabecalho: dict[str, str]) -> dict[str, int]:
    return {
        nome: ok(
            api.post("/fornecedores", headers=cabecalho, json={"nome": nome, "telefone": fone}),
            201,
        )["id"]
        for nome, fone in (
            ("Revenda Boa Safra", "(69) 90000-0001"),
            ("Cooperativa Vale Verde", "(69) 90000-0002"),
        )
    }


def test_do_pedido_sugerido_a_carga_recebida_pelo_funcionario(
    api_autenticada: httpx.Client, banco: Engine
) -> None:
    api = api_autenticada
    with banco.begin() as conexao:
        for nome, login, papel in (("Dona Maria", "maria", "dono"), ("Zé", "ze", "funcionario")):
            conexao.execute(
                text(
                    "INSERT INTO usuarios (nome, login, senha_hash, papel)"
                    " VALUES (:nome, :login, :hash, :papel)"
                ),
                {"nome": nome, "login": login, "hash": gerar_hash(SENHA), "papel": papel},
            )

    def entrar(login: str) -> dict[str, str]:
        token = ok(api.post("/auth/entrar", json={"login": login, "senha": SENHA}))["token"]
        return {"Authorization": f"Bearer {token}"}

    dono, ze = entrar("maria"), entrar("ze")
    revendas = cadastrar_revendas(api, dono)
    # O cadastro de fornecedor é do dono; consultar, o funcionário também pode.
    negado = api.post("/fornecedores", headers=ze, json={"nome": "Outra"})
    assert negado.status_code == 403
    assert len(ok(api.get("/fornecedores", headers=ze))) == 2
    ureia = ureia_que_vai_faltar(api, dono, revendas)
    antes = ok(api.get(f"/insumos/{ureia}", headers=dono))
    assert (antes["estoque_atual"], antes["em_pedido"]) == (170, 0)
    assert antes["analise"]["reposicao"]["pedir_hoje"] is True
    assert risco(antes)["gravidade"] == "alta"

    # A sugestão manda a ureia para a revenda da última compra, ao preço dela.
    sugestao = ok(api.get("/pedidos/sugestao", headers=ze))
    [grupo] = sugestao["grupos"]
    assert grupo["fornecedor"] == "Revenda Boa Safra"
    [item] = grupo["itens"]
    assert (item["quantidade"], item["valor_unitario"], item["custo_estimado"]) == (950, 3.6, 3420)
    assert item["justificativa"].startswith("Pedir hoje 19 embalagens de 50 kg (950 kg)")

    # O dono confirma; o funcionário não pode fazer o pedido.
    corpo = {"fornecedor_id": grupo["fornecedor_id"], "previsto_para": dia(5)}
    corpo["itens"] = grupo["itens"]
    assert api.post("/pedidos", headers=ze, json=corpo).status_code == 403
    pedido = ok(api.post("/pedidos", headers=dono, json=corpo), 201)
    assert (pedido["situacao"], pedido["feito_por"], pedido["custo_estimado"]) == (
        "aberto",
        "Dona Maria",
        3420,
    )

    # A análise sabe do pedido: não manda comprar de novo, e o risco deixa de ser urgente
    # (os 950 kg chegam em 5 dias, antes de o estoque acabar), sem esconder o estoque baixo.
    a_caminho = ok(api.get(f"/insumos/{ureia}", headers=ze))
    assert (a_caminho["estoque_atual"], a_caminho["em_pedido"]) == (170, 950)
    assert a_caminho["chegada_prevista"] == dia(5)
    assert a_caminho["analise"]["reposicao"]["pedir_hoje"] is False
    assert "Desconta os 950 kg já pedidos" in a_caminho["analise"]["reposicao"]["justificativa"]
    assert risco(a_caminho)["gravidade"] == "media"
    assert "Já foram pedidos 950 kg." in risco(a_caminho)["justificativa"]
    assert ok(api.get("/pedidos/sugestao", headers=dono))["grupos"] == []
    assert ok(api.get("/resumo", headers=ze))["compras_sugeridas"]["pedir_hoje"] == 0

    # Chega a primeira parte: o Zé confere e dá a entrada, com a validade do lote.
    linha = pedido["itens"][0]["id"]
    primeira = ok(
        api.post(
            f"/pedidos/{pedido['id']}/receber",
            headers=ze,
            json={"itens": [{"item_id": linha, "quantidade": 500, "validade": dia(400)}]},
        )
    )
    assert primeira["pedido"]["situacao"] == "parcial"
    [entrada] = primeira["movimentacoes"]
    assert (entrada["comprado_de"], entrada["valor_unitario"], entrada["validade"]) == (
        "Revenda Boa Safra",
        3.6,
        dia(400),
    )
    assert (entrada["lancado_por"], entrada["data"]) == ("Zé", dia(0))
    with banco.connect() as conexao:
        gravada = conexao.execute(
            text(
                "SELECT f.nome, m.item_do_pedido_id, u.login, i.quantidade_recebida"
                " FROM movimentacoes m JOIN fornecedores f ON f.id = m.fornecedor_id"
                " JOIN usuarios u ON u.id = m.usuario_id"
                " JOIN itens_do_pedido i ON i.id = m.item_do_pedido_id WHERE m.id = :id"
            ),
            {"id": entrada["id"]},
        ).one()
    assert tuple(gravada) == ("Revenda Boa Safra", linha, "ze", 500)
    parcial = ok(api.get(f"/insumos/{ureia}", headers=ze))
    # (1000 × 3 + 1000 × 3,6 + 500 × 3,6) / 2500 = 3,36.
    assert (parcial["estoque_atual"], parcial["em_pedido"], parcial["custo_medio"]) == (
        670,
        450,
        3.36,
    )
    # O pedido com parte recebida não se altera; o funcionário também não cancela.
    alterar = api.put(f"/pedidos/{pedido['id']}", headers=dono, json=corpo)
    assert alterar.status_code == 409
    assert api.post(f"/pedidos/{pedido['id']}/cancelar", headers=ze).status_code == 403

    # O resto chega com a nota mais cara.
    resto = {"item_id": linha, "quantidade": 450, "valor_unitario": 3.7}
    segunda = ok(api.post(f"/pedidos/{pedido['id']}/receber", headers=ze, json={"itens": [resto]}))
    assert segunda["pedido"]["situacao"] == "recebido"
    assert segunda["pedido"]["itens"][0]["falta_receber"] == 0
    recebida = ok(api.get(f"/insumos/{ureia}", headers=dono))
    assert (recebida["estoque_atual"], recebida["em_pedido"]) == (1120, 0)
    # O pedido que chegou inteiro não se cancela, nem se altera, nem recebe de novo.
    assert api.post(f"/pedidos/{pedido['id']}/cancelar", headers=dono).status_code == 409
    assert api.put(f"/pedidos/{pedido['id']}", headers=dono, json=corpo).status_code == 409
    de_novo = api.post(f"/pedidos/{pedido['id']}/receber", headers=ze, json={"itens": [resto]})
    assert de_novo.status_code == 409

    # O preço pago a cada revenda, lado a lado: (1000 × 3,6 + 500 × 3,6 + 450 × 3,7) / 1950.
    precos = recebida["precos_por_fornecedor"]
    assert [(p["fornecedor"], p["preco_medio"], p["ultimo_preco"]) for p in precos] == [
        ("Cooperativa Vale Verde", 3, 3),
        ("Revenda Boa Safra", 3.62, 3.7),
    ]
    revenda = ok(api.get(f"/fornecedores/{revendas['Revenda Boa Safra']}", headers=ze))
    assert [(p["insumo"], p["compras"], p["quantidade"]) for p in revenda["precos"]] == [
        ("Ureia agrícola", 3, 1950)
    ]
    assert [m["quantidade"] for m in revenda["compras"]] == [450, 500, 1000]
    # Com compra, a revenda não se apaga: desativa.
    assert (
        api.delete(f"/fornecedores/{revendas['Revenda Boa Safra']}", headers=dono).status_code
        == 409
    )


def test_cancelar_e_os_erros_de_uso(api: httpx.Client) -> None:
    revendas = cadastrar_revendas(api, {})
    ureia = ureia_que_vai_faltar(api, {}, revendas)
    item = {"insumo_id": ureia, "quantidade": 950, "valor_unitario": 3.6}
    pedido = ok(api.post("/pedidos", json={"previsto_para": dia(5), "itens": [item]}), 201)
    linha = pedido["itens"][0]["id"]

    cancelado = ok(api.post(f"/pedidos/{pedido['id']}/cancelar"))

    assert cancelado["situacao"] == "cancelado"
    # O cancelado sai da análise: a reposição volta a mandar pedir hoje.
    insumo = ok(api.get(f"/insumos/{ureia}"))
    assert (insumo["em_pedido"], insumo["analise"]["reposicao"]["pedir_hoje"]) == (0, True)
    recusa = api.post(
        f"/pedidos/{pedido['id']}/receber", json={"itens": [{"item_id": linha, "quantidade": 1}]}
    )
    assert recusa.status_code == 409
    assert "cancelado" in recusa.json()["detail"]
    assert api.put(f"/pedidos/{pedido['id']}", json={"itens": [item]}).status_code == 409

    outro = ok(api.post("/pedidos", json={"itens": [item]}), 201)
    for corpo in (
        {"itens": [{"item_id": outro["itens"][0]["id"], "quantidade": 1}], "data": dia(1)},
        {"itens": [{"item_id": linha, "quantidade": 1}]},  # linha de outro pedido
        {"itens": [{"item_id": outro["itens"][0]["id"], "quantidade": 0}]},
    ):
        assert api.post(f"/pedidos/{outro['id']}/receber", json=corpo).status_code == 422, corpo
    for corpo in (
        {"itens": [item], "feito_em": dia(1)},
        {"itens": [item], "feito_em": dia(-1), "previsto_para": dia(-2)},
        {"itens": [item, item]},
        {"itens": [item], "fornecedor_id": 999_999},
        {"itens": [item | {"insumo_id": 999_999}]},
    ):
        assert api.post("/pedidos", json=corpo).status_code == 422, corpo
    # Fornecedor só em compra: na retirada e no descarte é 422; inexistente também.
    base = {"insumo_id": ureia, "data": dia(0), "quantidade": 1}
    for corpo in (
        {"tipo": "saida", "fornecedor_id": revendas["Revenda Boa Safra"]},
        {"tipo": "saida", "motivo": "descarte", "fornecedor_id": revendas["Revenda Boa Safra"]},
        {"tipo": "entrada", "fornecedor_id": 999_999},
    ):
        assert api.post("/movimentacoes", json=base | corpo).status_code == 422, corpo
    assert api.get("/pedidos/999999").status_code == 404
    assert [p["situacao"] for p in ok(api.get("/pedidos"))] == ["aberto", "cancelado"]


def test_recebimentos_simultaneos_entram_um_de_cada_vez(api: httpx.Client) -> None:
    """Dois celulares dão a entrada da mesma carga ao mesmo tempo: a trava no pedido faz um
    esperar o outro, e o recebido nunca se perde nem passa do que chegou."""
    insumo = ok(api.post("/insumos", json={"nome": "Sal mineral", "unidade": "kg"}), 201)["id"]
    pedido = ok(
        api.post("/pedidos", json={"itens": [{"insumo_id": insumo, "quantidade": 100}]}), 201
    )
    corpo = {"itens": [{"item_id": pedido["itens"][0]["id"], "quantidade": 20}]}

    def receber(_: int) -> int:
        with httpx.Client(base_url=str(api.base_url), timeout=30) as cliente:
            return cliente.post(f"/pedidos/{pedido['id']}/receber", json=corpo).status_code

    with ThreadPoolExecutor(max_workers=8) as threads:
        respostas = list(threads.map(receber, range(8)))

    # Cinco cargas de 20 kg completam os 100; as outras três encontram o pedido recebido.
    assert sorted(respostas) == [200] * 5 + [409] * 3
    final = ok(api.get(f"/pedidos/{pedido['id']}"))
    assert (final["situacao"], final["itens"][0]["quantidade_recebida"]) == ("recebido", 100)
    assert ok(api.get(f"/insumos/{insumo}"))["estoque_atual"] == 100


def test_relatorio_com_pedidos_e_copia_que_devolve_os_fornecedores(api: httpx.Client) -> None:
    revendas = cadastrar_revendas(api, {})
    ureia = ureia_que_vai_faltar(api, {}, revendas)
    pedido = ok(
        api.post(
            "/pedidos",
            json={
                "fornecedor_id": revendas["Revenda Boa Safra"],
                "feito_em": dia(-2),
                "previsto_para": dia(5),
                "itens": [{"insumo_id": ureia, "quantidade": 950, "valor_unitario": 3.6}],
            },
        ),
        201,
    )
    linha = pedido["itens"][0]["id"]
    ok(
        api.post(
            f"/pedidos/{pedido['id']}/receber",
            json={"itens": [{"item_id": linha, "quantidade": 300}]},
        )
    )
    compras = {
        (m["data"], m["quantidade"], m["comprado_de"])
        for m in ok(api.get("/movimentacoes", params={"limite": 1000}))
        if m["tipo"] == "entrada"
    }

    relatorio = api.get("/relatorio.xlsx")
    assert relatorio.status_code == 200
    livro = load_workbook(BytesIO(relatorio.content))

    [linha_do_pedido] = list(livro["Pedidos"].iter_rows(min_row=2, values_only=True))
    meia_noite = datetime.min.time()
    assert linha_do_pedido == (
        pedido["id"],
        datetime.combine(HOJE - timedelta(days=2), meia_noite),
        "Revenda Boa Safra",
        datetime.combine(HOJE + timedelta(days=5), meia_noite),
        "Ureia agrícola",
        "kg",
        950,
        300,
        650,
        3.6,
        "Parcial",
    )
    na_copia = {linha[8] for linha in livro["Movimentações"].iter_rows(min_row=2, values_only=True)}
    assert na_copia == {"Revenda Boa Safra", "Cooperativa Vale Verde", None}

    # Perdeu tudo: remove o insumo (o pedido vazio sai junto) e as revendas, e restaura.
    ok(api.delete(f"/insumos/{ureia}"), 204)
    assert ok(api.get("/pedidos")) == []
    for fornecedor_id in revendas.values():
        ok(api.delete(f"/fornecedores/{fornecedor_id}"), 204)
    xlsx = base64.b64encode(relatorio.content).decode()
    previa = ok(api.post("/importacoes", json={"xlsx": xlsx}))
    assert sorted(previa["fornecedores_novos"]) == ["Cooperativa Vale Verde", "Revenda Boa Safra"]
    assert ok(api.post("/importacoes", json={"xlsx": xlsx, "confirmar": True}))["importado"]

    restauradas = {
        (m["data"], m["quantidade"], m["comprado_de"])
        for m in ok(api.get("/movimentacoes", params={"limite": 1000}))
        if m["tipo"] == "entrada"
    }
    assert restauradas == compras
    assert sorted(f["nome"] for f in ok(api.get("/fornecedores"))) == [
        "Cooperativa Vale Verde",
        "Revenda Boa Safra",
    ]
    # A mesma cópia de novo: tudo já está lá, com os mesmos fornecedores.
    de_novo = ok(api.post("/importacoes", json={"xlsx": xlsx}))
    assert (de_novo["ja_lancadas"], de_novo["fornecedores_novos"]) == (de_novo["linhas"], [])


def test_aba_pedidos_com_mais_de_200_a_caminho(api: httpx.Client, banco: Engine) -> None:
    # 201 pedidos abertos e 1 parcial, gravados direto no banco (pela API seriam 202
    # pedidos HTTP só de preparo). A lista da tela para em 200 por situação; a planilha
    # traz todos, e o mais antigo (o mais atrasado) também.
    ureia = ok(api.post("/insumos", json={"nome": "Ureia", "unidade": "kg"}), 201)["id"]
    with banco.begin() as conexao:
        conexao.execute(
            text(
                "INSERT INTO pedidos (situacao, feito_em, previsto_para)"
                " SELECT CASE WHEN n = 0 THEN 'parcial' ELSE 'aberto' END,"
                " CURRENT_DATE - 400 + n, CURRENT_DATE - 300 + n"
                " FROM generate_series(0, 201) AS n"
            )
        )
        conexao.execute(
            text(
                "INSERT INTO itens_do_pedido (pedido_id, insumo_id, quantidade,"
                " quantidade_recebida)"
                " SELECT id, :ureia, 10, CASE WHEN situacao = 'parcial' THEN 4 ELSE 0 END"
                " FROM pedidos"
            ),
            {"ureia": ureia},
        )
        mais_antigo = conexao.execute(
            text("SELECT id FROM pedidos WHERE situacao = 'aberto' ORDER BY feito_em LIMIT 1")
        ).scalar_one()
    assert len(ok(api.get("/pedidos", params={"situacao": "aberto"}))) == 200

    relatorio = api.get("/relatorio.xlsx")

    assert relatorio.status_code == 200
    linhas = list(
        load_workbook(BytesIO(relatorio.content))["Pedidos"].iter_rows(min_row=2, values_only=True)
    )
    assert len(linhas) == 202
    assert mais_antigo in {linha[0] for linha in linhas}
    assert sorted(linha[8] for linha in linhas)[0] == 6  # o parcial: faltam 6 dos 10


def _pedido_meio_recebido(
    api: httpx.Client, real: int, outro: int, recebe: int, quantidade: float
) -> dict[str, Any]:
    """Pedido atrasado com duas linhas de 10, das quais `recebe` recebe `quantidade`."""
    feito = ok(
        api.post(
            "/pedidos",
            json={
                "feito_em": dia(-10),
                "previsto_para": dia(-2),
                "itens": [
                    {"insumo_id": real, "quantidade": 10},
                    {"insumo_id": outro, "quantidade": 10},
                ],
            },
        ),
        201,
    )
    linha = next(i["id"] for i in feito["itens"] if i["insumo_id"] == recebe)
    recebido = ok(
        api.post(
            f"/pedidos/{feito['id']}/receber",
            json={"itens": [{"item_id": linha, "quantidade": quantidade}]},
        )
    )
    assert recebido["pedido"]["situacao"] == "parcial"
    return feito


def _conferir_o_que_sobrou(
    api: httpx.Client, chegou_o_resto: dict[str, Any], nada_chegou: dict[str, Any], real: int
) -> None:
    recebido = ok(api.get(f"/pedidos/{chegou_o_resto['id']}"))
    assert (recebido["situacao"], recebido["atrasado"]) == ("recebido", False)
    [linha] = recebido["itens"]
    de_novo = api.post(
        f"/pedidos/{chegou_o_resto['id']}/receber",
        json={"itens": [{"item_id": linha["id"], "quantidade": 1}]},
    )
    assert de_novo.status_code == 409
    assert ok(api.get(f"/pedidos/{nada_chegou['id']}"))["situacao"] == "aberto"
    ok(
        api.put(
            f"/pedidos/{nada_chegou['id']}",
            json={"itens": [{"insumo_id": real, "quantidade": 12}]},
        )
    )


def test_remover_insumo_recalcula_a_situacao_do_pedido_que_fica(api: httpx.Client) -> None:
    adubo = ok(api.post("/insumos", json={"nome": "Adubo", "unidade": "kg"}), 201)["id"]
    calcario = ok(api.post("/insumos", json={"nome": "Calcário", "unidade": "t"}), 201)["id"]
    sal = ok(api.post("/insumos", json={"nome": "Sal mineral", "unidade": "kg"}), 201)["id"]
    chegou_o_resto = _pedido_meio_recebido(api, adubo, sal, recebe=adubo, quantidade=10)
    nada_chegou = _pedido_meio_recebido(api, calcario, sal, recebe=sal, quantidade=4)

    ok(api.delete(f"/insumos/{sal}"), 204)

    _conferir_o_que_sobrou(api, chegou_o_resto, nada_chegou, calcario)


def test_limpar_a_demonstracao_recalcula_a_situacao_do_pedido_real(
    api: httpx.Client,
) -> None:
    # O comando de verdade, com o papel da aplicação, como o `make demo-limpar`.
    def demo(acao: str) -> None:
        rodou = subprocess.run(  # noqa: S603 (comando fixo)
            [sys.executable, "-m", "app.demo", acao],
            cwd=BACKEND,
            env=ambiente_do_entrypoint(),
            capture_output=True,
            text=True,
            check=False,
        )
        assert rodou.returncode == 0, rodou.stderr

    demo("carregar")
    sal = next(i["id"] for i in ok(api.get("/insumos")) if i["nome"] == "Sal mineral 80P")
    adubo = ok(api.post("/insumos", json={"nome": "Adubo da fazenda", "unidade": "kg"}), 201)["id"]
    calcario = ok(api.post("/insumos", json={"nome": "Calcário", "unidade": "t"}), 201)["id"]
    chegou_o_resto = _pedido_meio_recebido(api, adubo, sal, recebe=adubo, quantidade=10)
    nada_chegou = _pedido_meio_recebido(api, calcario, sal, recebe=sal, quantidade=4)

    demo("limpar")

    _conferir_o_que_sobrou(api, chegou_o_resto, nada_chegou, calcario)
