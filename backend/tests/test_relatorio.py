"""Relatório em Excel (GET /relatorio.xlsx), com a fazenda de demonstração."""

import base64
from datetime import datetime, timedelta
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.core.relogio import hoje
from app.demo import carregar


def baixar(client: TestClient) -> tuple[bytes, Workbook]:
    resposta = client.get("/api/v1/relatorio.xlsx")
    assert resposta.status_code == 200
    nome = f'filename="relatorio-campocerto-{hoje()}.xlsx"'
    assert nome in resposta.headers["content-disposition"]
    return resposta.content, load_workbook(BytesIO(resposta.content))


def test_relatorio_traz_alertas_reposicao_e_estoque(client: TestClient, db: Session) -> None:
    carregar(db, hoje())

    _, livro = baixar(client)

    assert livro.sheetnames == [
        "Resumo",
        "Alertas",
        "Reposição",
        "Pedidos",
        "Estoque",
        "Talhões",
        "Por mês",
        "Movimentações",
    ]
    alertas = list(livro["Alertas"].iter_rows(min_row=2, values_only=True))
    # Os mais graves primeiro, cada um com a conta que o justifica.
    assert [a[0] for a in alertas] == ["Urgente"] * 2 + ["Atenção"] * 3 + ["Observar"] * 2
    assert ("Urgente", "Ureia agrícola", "Risco de falta") in {a[:3] for a in alertas}
    assert all(a[3] for a in alertas)
    reposicao = list(livro["Reposição"].iter_rows(min_row=2, values_only=True))
    assert reposicao[0][:3] == ("Ureia agrícola", 700, "kg")  # 14 sacos de 50 kg
    estoque = {linha[0]: linha for linha in livro["Estoque"].iter_rows(min_row=2, values_only=True)}
    assert len(estoque) == 7
    assert estoque["Semente de milho híbrido"][9] == "Sem alertas"
    assert livro["Resumo"]["A1"].value == "Situação do estoque"


def test_relatorio_fala_em_reais(client: TestClient) -> None:
    ureia = client.post(
        "/api/v1/insumos",
        json={"nome": "Ureia", "unidade": "kg", "estoque_minimo": 10, "prazo_reposicao_dias": 15},
    ).json()

    def lancar(**corpo: object) -> None:
        resposta = client.post("/api/v1/movimentacoes", json={"insumo_id": ureia["id"]} | corpo)
        assert resposta.status_code == 201, resposta.json()

    lancar(tipo="entrada", data=str(hoje() - timedelta(days=30)), quantidade=80, valor_unitario=3)
    for atras in range(30):
        lancar(tipo="saida", data=str(hoje() - timedelta(days=atras)), quantidade=2)
    lancar(tipo="saida", motivo="descarte", data=str(hoje()), quantidade=5)
    [insumo] = client.get("/api/v1/insumos").json()

    _, livro = baixar(client)

    # 80 − 60 − 5 = 15 kg a R$ 3,00.
    estoque = list(livro["Estoque"].iter_rows(values_only=True))
    assert estoque[0][7:9] == ("custo médio (R$)", "valor em estoque (R$)")
    assert estoque[1][7:9] == (3, 45)
    assert livro["Estoque"]["I2"].number_format == '"R$" #,##0.00'
    reposicao = list(livro["Reposição"].iter_rows(values_only=True))
    assert reposicao[0][4] == "custo estimado (R$)"
    assert reposicao[1][4] == insumo["analise"]["reposicao"]["custo_estimado"]
    resumo = {
        linha[0]: linha[1] for linha in livro["Resumo"].iter_rows(values_only=True) if linha[0]
    }
    assert resumo["Valor em estoque"] == 45
    assert resumo["Perdas nos últimos 30 dias"] == 15
    assert resumo["Perdas no total"] == 15


def test_movimentacoes_do_relatorio_voltam_pela_importacao(client: TestClient, db: Session) -> None:
    carregar(db, hoje())
    [ivermectina] = [i for i in client.get("/api/v1/insumos").json() if "Ivermectina" in i["nome"]]
    descarte = {"insumo_id": ivermectina["id"], "tipo": "saida", "motivo": "descarte"}
    client.post("/api/v1/movimentacoes", json=descarte | {"data": str(hoje()), "quantidade": 4})
    agrovet = client.post("/api/v1/fornecedores", json={"nome": "AgroVet do Teste"}).json()
    compra = {
        "insumo_id": ivermectina["id"],
        "tipo": "entrada",
        "valor_unitario": 95.1234,
        "fornecedor_id": agrovet["id"],
    }
    client.post("/api/v1/movimentacoes", json=compra | {"data": str(hoje()), "quantidade": 2})
    insumos = {i["nome"]: i for i in client.get("/api/v1/insumos").json()}
    ureia, sal = insumos["Ureia agrícola"], insumos["Sal mineral 80P"]
    for insumo, diferenca in ((ureia, 5), (sal, -5)):
        contagem = {"quantidade_contada": insumo["estoque_atual"] + diferenca}
        client.post(f"/api/v1/insumos/{insumo['id']}/inventario", json=contagem)
    conteudo, livro = baixar(client)
    linhas = livro["Movimentações"].max_row - 1
    tipos = {linha[3] for linha in livro["Movimentações"].iter_rows(min_row=2, values_only=True)}
    assert tipos == {"entrada", "saída", "descarte", "sobra", "falta"}
    assert livro["Movimentações"]["G1"].value == "valor"
    precos = {linha[6] for linha in livro["Movimentações"].iter_rows(min_row=2, values_only=True)}
    assert 95.1234 in precos
    assert livro["Movimentações"]["I1"].value == "fornecedor"
    fornecedores = {
        linha[8] for linha in livro["Movimentações"].iter_rows(min_row=2, values_only=True)
    }
    assert "AgroVet do Teste" in fornecedores

    resposta = client.post(
        "/api/v1/importacoes", json={"xlsx": base64.b64encode(conteudo).decode()}
    )

    # É a cópia de segurança: tudo já está no sistema, e a prévia avisa.
    previa = resposta.json()
    assert (previa["erros"], previa["linhas"], previa["ja_lancadas"]) == ([], linhas, linhas)


def test_relatorio_sem_nada_cadastrado(client: TestClient) -> None:
    _, livro = baixar(client)

    assert livro["Alertas"]["A2"].value == "Nenhum alerta hoje."


def test_aba_pedidos_mostra_o_que_falta_chegar(client: TestClient) -> None:
    ureia = client.post("/api/v1/insumos", json={"nome": "Ureia", "unidade": "kg"}).json()
    sal = client.post("/api/v1/insumos", json={"nome": "Sal mineral", "unidade": "kg"}).json()
    revenda = client.post("/api/v1/fornecedores", json={"nome": "Revenda Boa Safra"}).json()
    atrasado = client.post(
        "/api/v1/pedidos",
        json={
            "fornecedor_id": revenda["id"],
            "feito_em": str(hoje() - timedelta(days=10)),
            "previsto_para": str(hoje() - timedelta(days=1)),
            "itens": [
                {"insumo_id": ureia["id"], "quantidade": 700, "valor_unitario": 3.3},
                {"insumo_id": sal["id"], "quantidade": 200},
            ],
        },
    ).json()
    linha_da_ureia = atrasado["itens"][0]["id"]
    client.post(
        f"/api/v1/pedidos/{atrasado['id']}/receber",
        json={"itens": [{"item_id": linha_da_ureia, "quantidade": 300}]},
    )
    recebido = client.post(
        "/api/v1/pedidos", json={"itens": [{"insumo_id": sal["id"], "quantidade": 5}]}
    ).json()
    client.post(
        f"/api/v1/pedidos/{recebido['id']}/receber",
        json={"itens": [{"item_id": recebido["itens"][0]["id"], "quantidade": 5}]},
    )

    _, livro = baixar(client)

    linhas = list(livro["Pedidos"].iter_rows(values_only=True))
    assert linhas[0] == (
        "pedido",
        "feito em",
        "fornecedor",
        "previsto para",
        "insumo",
        "unidade",
        "pedido (qtd)",
        "recebido",
        "falta chegar",
        "preço combinado (R$)",
        "situação",
    )
    # Só o que ainda falta chegar: o pedido recebido não aparece.
    feito_em = datetime.combine(hoje() - timedelta(days=10), datetime.min.time())
    previsto = datetime.combine(hoje() - timedelta(days=1), datetime.min.time())
    assert linhas[1:] == [
        (
            atrasado["id"],
            feito_em,
            "Revenda Boa Safra",
            previsto,
            "Sal mineral",
            "kg",
            200,
            0,
            200,
            None,
            "Atrasado",
        ),
        (
            atrasado["id"],
            feito_em,
            "Revenda Boa Safra",
            previsto,
            "Ureia",
            "kg",
            700,
            300,
            400,
            3.3,
            "Atrasado",
        ),
    ]
    assert livro["Pedidos"]["J3"].number_format == '"R$" #,##0.00'


def test_aba_pedidos_sem_pedido(client: TestClient) -> None:
    _, livro = baixar(client)

    assert livro["Pedidos"]["A2"].value == "Nenhum pedido a caminho."


def test_aba_por_mes_com_os_ultimos_doze_meses(client: TestClient, db: Session) -> None:
    carregar(db, hoje())
    periodo = client.get("/api/v1/relatorios/periodo").json()

    _, livro = baixar(client)

    aba = livro["Por mês"]
    assert [c.value for c in aba[1]] == [
        "mês",
        "gasto em compras (R$)",
        "compras sem preço",
        "consumo (R$)",
        "perdas (R$)",
    ]
    linhas = list(aba.iter_rows(min_row=2, max_row=13, values_only=True))
    # Os mesmos números do GET /relatorios/periodo: um mês por linha, sem buraco.
    assert [linha[0] for linha in linhas] == [p["rotulo"] for p in periodo["periodos"]]
    assert [linha[1] for linha in linhas] == [p["gasto"] for p in periodo["periodos"]]
    assert [linha[4] for linha in linhas] == [p["perdas"] for p in periodo["periodos"]]
    consumo = [round(sum(c["consumo"] for c in p["por_categoria"]), 2) for p in periodo["periodos"]]
    assert [linha[3] for linha in linhas] == consumo
    assert sum(linha[1] for linha in linhas) > 0
