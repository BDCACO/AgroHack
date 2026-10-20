"""Planilha modelo da importação (.xlsx), gerada na hora do download.

Gerada, e não um arquivo parado no repositório: já vem com os insumos cadastrados na
fazenda para escolher da lista, e os exemplos têm datas recentes. A formatação — datas
no padrão brasileiro, listas de escolha, avisos do próprio Excel — ajuda o produtor a
preencher certo antes mesmo de importar.
"""

import math
from dataclasses import dataclass
from datetime import date, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.relogio import hoje
from app.models import Insumo

# Cores da marca, as mesmas de frontend/src/styles/main.css.
VERDE = "3F7A2C"  # campo-600
VERDE_ESCURO = "1F3B17"  # campo-900
CINZA = "57534E"  # stone-600
CINZA_CLARO = "E7E5E4"  # stone-200
FORMATO_DATA = "DD/MM/YYYY"  # o Excel em português mostra como DD/MM/AAAA
UNIDADES = ("kg", "L", "t", "sc", "dose", "un")
# Linhas da aba Movimentações que já saem com formato de data e validação.
LINHAS_PRONTAS = 1000


@dataclass(frozen=True)
class Coluna:
    titulo: str
    largura: int
    obrigatoria: str
    explicacao: str


COLUNAS = (
    Coluna(
        "data",
        14,
        "Sim",
        "Dia em que a compra chegou ou em que o produto saiu do galpão. Não pode ser uma "
        "data futura.",
    ),
    Coluna(
        "insumo",
        34,
        "Sim",
        "Nome do produto. Escolha da lista um insumo já cadastrado, ou digite um nome novo: "
        "ele será cadastrado na importação.",
    ),
    Coluna(
        "unidade",
        11,
        "Só para insumo novo",
        "kg, L, t, sc, dose, un… Para um insumo já cadastrado, pode ficar em branco; se "
        "preenchida, tem de ser a mesma do cadastro.",
    ),
    Coluna("tipo", 12, "Sim", "entrada (compra que chegou) ou saída (retirada do galpão)."),
    Coluna(
        "quantidade",
        14,
        "Sim",
        "Maior que zero, na unidade do insumo. Decimais com vírgula: 25,5.",
    ),
    Coluna(
        "validade",
        14,
        "Não",
        "Validade do lote, só em entradas. Deixe em branco se o produto não vence.",
    ),
)


def gerar(db: Session) -> bytes:
    insumos = list(db.scalars(select(Insumo).order_by(Insumo.nome)))
    referencia = hoje()
    livro = Workbook()
    livro.properties.creator = "CampoCerto"
    livro.properties.title = "Movimentações de insumos para importar"

    # A primeira aba é a que a importação lê: tem de ser a de preencher.
    movimentacoes = livro.active
    movimentacoes.title = "Movimentações"
    movimentacoes.sheet_properties.tabColor = VERDE
    _aba_de_movimentacoes(movimentacoes, total_de_insumos=len(insumos))
    _aba_de_instrucoes(livro.create_sheet("Como preencher"))
    exemplo = livro.create_sheet("Exemplo")
    _aba_de_movimentacoes(exemplo, total_de_insumos=0, linhas=_exemplos(referencia))
    _aba_de_insumos(livro.create_sheet("Insumos cadastrados"), insumos)

    arquivo = BytesIO()
    livro.save(arquivo)
    return arquivo.getvalue()


def _aba_de_movimentacoes(
    aba: Worksheet, *, total_de_insumos: int, linhas: list[tuple] | None = None
) -> None:
    _titulos(aba, [coluna.titulo for coluna in COLUNAS])
    for indice, coluna in enumerate(COLUNAS, start=1):
        aba.column_dimensions[get_column_letter(indice)].width = coluna.largura
        # O comentário aparece ao passar o mouse no título: a explicação fica à mão.
        aba.cell(1, indice).comment = Comment(coluna.explicacao, "CampoCerto", 90, 260)
    aba.freeze_panes = "A2"
    _para_imprimir(aba)

    for linha in linhas or []:
        aba.append(linha)
    ultima = 1 + (LINHAS_PRONTAS if linhas is None else len(linhas))
    for data, *_, validade in aba.iter_rows(min_row=2, max_row=ultima, max_col=len(COLUNAS)):
        data.number_format = validade.number_format = FORMATO_DATA
        data.alignment = validade.alignment = Alignment(horizontal="center")
    if linhas is not None:
        return  # o exemplo é para ler, não para preencher

    def validar(coluna: str, **regra: object) -> None:
        letra = get_column_letter([c.titulo for c in COLUNAS].index(coluna) + 1)
        validacao = DataValidation(allow_blank=True, **regra)
        validacao.add(f"{letra}2:{letra}{ultima}")
        aba.add_data_validation(validacao)

    validar(
        "data",
        type="date",
        operator="between",
        formula1="DATE(2000,1,1)",
        formula2="TODAY()",
        errorTitle="Data inválida",
        error="Use uma data como 05/03/2026, de 2000 até hoje.",
    )
    if total_de_insumos:
        validar(
            "insumo",
            type="list",
            formula1=f"'Insumos cadastrados'!$A$2:$A${total_de_insumos + 1}",
            # Insumo novo também vale: a lista sugere, não obriga.
            showErrorMessage=False,
            promptTitle="Insumo",
            prompt="Escolha da lista ou digite um insumo novo.",
            showInputMessage=True,
        )
    validar(
        "unidade",
        type="list",
        formula1=f'"{",".join(UNIDADES)}"',
        showErrorMessage=False,
        promptTitle="Unidade",
        prompt="Só para insumo novo. Escolha ou digite outra.",
        showInputMessage=True,
    )
    validar(
        "tipo",
        type="list",
        formula1='"entrada,saída"',
        errorTitle="Tipo inválido",
        error="Escolha entrada (compra que chegou) ou saída (retirada do galpão).",
    )
    validar(
        "quantidade",
        type="decimal",
        operator="greaterThan",
        formula1="0",
        errorTitle="Quantidade inválida",
        error="A quantidade deve ser um número maior que zero.",
    )
    validar(
        "validade",
        type="date",
        operator="greaterThan",
        formula1="DATE(2000,1,1)",
        errorTitle="Validade inválida",
        error="Use uma data como 30/06/2027.",
    )


def _exemplos(referencia: date) -> list[tuple]:
    """Datas recentes, para o exemplo nunca parecer velho nem cair no futuro."""

    def dia(dias_atras: int) -> date:
        return referencia - timedelta(days=dias_atras)

    return [
        (dia(20), "Adubo NPK 20-05-20", "kg", "entrada", 1000, None),
        (dia(13), "Adubo NPK 20-05-20", None, "saída", 25.5, None),
        (dia(6), "Vacina contra aftosa", "dose", "entrada", 50, referencia + timedelta(days=250)),
    ]


def _aba_de_instrucoes(aba: Worksheet) -> None:
    aba.sheet_view.showGridLines = False
    larguras = {"A": 18, "B": 22, "C": 84}
    for letra, largura in larguras.items():
        aba.column_dimensions[letra].width = largura

    aba["A1"] = "Como preencher a planilha de movimentações"
    aba["A1"].font = _fonte(bold=True, size=16, color=VERDE_ESCURO)
    aba["A2"] = (
        "Uma linha por movimentação: cada compra que chegou (entrada) ou retirada do galpão "
        "(saída). Preencha a aba Movimentações; a aba Exemplo mostra como fica."
    )
    aba["A2"].font = _fonte(color=CINZA)

    _titulos(aba, ["Coluna", "Obrigatória", "Como preencher"], linha=4)
    for numero, coluna in enumerate(COLUNAS, start=5):
        aba.cell(numero, 1, coluna.titulo).font = _fonte(bold=True, color=VERDE)
        aba.cell(numero, 2, coluna.obrigatoria)
        aba.cell(numero, 3, coluna.explicacao)
        _texto_corrido(aba, numero, coluna.explicacao, larguras["C"])
        for indice in range(1, 4):
            aba.cell(numero, indice).border = Border(bottom=Side(color=CINZA_CLARO))

    proxima = 5 + len(COLUNAS) + 1
    for titulo, itens in (
        (
            "Antes de importar",
            [
                "• Comece na linha 2 da aba Movimentações e não mude os títulos da linha 1.",
                "• Lance também o estoque que já estava no galpão, como uma entrada: sem ela, "
                "as saídas deixam o estoque negativo.",
                "• Cada importação soma ao que já está no sistema. Importe cada período uma vez "
                "só — o CampoCerto avisa se as linhas já estiverem lá.",
                "• O Excel avisa na hora se a data, o tipo ou a quantidade estiverem fora do "
                "esperado.",
            ],
        ),
        (
            "Para importar",
            [
                "1. Salve este arquivo como está (.xlsx; não precisa converter para CSV).",
                "2. No CampoCerto, abra Importar dados e escolha o arquivo.",
                "3. Confira o resumo: quanto entra e sai de cada insumo. Nada é gravado antes "
                "de você confirmar.",
                "4. Se houver erro, a tela diz a linha e o que corrigir, e nada é importado até "
                "a planilha estar certa.",
            ],
        ),
    ):
        aba.cell(proxima, 1, titulo).font = _fonte(bold=True, size=12, color=VERDE_ESCURO)
        for item in itens:
            proxima += 1
            aba.cell(proxima, 1, item)
            aba.merge_cells(start_row=proxima, start_column=1, end_row=proxima, end_column=3)
            _texto_corrido(aba, proxima, item, sum(larguras.values()))
        proxima += 2
    _para_imprimir(aba)


def _aba_de_insumos(aba: Worksheet, insumos: list[Insumo]) -> None:
    _titulos(aba, ["insumo", "unidade"])
    aba.column_dimensions["A"].width = 34
    aba.column_dimensions["B"].width = 11
    aba.freeze_panes = "A2"
    for insumo in insumos:
        aba.append([insumo.nome, insumo.unidade])
    if not insumos:
        aba["A2"] = "Nenhum insumo cadastrado ainda: os da planilha serão cadastrados."
        aba["A2"].font = _fonte(italic=True, color=CINZA)
    # É referência (e a fonte da lista de escolha): protegida contra edição sem querer.
    aba.protection.sheet = True


def _fonte(**estilo: object) -> Font:
    """Calibri explícita: fonte sem nome cada programa de planilha resolve de um jeito."""
    return Font(name="Calibri", **{"size": 11, **estilo})


def _titulos(aba: Worksheet, titulos: list[str], linha: int = 1) -> None:
    for indice, titulo in enumerate(titulos, start=1):
        celula = aba.cell(linha, indice, titulo)
        celula.font = _fonte(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor=VERDE)
        celula.alignment = Alignment(horizontal="center", vertical="center")
    aba.row_dimensions[linha].height = 22


def _texto_corrido(aba: Worksheet, linha: int, texto: str, largura: int) -> None:
    """Quebra o texto na largura da coluna e ajusta a altura da linha: nem todo programa
    de planilha recalcula a altura sozinho."""
    for celula in aba[linha]:
        celula.alignment = Alignment(wrap_text=True, vertical="top")
    aba.row_dimensions[linha].height = 15 * math.ceil(len(texto) / (largura * 1.1))


def _para_imprimir(aba: Worksheet) -> None:
    aba.page_setup.orientation = "landscape"
    aba.page_setup.fitToWidth = 1
    aba.page_setup.fitToHeight = 0
    aba.sheet_properties.pageSetUpPr.fitToPage = True
    aba.print_title_rows = "1:1"
