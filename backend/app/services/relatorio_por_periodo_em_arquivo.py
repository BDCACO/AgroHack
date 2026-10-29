"""O relatório por período em arquivo, para baixar da tela de Relatórios: em Excel, para
recortar e somar de novo; em PDF, para imprimir ou mandar ao agrônomo e ao contador.

Os dois saem do mesmo `RelatorioPorPeriodo` que a tela mostra (as somas são do domínio,
`app/domain/relatorio_por_periodo.py`): o que se baixa é o que se viu, com os mesmos
filtros. Aqui só se arruma na página.
"""

from collections.abc import Sequence
from datetime import date
from io import BytesIO

from fpdf import FPDF
from fpdf.enums import XPos, YPos
from fpdf.fonts import FontFace
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from app.domain.relatorio_por_periodo import (
    ConsumoDoInsumo,
    RelatorioPorPeriodo,
    Totais,
)
from app.models import Categoria, Fornecedor, Insumo, Talhao
from app.services.planilha_modelo import (
    CINZA,
    FORMATO_DATA,
    FORMATO_DINHEIRO,
    VERDE_ESCURO,
    fonte,
    para_imprimir,
    titulos,
)

AGRUPAMENTOS = {"mes": "por mês", "semana": "por semana", "dia": "por dia"}
# As cores do gráfico da tela (GraficoPorPeriodo.vue): gasto, consumo e perdas.
CORES = {"gasto": (63, 122, 44), "consumo": (168, 127, 72), "perdas": (220, 38, 38)}
# Por dia, um ano tem 365 barras de três: finas demais para ler. Acima disso, só a tabela.
MAXIMO_NO_GRAFICO = 62


def filtros_em_texto(
    db: Session,
    *,
    categoria_id: int | None = None,
    insumo_id: int | None = None,
    talhao_id: int | None = None,
    fornecedor_id: int | None = None,
) -> list[str]:
    """O recorte em palavras, para o cabeçalho: "Categoria: Fertilizante". Um id que não
    existe mais aparece como o número, e não some: o relatório saiu filtrado por ele."""
    filtros = []
    for rotulo, modelo, chave in (
        ("Categoria", Categoria, categoria_id),
        ("Insumo", Insumo, insumo_id),
        ("Talhão", Talhao, talhao_id),
        ("Fornecedor", Fornecedor, fornecedor_id),
    ):
        if chave is not None:
            registro = db.get(modelo, chave)
            filtros.append(f"{rotulo}: {registro.nome if registro else f'nº {chave}'}")
    return filtros


def nome_do_arquivo(relatorio: RelatorioPorPeriodo, extensao: str) -> str:
    return f"relatorio-campocerto-{relatorio.desde}-a-{relatorio.ate}.{extensao}"


def consumo_em_reais(totais: Totais) -> float:
    """Em reais, e não em quantidade: kg de um insumo e litros de outro não se somam."""
    return round(sum(c.consumo for c in totais.por_categoria), 2)


def consumo_por_insumo(totais: Totais) -> list[ConsumoDoInsumo]:
    """O que mais custou primeiro, como na tela: é por onde se começa a cortar."""
    usados = [c for c in totais.consumo if c.quantidade > 0]
    return sorted(usados, key=lambda c: -1 if c.valor is None else c.valor, reverse=True)


# ---------------------------------------------------------------------------- Excel


def planilha(relatorio: RelatorioPorPeriodo, filtros: Sequence[str], gerado_em: date) -> bytes:
    livro = Workbook()
    _aba_de_resumo(livro.active, relatorio, filtros, gerado_em)
    _aba_por_periodo(livro.create_sheet("Por período"), relatorio)
    _aba_de_consumo(livro.create_sheet("Consumo por insumo"), relatorio.total)
    _aba_por_categoria(livro.create_sheet("Por categoria"), relatorio.total)
    _aba_por_fornecedor(livro.create_sheet("Por fornecedor"), relatorio.total)
    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()


def _aba_de_resumo(
    aba: Worksheet, relatorio: RelatorioPorPeriodo, filtros: Sequence[str], gerado_em: date
) -> None:
    aba.title = "Resumo"
    aba.column_dimensions["A"].width = 30
    aba.column_dimensions["B"].width = 22
    aba["A1"] = "CampoCerto · Relatório por período"
    aba["A1"].font = fonte(bold=True, size=16, color=VERDE_ESCURO)
    aba["A2"] = (
        f"{_data(relatorio.desde)} a {_data(relatorio.ate)}, "
        f"{AGRUPAMENTOS[relatorio.agrupar]} · gerado em {_data(gerado_em)}"
    )
    aba["A2"].font = fonte(color=CINZA)
    linha = 3
    for filtro in filtros:
        aba.cell(linha, 1, _limpo(filtro)).font = fonte(color=CINZA)
        linha += 1

    linha += 1
    titulos(aba, ["No período inteiro", "R$"], linha)
    total = relatorio.total
    for rotulo, valor in (
        ("Gasto em compras", total.gasto),
        ("Consumo", consumo_em_reais(total)),
        ("Perdas (descarte e falta)", total.perdas),
    ):
        linha += 1
        aba.cell(linha, 1, rotulo)
        aba.cell(linha, 2, valor).number_format = FORMATO_DINHEIRO
    linha += 1
    aba.cell(linha, 1, "Compras sem preço (fora do gasto)")
    aba.cell(linha, 2, total.compras_sem_preco)

    linha += 2
    for limitacao in relatorio.limitacoes:
        celula = aba.cell(linha, 1, f"• {limitacao}")
        celula.font = fonte(italic=True, color=CINZA)
        linha += 1
    para_imprimir(aba)
    aba.print_title_rows = None


def _aba_por_periodo(aba: Worksheet, relatorio: RelatorioPorPeriodo) -> None:
    _tabela(
        aba,
        [
            ("período", 22),
            ("de", 12),
            ("até", 12),
            ("gasto em compras (R$)", 22),
            ("consumo (R$)", 16),
            ("perdas (R$)", 14),
            ("compras sem preço", 18),
        ],
    )
    for numero, periodo in enumerate(relatorio.periodos, start=2):
        valores = (
            periodo.rotulo,
            periodo.inicio,
            periodo.fim,
            periodo.gasto,
            consumo_em_reais(periodo),
            periodo.perdas,
            periodo.compras_sem_preco,
        )
        for coluna, valor in enumerate(valores, start=1):
            celula = aba.cell(numero, coluna, valor)
            if coluna in (2, 3):
                celula.number_format = FORMATO_DATA
            elif coluna in (4, 5, 6):
                celula.number_format = FORMATO_DINHEIRO
    if relatorio.periodos and len(relatorio.periodos) <= MAXIMO_NO_GRAFICO:
        ultima = len(relatorio.periodos) + 1
        grafico = BarChart()
        grafico.title = "Gasto, consumo e perdas"
        grafico.y_axis.title = "R$"
        grafico.height, grafico.width = 9, 24
        grafico.add_data(
            Reference(aba, min_col=4, max_col=6, min_row=1, max_row=ultima), titles_from_data=True
        )
        grafico.set_categories(Reference(aba, min_col=1, min_row=2, max_row=ultima))
        for serie, (r, g, b) in zip(grafico.series, CORES.values(), strict=True):
            serie.graphicalProperties.solidFill = f"{r:02X}{g:02X}{b:02X}"
            serie.graphicalProperties.line.noFill = True
        aba.add_chart(grafico, "I2")


def _aba_de_consumo(aba: Worksheet, total: Totais) -> None:
    _tabela(aba, [("insumo", 34), ("quantidade", 14), ("unidade", 10), ("valor (R$)", 16)])
    linhas = consumo_por_insumo(total)
    for numero, item in enumerate(linhas, start=2):
        aba.cell(numero, 1, _limpo(item.insumo)).data_type = "s"
        aba.cell(numero, 2, item.quantidade)
        aba.cell(numero, 3, _limpo(item.unidade)).data_type = "s"
        # Sem preço, a célula vazia: R$ 0,00 diria que saiu de graça.
        aba.cell(numero, 4, item.valor).number_format = FORMATO_DINHEIRO
    if not linhas:
        _vazio(aba, "Nenhum consumo no período.")


def _aba_por_categoria(aba: Worksheet, total: Totais) -> None:
    _tabela(aba, [("categoria", 30), ("gasto (R$)", 16), ("consumo (R$)", 16), ("perdas (R$)", 16)])
    linhas = [c for c in total.por_categoria if c.gasto or c.consumo or c.perdas]
    for numero, item in enumerate(linhas, start=2):
        aba.cell(numero, 1, _limpo(item.categoria or "Sem categoria")).data_type = "s"
        for coluna, valor in enumerate((item.gasto, item.consumo, item.perdas), start=2):
            aba.cell(numero, coluna, valor).number_format = FORMATO_DINHEIRO
    if not linhas:
        _vazio(aba, "Nenhum movimento com valor no período.")


def _aba_por_fornecedor(aba: Worksheet, total: Totais) -> None:
    _tabela(aba, [("fornecedor", 30), ("gasto (R$)", 16), ("compras", 10), ("sem preço", 12)])
    linhas = [f for f in total.por_fornecedor if f.compras]
    for numero, item in enumerate(linhas, start=2):
        aba.cell(numero, 1, _limpo(item.fornecedor or "Sem fornecedor")).data_type = "s"
        aba.cell(numero, 2, item.gasto).number_format = FORMATO_DINHEIRO
        aba.cell(numero, 3, item.compras)
        aba.cell(numero, 4, item.compras_sem_preco)
    if not linhas:
        _vazio(aba, "Nenhuma compra no período.")


def _tabela(aba: Worksheet, colunas: list[tuple[str, int]]) -> None:
    titulos(aba, [titulo for titulo, _ in colunas])
    for indice, (_, largura) in enumerate(colunas, start=1):
        aba.column_dimensions[get_column_letter(indice)].width = largura
    aba.freeze_panes = "A2"
    para_imprimir(aba)


def _vazio(aba: Worksheet, texto: str) -> None:
    aba["A2"] = texto
    aba["A2"].font = fonte(italic=True, color=CINZA)
    aba["A2"].alignment = Alignment(vertical="top")


def _limpo(texto: str) -> str:
    """Nome digitado pelo usuário, sem o caractere de controle que o Excel recusa."""
    return ILLEGAL_CHARACTERS_RE.sub("", texto)


# ---------------------------------------------------------------------------- PDF

_VERDE = (63, 122, 44)  # campo-600
_VERDE_ESCURO = (31, 59, 23)  # campo-900
_TINTA = (41, 37, 36)  # stone-800
_CINZA = (87, 83, 78)  # stone-600
_CINZA_CLARO = (231, 229, 228)  # stone-200
_FUNDO = (245, 245, 244)  # stone-100


class _Pdf(FPDF):
    def __init__(self, rodape: str) -> None:
        super().__init__(format="A4")
        self._rodape = rodape
        # As fontes de base do PDF (Helvetica) vêm em WinAnsi, o cp1252: acentos, "—", "…"
        # e "•" cabem, sem embutir arquivo de fonte.
        self.core_fonts_encoding = "windows-1252"
        self.set_margins(15, 15, 15)
        self.set_auto_page_break(auto=True, margin=18)
        self.alias_nb_pages()

    def footer(self) -> None:
        self.set_y(-12)
        self.set_font("Helvetica", size=8)
        self.set_text_color(*_CINZA)
        self.cell(0, 5, self._rodape)
        self.cell(0, 5, f"Página {self.page_no()} de {{nb}}", align="R")


def pdf(relatorio: RelatorioPorPeriodo, filtros: Sequence[str], gerado_em: date) -> bytes:
    documento = _Pdf(f"CampoCerto · gerado em {_data(gerado_em)}")
    documento.set_title("Relatório por período")
    documento.set_author("CampoCerto")
    documento.add_page()

    documento.set_font("Helvetica", "B", 18)
    documento.set_text_color(*_VERDE_ESCURO)
    documento.cell(0, 9, "Relatório por período", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    documento.set_font("Helvetica", size=10)
    documento.set_text_color(*_CINZA)
    documento.cell(
        0,
        6,
        _pdf_texto(
            f"{_data(relatorio.desde)} a {_data(relatorio.ate)}, {AGRUPAMENTOS[relatorio.agrupar]}"
            + (f" · {' · '.join(filtros)}" if filtros else "")
        ),
        new_x=XPos.LMARGIN,
        new_y=YPos.NEXT,
    )
    documento.ln(4)

    total = relatorio.total
    _cartoes(
        documento,
        [
            ("Gasto em compras", _reais(total.gasto), _sem_preco(total.compras_sem_preco)),
            ("Consumo", _reais(consumo_em_reais(total)), "pelo custo médio"),
            ("Perdas", _reais(total.perdas), "descarte e falta da contagem"),
        ],
    )

    if relatorio.periodos and len(relatorio.periodos) <= MAXIMO_NO_GRAFICO:
        _titulo(documento, "Gasto, consumo e perdas")
        _grafico(documento, relatorio)

    _titulo(documento, "Por período")
    _tabela_pdf(
        documento,
        ["Período", "Gasto", "Consumo", "Perdas"],
        [
            (p.rotulo, _reais(p.gasto), _reais(consumo_em_reais(p)), _reais(p.perdas))
            for p in relatorio.periodos
        ],
        (2, 1, 1, 1),
    )

    consumo = consumo_por_insumo(total)
    _titulo(documento, "Consumo por insumo")
    _tabela_pdf(
        documento,
        ["Insumo", "Quantidade", "Valor"],
        [
            (
                c.insumo,
                f"{_numero(c.quantidade)} {c.unidade}",
                "sem preço" if c.valor is None else _reais(c.valor),
            )
            for c in consumo
        ],
        (3, 1.3, 1.2),
        vazio="Nenhum consumo no período.",
    )

    categorias = [c for c in total.por_categoria if c.gasto or c.consumo or c.perdas]
    _titulo(documento, "Por categoria")
    _tabela_pdf(
        documento,
        ["Categoria", "Gasto", "Consumo", "Perdas"],
        [
            (c.categoria or "Sem categoria", _reais(c.gasto), _reais(c.consumo), _reais(c.perdas))
            for c in categorias
        ],
        (2, 1, 1, 1),
        vazio="Nenhum movimento com valor no período.",
    )

    fornecedores = [f for f in total.por_fornecedor if f.compras]
    _titulo(documento, "Por fornecedor")
    _tabela_pdf(
        documento,
        ["Fornecedor", "Gasto", "Compras"],
        [
            (
                f.fornecedor or "Sem fornecedor",
                _reais(f.gasto),
                str(f.compras)
                + (f" ({f.compras_sem_preco} sem preço)" if f.compras_sem_preco else ""),
            )
            for f in fornecedores
        ],
        (2.4, 1.2, 1.4),
        vazio="Nenhuma compra no período.",
    )

    if relatorio.limitacoes:
        documento.ln(3)
        documento.set_font("Helvetica", "I", 8.5)
        documento.set_text_color(*_CINZA)
        for limitacao in relatorio.limitacoes:
            documento.multi_cell(
                0, 4.5, _pdf_texto(f"• {limitacao}"), new_x=XPos.LMARGIN, new_y=YPos.NEXT
            )
    return bytes(documento.output())


def _titulo(documento: FPDF, texto: str) -> None:
    # Título sozinho no pé da página, longe da tabela dele, não: passa para a próxima.
    if documento.will_page_break(30):
        documento.add_page()
    documento.ln(5)
    documento.set_font("Helvetica", "B", 12)
    documento.set_text_color(*_VERDE_ESCURO)
    documento.cell(0, 7, _pdf_texto(texto), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    documento.ln(1)


def _cartoes(documento: FPDF, cartoes: list[tuple[str, str, str]]) -> None:
    largura = (documento.epw - 4 * (len(cartoes) - 1)) / len(cartoes)
    topo = documento.get_y()
    for indice, (rotulo, valor, nota) in enumerate(cartoes):
        x = documento.l_margin + indice * (largura + 4)
        documento.set_fill_color(*_FUNDO)
        documento.set_draw_color(*_CINZA_CLARO)
        documento.rect(x, topo, largura, 22, style="DF", round_corners=True, corner_radius=2)
        documento.set_xy(x + 4, topo + 3)
        documento.set_font("Helvetica", size=8.5)
        documento.set_text_color(*_CINZA)
        documento.cell(largura - 8, 4, _pdf_texto(rotulo))
        documento.set_xy(x + 4, topo + 8)
        documento.set_font("Helvetica", "B", 13)
        documento.set_text_color(*_TINTA)
        documento.cell(largura - 8, 7, _pdf_texto(valor))
        documento.set_xy(x + 4, topo + 15.5)
        documento.set_font("Helvetica", size=7.5)
        documento.set_text_color(*_CINZA)
        documento.cell(largura - 8, 4, _pdf_texto(nota))
    documento.set_xy(documento.l_margin, topo + 24)


def _grafico(documento: FPDF, relatorio: RelatorioPorPeriodo) -> None:
    """Barras de gasto, consumo e perdas por período, as cores da tela, com a legenda."""
    altura, legenda, eixo = 52, 7, 12
    if documento.will_page_break(altura + legenda + eixo):
        documento.add_page()
    x0, topo = documento.l_margin + 16, documento.get_y() + legenda
    largura = documento.epw - 16
    series = [(p.gasto, consumo_em_reais(p), p.perdas) for p in relatorio.periodos]
    maximo = max((v for trio in series for v in trio), default=0) or 1

    # Legenda
    documento.set_font("Helvetica", size=8)
    x = documento.l_margin
    for nome, cor in zip(("Gasto", "Consumo", "Perdas"), CORES.values(), strict=True):
        documento.set_fill_color(*cor)
        documento.rect(x, documento.get_y() + 1, 3, 3, style="F")
        documento.set_xy(x + 4, documento.get_y())
        documento.set_text_color(*_CINZA)
        documento.cell(20, 5, nome)
        x += 24

    # Linhas de grade com o valor, em quatro faixas
    documento.set_draw_color(*_CINZA_CLARO)
    documento.set_font("Helvetica", size=6.5)
    for faixa in range(5):
        y = topo + altura - altura * faixa / 4
        documento.line(x0, y, x0 + largura, y)
        documento.set_xy(documento.l_margin - 2, y - 2)
        documento.cell(17, 4, _reais_curto(maximo * faixa / 4), align="R")

    passo = largura / len(series)
    barra = min(passo * 0.8 / 3, 6)
    for indice, (trio, periodo) in enumerate(zip(series, relatorio.periodos, strict=True)):
        inicio = x0 + indice * passo + (passo - 3 * barra) / 2
        for posicao, (valor, cor) in enumerate(zip(trio, CORES.values(), strict=True)):
            if valor > 0:
                h = altura * valor / maximo
                documento.set_fill_color(*cor)
                documento.rect(inicio + posicao * barra, topo + altura - h, barra, h, style="F")
        # O rótulo de cada período, se couber; senão, um sim e outros não.
        a_cada = max(1, round(18 / passo))
        if indice % a_cada == 0:
            documento.set_xy(x0 + indice * passo, topo + altura + 1)
            documento.set_text_color(*_CINZA)
            documento.cell(
                passo * a_cada,
                4,
                _pdf_texto(periodo.rotulo.removeprefix("semana de ")),
                align="C" if a_cada == 1 else "L",
            )
    documento.set_xy(documento.l_margin, topo + altura + eixo - 4)


def _tabela_pdf(
    documento: FPDF,
    cabecalho: list[str],
    linhas: list[tuple[str, ...]],
    larguras: tuple[float, ...],
    vazio: str = "Nada no período.",
) -> None:
    if not linhas:
        documento.set_font("Helvetica", "I", 9)
        documento.set_text_color(*_CINZA)
        documento.cell(0, 6, _pdf_texto(vazio), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        return
    documento.set_font("Helvetica", size=9)
    documento.set_text_color(*_TINTA)
    documento.set_draw_color(*_CINZA_CLARO)
    alinhamentos = ("LEFT",) + ("RIGHT",) * (len(cabecalho) - 1)
    with documento.table(
        col_widths=larguras,
        text_align=alinhamentos,
        line_height=6,
        headings_style=FontFace(emphasis="BOLD", color=(255, 255, 255), fill_color=_VERDE),
        borders_layout="HORIZONTAL_LINES",
        first_row_as_headings=True,
        repeat_headings=1,
    ) as tabela:
        # A faixa de cada linha com a cor dela: o modo automático do fpdf pinta a linha sem
        # faixa com a última cor usada, a do gráfico.
        for indice, linha in enumerate([tuple(cabecalho), *linhas]):
            faixa = _FUNDO if indice % 2 else (255, 255, 255)
            fila = tabela.row(style=None if indice == 0 else FontFace(fill_color=faixa))
            for texto in linha:
                fila.cell(_pdf_texto(texto))


# ---------------------------------------------------------------------------- formatos


def _data(dia: date) -> str:
    return dia.strftime("%d/%m/%Y")


def _numero(valor: float, casas: int = 2) -> str:
    """1234.5 vira "1.234,5": o número como se escreve no Brasil, sem zeros à toa."""
    texto = f"{valor:,.{casas}f}"
    if casas:
        texto = texto.rstrip("0").rstrip(".")
    return texto.replace(",", "_").replace(".", ",").replace("_", ".")


def _reais(valor: float) -> str:
    texto = f"{valor:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")
    return f"R$ {texto}"


def _reais_curto(valor: float) -> str:
    """No eixo do gráfico: "R$ 12 mil", e não "R$ 12.000,00", que não cabe."""
    if valor >= 1000:
        return f"R$ {_numero(valor / 1000, 1)} mil"
    return f"R$ {_numero(valor, 0)}"


def _sem_preco(quantas: int) -> str:
    if not quantas:
        return "todas as compras com preço"
    return (
        "1 compra sem preço, fora da soma"
        if quantas == 1
        else f"{quantas} compras sem preço, fora da soma"
    )


def _pdf_texto(texto: str) -> str:
    """O que a fonte de base do PDF não tem (emoji, outras escritas, num nome digitado) vira
    "?", em vez de derrubar o relatório inteiro."""
    return texto.encode("cp1252", "replace").decode("cp1252")
