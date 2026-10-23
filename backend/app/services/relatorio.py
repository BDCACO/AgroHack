"""Relatório do estoque em Excel (.xlsx): o que o Painel mostra, para imprimir, mandar ao
agrônomo ou guardar.

A aba Movimentações tem o formato da importação: o relatório também é a cópia de
segurança de todas as entradas e saídas (com o talhão de cada aplicação e o fornecedor de
cada compra), e volta para o sistema pela mesma tela. Os talhões precisam estar
cadastrados antes: a aba Talhões tem a área e a cultura de cada um. A aba Pedidos mostra o
que foi pedido e ainda não chegou; ela não volta pela importação.
"""

import math
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, Cell
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.relogio import hoje
from app.domain.estoque import Gravidade, TipoAlerta
from app.models import Fornecedor, Insumo, Movimentacao, Talhao
from app.schemas.compras import PedidoResposta
from app.schemas.estoque import InsumoResposta, Resumo
from app.schemas.talhao import TalhaoResposta
from app.services import estoque, pedidos, resumo, talhoes
from app.services.planilha_modelo import (
    CINZA,
    FORMATO_DATA,
    FORMATO_DINHEIRO,
    VERDE_ESCURO,
    aba_de_movimentacoes,
    fonte,
    para_imprimir,
    texto_corrido,
    titulos,
)

# Os mesmos rótulos e cores da tela (SeloDeGravidade.vue, CartaoDeAlerta.vue).
GRAVIDADES: dict[Gravidade | None, tuple[str, str, str]] = {
    "alta": ("Urgente", "FEE2E2", "991B1B"),
    "media": ("Atenção", "FEF3C7", "92400E"),
    "baixa": ("Observar", "E0F2FE", "075985"),
    None: ("Sem alertas", "E0EFD9", "326223"),
}
ROTULOS: dict[TipoAlerta, str] = {
    "abaixo_do_minimo": "Abaixo do mínimo",
    "risco_de_falta": "Risco de falta",
    "vencido": "Lote vencido",
    "vencimento_proximo": "Vencimento próximo",
    "excesso": "Excesso de estoque",
    "baixa_movimentacao": "Baixa movimentação",
    "consumo_incomum": "Consumo incomum",
    "carencia": "Carência",
}
_ORDEM: dict[Gravidade, int] = {"alta": 0, "media": 1, "baixa": 2}
# Como a importação lê: o descarte volta como descarte, e não como consumo.
_TIPO_NA_PLANILHA = {
    ("entrada", None): "entrada",
    ("saida", None): "saída",
    ("saida", "descarte"): "descarte",
    ("entrada", "inventario"): "sobra",
    ("saida", "inventario"): "falta",
}


def gerar(db: Session) -> bytes:
    insumos = estoque.listar_insumos(db)
    referencia = hoje()
    livro = Workbook()
    livro.properties.creator = "CampoCerto"
    livro.properties.title = f"Situação do estoque em {referencia:%d/%m/%Y}"

    aba_resumo = livro.active
    aba_resumo.title = "Resumo"
    _aba_de_resumo(aba_resumo, insumos, resumo.resumir(db, insumos, referencia), referencia)
    _aba_de_alertas(livro.create_sheet("Alertas"), insumos)
    _aba_de_reposicao(livro.create_sheet("Reposição"), insumos)
    _aba_de_pedidos(
        livro.create_sheet("Pedidos"),
        [p for situacao in pedidos.EM_ANDAMENTO for p in pedidos.listar_pedidos(db, situacao)],
    )
    _aba_de_estoque(livro.create_sheet("Estoque"), insumos)
    _aba_de_talhoes(livro.create_sheet("Talhões"), talhoes.listar_talhoes(db))
    aba_de_movimentacoes(
        livro.create_sheet("Movimentações"), total_de_insumos=0, linhas=_movimentacoes(db)
    )

    arquivo = BytesIO()
    livro.save(arquivo)
    return arquivo.getvalue()


def _aba_de_resumo(
    aba: Worksheet, insumos: list[InsumoResposta], geral: Resumo, referencia: date
) -> None:
    aba.sheet_view.showGridLines = False
    aba.column_dimensions["A"].width = 26
    aba.column_dimensions["B"].width = 16
    aba.column_dimensions["C"].width = 70

    aba["A1"] = "Situação do estoque"
    aba["A1"].font = fonte(bold=True, size=16, color=VERDE_ESCURO)
    aba["A2"] = f"CampoCerto · gerado em {referencia:%d/%m/%Y}"
    aba["A2"].font = fonte(color=CINZA)
    linha = 4
    if any(insumo.demonstracao for insumo in insumos):
        aba["A3"] = "Dados de demonstração: parte dos insumos não vem de uma fazenda real."
        aba["A3"].font = fonte(bold=True, color="92400E")
        linha = 5

    titulos(aba, ["Alertas", "Quantos"], linha=linha)
    for gravidade, total in geral.alertas.model_dump().items():
        linha += 1
        _selo(aba.cell(linha, 1), gravidade)
        aba.cell(linha, 2, total).font = fonte(bold=True)
    pedir_hoje = geral.compras_sugeridas.pedir_hoje
    for rotulo, valor in (
        ("Pedir hoje", pedir_hoje),
        ("Insumos cadastrados", len(insumos)),
        # À parte dos alertas, como no GET /resumo: a carência é do talhão.
        ("Talhões em carência", geral.talhoes_em_carencia),
    ):
        linha += 1
        aba.cell(linha, 1, rotulo)
        aba.cell(linha, 2, valor).font = fonte(bold=True)

    linha += 2
    titulos(aba, ["Em reais", "R$"], linha=linha)
    for rotulo, valor in (
        ("Valor em estoque", geral.valor_em_estoque),
        ("Perdas nos últimos 30 dias", geral.perdas.ultimos_30_dias),
        ("Perdas no total", geral.perdas.total),
        ("Compras dos próximos 30 dias", geral.compras_sugeridas.valor),
    ):
        linha += 1
        aba.cell(linha, 1, rotulo)
        # Sem preço nenhum, a célula diz isso, em vez de um zero que pareceria estoque de graça.
        celula = aba.cell(linha, 2, "sem preço" if valor is None else valor)
        celula.font = fonte(bold=True)
        celula.number_format = FORMATO_DINHEIRO
    for limitacao in geral.limitacoes:
        linha += 1
        aba.cell(linha, 1, f"• {limitacao}").font = fonte(italic=True, color=CINZA)

    linha += 2
    aba.cell(linha, 1, "O que tem em cada aba").font = fonte(bold=True, size=12, color=VERDE_ESCURO)
    for nome, explicacao in (
        ("Alertas", "Cada alerta, dos mais graves aos mais leves, com a conta que o justifica."),
        ("Reposição", "Quanto pedir de cada insumo, até quando e quanto custa, com a conta."),
        (
            "Pedidos",
            "O que foi pedido às revendas e ainda não chegou, com a previsão de chegada; os "
            "atrasados em destaque.",
        ),
        (
            "Estoque",
            "A situação de cada insumo, o custo médio e o valor em estoque, e o que a análise "
            "não sabe (limitações).",
        ),
        (
            "Talhões",
            "Cada talhão com a área, a cultura, até quando não colher (a carência do que foi "
            "aplicado, com a conta) e o custo do que foi aplicado nele, total e por hectare.",
        ),
        (
            "Movimentações",
            "Todas as entradas e saídas, no formato da importação: este arquivo também é a "
            "cópia de segurança, e volta pela tela Importar dados, com os preços, o talhão de "
            "cada aplicação e o fornecedor de cada compra (o que faltar é cadastrado). Os "
            "talhões precisam estar cadastrados antes (a aba Talhões tem a "
            "área e a cultura). O mínimo, o prazo, a embalagem, o pedido mínimo e a carência "
            "de cada insumo não voltam: lance-os de novo na tela de Insumos.",
        ),
    ):
        linha += 1
        aba.cell(linha, 1, nome).font = fonte(bold=True)
        aba.cell(linha, 3, explicacao)
        texto_corrido(aba, linha, explicacao, 70)
    para_imprimir(aba)


def _aba_de_alertas(aba: Worksheet, insumos: list[InsumoResposta]) -> None:
    _tabela(aba, [("gravidade", 12), ("insumo", 28), ("alerta", 22), ("justificativa", 90)])
    alertas = sorted(
        ((insumo, alerta) for insumo in insumos for alerta in insumo.analise.alertas),
        key=lambda par: _ORDEM[par[1].gravidade],
    )
    if not alertas:
        _vazio(aba, "Nenhum alerta hoje.")
    for numero, (insumo, alerta) in enumerate(alertas, start=2):
        _texto(aba.cell(numero, 2), insumo.nome)
        aba.cell(numero, 3, ROTULOS[alerta.tipo])
        aba.cell(numero, 4, alerta.justificativa)
        texto_corrido(aba, numero, alerta.justificativa, 90)
        _selo(aba.cell(numero, 1), alerta.gravidade)
    _filtro(aba, len(alertas))


def _aba_de_reposicao(aba: Worksheet, insumos: list[InsumoResposta]) -> None:
    _tabela(
        aba,
        [
            ("insumo", 28),
            ("quantidade", 13),
            ("unidade", 10),
            ("pedir até", 13),
            ("custo estimado (R$)", 19),
            ("conta", 80),
        ],
    )
    pedidos = sorted(
        (i for i in insumos if i.analise.reposicao), key=lambda i: i.analise.reposicao.pedir_em
    )
    if not pedidos:
        _vazio(aba, "Nada a pedir: nenhum insumo chega ao ponto de reposição.")
    for numero, insumo in enumerate(pedidos, start=2):
        reposicao = insumo.analise.reposicao
        _texto(aba.cell(numero, 1), insumo.nome)
        aba.cell(numero, 2, reposicao.quantidade)
        _texto(aba.cell(numero, 3), insumo.unidade)
        pedir = aba.cell(numero, 4, reposicao.pedir_em)
        pedir.number_format = FORMATO_DATA
        aba.cell(numero, 5, reposicao.custo_estimado).number_format = FORMATO_DINHEIRO
        aba.cell(numero, 6, reposicao.justificativa)
        texto_corrido(aba, numero, reposicao.justificativa, 80)
        if reposicao.pedir_hoje:
            # O que é para pedir hoje salta aos olhos, como na tela.
            pedir.fill = PatternFill("solid", fgColor=GRAVIDADES["alta"][1])
            pedir.font = fonte(bold=True, color=GRAVIDADES["alta"][2])
        pedir.alignment = Alignment(horizontal="center", vertical="top")
    _filtro(aba, len(pedidos))


def _aba_de_pedidos(aba: Worksheet, lista: list[PedidoResposta]) -> None:
    _tabela(
        aba,
        [
            ("pedido", 9),
            ("feito em", 12),
            ("fornecedor", 26),
            ("previsto para", 13),
            ("insumo", 28),
            ("unidade", 10),
            ("pedido (qtd)", 13),
            ("recebido", 11),
            ("falta chegar", 13),
            ("preço combinado (R$)", 19),
            ("situação", 12),
        ],
    )
    # Uma linha por insumo que ainda falta chegar; a previsão mais próxima primeiro.
    linhas = [
        (pedido, item)
        for pedido in sorted(
            lista, key=lambda p: (p.previsto_para is None, p.previsto_para or p.feito_em, p.id)
        )
        for item in sorted(pedido.itens, key=lambda i: i.insumo)
        if item.falta_receber > 0
    ]
    if not linhas:
        _vazio(aba, "Nenhum pedido a caminho.")
    for numero, (pedido, item) in enumerate(linhas, start=2):
        aba.cell(numero, 1, pedido.id)
        aba.cell(numero, 2, pedido.feito_em).number_format = FORMATO_DATA
        if pedido.fornecedor is not None:
            _texto(aba.cell(numero, 3), pedido.fornecedor)
        aba.cell(numero, 4, pedido.previsto_para).number_format = FORMATO_DATA
        _texto(aba.cell(numero, 5), item.insumo)
        _texto(aba.cell(numero, 6), item.unidade)
        for coluna, valor in enumerate(
            (item.quantidade, item.quantidade_recebida, item.falta_receber), start=7
        ):
            aba.cell(numero, coluna, valor)
        aba.cell(numero, 10, item.valor_unitario).number_format = FORMATO_DINHEIRO
        situacao = aba.cell(numero, 11, _SITUACAO_DO_PEDIDO[pedido.situacao])
        if pedido.atrasado:
            # O atrasado salta aos olhos, nas cores do urgente: é hora de cobrar a revenda.
            situacao.value = "Atrasado"
            situacao.fill = PatternFill("solid", fgColor=GRAVIDADES["alta"][1])
            situacao.font = fonte(bold=True, color=GRAVIDADES["alta"][2])
    _filtro(aba, len(linhas))


_SITUACAO_DO_PEDIDO = {"aberto": "Aberto", "parcial": "Parcial"}


def _aba_de_estoque(aba: Worksheet, insumos: list[InsumoResposta]) -> None:
    _tabela(
        aba,
        [
            ("insumo", 28),
            ("unidade", 10),
            ("estoque atual", 14),
            ("mínimo", 11),
            ("consumo por dia", 16),
            ("dura (dias)", 12),
            ("entrega (dias)", 14),
            ("custo médio (R$)", 16),
            ("valor em estoque (R$)", 20),
            ("situação", 13),
            ("limitações da análise", 80),
        ],
    )
    for numero, insumo in enumerate(insumos, start=2):
        analise = insumo.analise
        consumo, cobertura = analise.consumo_diario, analise.cobertura_dias
        _texto(aba.cell(numero, 1), insumo.nome)
        _texto(aba.cell(numero, 2), insumo.unidade)
        valores = (
            insumo.estoque_atual,
            insumo.estoque_minimo,
            None if consumo is None else round(consumo, 2),
            None if cobertura is None else round(cobertura),
            insumo.prazo_reposicao_dias,
            insumo.custo_medio,
            insumo.valor_em_estoque,
        )
        for coluna, valor in enumerate(valores, start=3):
            aba.cell(numero, coluna, valor)
        for coluna in (8, 9):
            aba.cell(numero, coluna).number_format = FORMATO_DINHEIRO
        limitacoes = [f"• {texto}" for texto in analise.limitacoes]
        aba.cell(numero, 11, "\n".join(limitacoes))
        texto_corrido(aba, numero, "", 80)
        # Uma limitação por linha do texto, cada uma quebrada na largura da coluna.
        aba.row_dimensions[numero].height = 15 * max(
            1, sum(math.ceil(len(texto) / 88) for texto in limitacoes)
        )
        # Depois do texto_corrido, que alinha a linha inteira. A API já ordena os
        # alertas do mais grave ao mais leve.
        _selo(aba.cell(numero, 10), analise.alertas[0].gravidade if analise.alertas else None)
    _filtro(aba, len(insumos))


def _aba_de_talhoes(aba: Worksheet, lista: list[TalhaoResposta]) -> None:
    _tabela(
        aba,
        [
            ("talhão", 28),
            ("área (ha)", 11),
            ("cultura", 14),
            ("situação", 13),
            ("colheita liberada em", 14),
            ("custo total (R$)", 16),
            ("custo por ha (R$)", 16),
            ("carência", 70),
            ("limitações", 80),
        ],
    )
    if not lista:
        _vazio(aba, "Nenhum talhão cadastrado.")
    for numero, talhao in enumerate(lista, start=2):
        _texto(aba.cell(numero, 1), talhao.nome)
        aba.cell(numero, 2, talhao.area_ha)
        if talhao.cultura is not None:
            _texto(aba.cell(numero, 3), talhao.cultura)
        aba.cell(numero, 5, talhao.colheita_liberada_em).number_format = FORMATO_DATA
        for coluna, valor in ((6, talhao.custo_total), (7, talhao.custo_por_ha)):
            # Sem preço nenhum, a célula diz isso, em vez de um zero de insumo de graça.
            celula = aba.cell(numero, coluna, "sem preço" if valor is None else valor)
            celula.number_format = FORMATO_DINHEIRO
        carencia = " ".join(alerta.justificativa for alerta in talhao.alertas) or None
        aba.cell(numero, 8, carencia)
        limitacoes = [f"• {texto}" for texto in talhao.limitacoes]
        aba.cell(numero, 9, "\n".join(limitacoes) or None)
        texto_corrido(aba, numero, "", 80)
        aba.row_dimensions[numero].height = 15 * max(
            1,
            math.ceil(len(carencia or "") / 77),
            sum(math.ceil(len(texto) / 88) for texto in limitacoes),
        )
        # Depois do texto_corrido, que alinha a linha inteira. Em carência, nas cores do
        # urgente da tela: colher antes é proibido.
        _selo(aba.cell(numero, 4), talhao.alertas[0].gravidade if talhao.alertas else None)
        aba.cell(numero, 5).alignment = Alignment(horizontal="center", vertical="top")
    _filtro(aba, len(lista))


def _movimentacoes(db: Session) -> list[tuple]:
    consulta = (
        select(Movimentacao, Insumo, Talhao.nome, Fornecedor.nome)
        .join(Insumo)
        .outerjoin(Talhao)
        .outerjoin(Fornecedor, Movimentacao.fornecedor_id == Fornecedor.id)
        .order_by(Movimentacao.data, Insumo.nome, Movimentacao.id)
    )
    return [
        (
            m.data,
            ILLEGAL_CHARACTERS_RE.sub("", insumo.nome),
            ILLEGAL_CHARACTERS_RE.sub("", insumo.unidade),
            _TIPO_NA_PLANILHA[(m.tipo, m.motivo)],
            m.quantidade,
            m.validade,
            m.valor_unitario,
            talhao and ILLEGAL_CHARACTERS_RE.sub("", talhao),
            fornecedor and ILLEGAL_CHARACTERS_RE.sub("", fornecedor),
        )
        for m, insumo, talhao, fornecedor in db.execute(consulta)
    ]


def _tabela(aba: Worksheet, colunas: list[tuple[str, int]]) -> None:
    titulos(aba, [titulo for titulo, _ in colunas])
    for indice, (_, largura) in enumerate(colunas, start=1):
        aba.column_dimensions[get_column_letter(indice)].width = largura
    aba.freeze_panes = "A2"
    para_imprimir(aba)


def _filtro(aba: Worksheet, linhas: int) -> None:
    """Setas de filtro nos títulos, para o leitor recortar a tabela no próprio Excel."""
    if linhas:
        aba.auto_filter.ref = f"A1:{get_column_letter(aba.max_column)}{linhas + 1}"


def _vazio(aba: Worksheet, texto: str) -> None:
    aba["A2"] = texto
    aba["A2"].font = fonte(italic=True, color=CINZA)


def _selo(celula: Cell, gravidade: Gravidade | None) -> None:
    rotulo, fundo, tinta = GRAVIDADES[gravidade]
    celula.value = rotulo
    celula.fill = PatternFill("solid", fgColor=fundo)
    celula.font = fonte(bold=True, color=tinta)
    celula.alignment = Alignment(horizontal="center", vertical="top")


def _texto(celula: Cell, valor: str) -> None:
    """Texto digitado pelo usuário: sempre como texto ("=1+1" não vira fórmula) e sem
    caractere que o Excel recusa (nome cadastrado antes da regra que os proíbe)."""
    celula.value = ILLEGAL_CHARACTERS_RE.sub("", valor)
    celula.data_type = "s"
