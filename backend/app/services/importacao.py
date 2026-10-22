"""Importação de movimentações a partir de uma planilha: do Excel (.xlsx) ou CSV.

Para o produtor que já anota compras e retiradas numa planilha não relançar tudo pela
tela. Tudo ou nada: com um erro sequer, nada entra. Meia planilha importada é pior que
nenhuma, porque o estoque fica errado sem que ninguém perceba.

Os dois formatos viram a mesma coisa — cabeçalho e linhas de células — e passam pelas
mesmas regras. No .xlsx, data e número já chegam com tipo; no CSV, tudo é texto, e o
formato dos números é deduzido da planilha inteira.
"""

import _csv
import csv
import io
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.relogio import hoje
from app.models import Insumo, Movimentacao, Talhao
from app.schemas.estoque import (
    QUANTIDADE_MAXIMA,
    SEM_CARACTERE_DE_CONTROLE,
    VALOR_UNITARIO_MAXIMO,
)

COLUNAS_OBRIGATORIAS = ("data", "insumo", "tipo", "quantidade")
# Colunas de número, que decidem juntas se a vírgula é decimal ou milhar. `valor` (o preço
# por unidade, só em compras) é opcional.
COLUNAS_DE_NUMERO = ("quantidade", "valor")
# Tipo escrito → (tipo, motivo). Descarte é saída que não é consumo: uma palavra na
# mesma coluna, em vez de uma coluna a mais para preencher.
TIPOS = {
    "entrada": ("entrada", None),
    "compra": ("entrada", None),
    "saida": ("saida", None),
    "saída": ("saida", None),
    "retirada": ("saida", None),
    "descarte": ("saida", "descarte"),
    # Ajustes de inventário, como o relatório (a cópia de segurança) os escreve.
    "sobra": ("entrada", "inventario"),
    "falta": ("saida", "inventario"),
}
# Título escrito com acento → o nome da coluna. Sem acento é o do modelo; com acento é o que
# o produtor escreve à mão.
SINONIMOS_DE_COLUNA = {"talhão": "talhao"}
# O Excel em português grava DD/MM/AAAA (ou DD/MM/AA, se a célula estiver assim).
FORMATOS_DE_DATA = ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y")
# Anos de registro de uma propriedade cabem com folga. O teto protege a API de um
# arquivo que, compactado, parece pequeno e se abre em milhões de linhas.
LINHAS_MAXIMAS = 50_000
# Tetos contra .xlsx montado para travar a API: pequeno compactado, enorme aberto.
# Uma planilha real de 50 mil linhas descompacta em uns 20 MB e tem centenas de
# milhares de células.
DESCOMPACTADO_MAXIMO = 50_000_000
CELULAS_MAXIMAS = 5_000_000
LINHA_MAXIMA_DO_EXCEL = 1_048_576

_XLSX_ILEGIVEL = (
    "Não deu para abrir a planilha do Excel. Abra o arquivo no Excel, salve de novo como "
    "Pasta de Trabalho do Excel (.xlsx) e tente outra vez."
)
_GRANDE_DEMAIS = (
    "A planilha é grande demais para importar de uma vez, ou tem linhas e colunas muito além "
    "dos dados. Apague as linhas e colunas vazias depois dos dados, ou divida em partes."
)
_MUITAS_LINHAS = (
    f"A planilha tem mais de {LINHAS_MAXIMAS:_} movimentações. Divida em partes (um ano por "
    "arquivo, por exemplo) e importe uma de cada vez."
).replace("_", ".")

# O CSV só dá texto; o .xlsx dá também número e data, já com tipo.
type Celula = str | float | date


@dataclass
class TotalDoInsumo:
    nome: str
    unidade: str
    novo: bool
    entradas: float = 0.0
    saidas: float = 0.0


@dataclass
class ResultadoImportacao:
    linhas: int = 0
    entradas: int = 0
    saidas: int = 0
    # Quanto entra e sai de cada insumo, na ordem da planilha: é onde um erro de
    # escala (mil lido como um) salta aos olhos antes de gravar.
    insumos: list[TotalDoInsumo] = field(default_factory=list)
    primeira_data: date | None = None
    ultima_data: date | None = None
    # Linhas idênticas a movimentações que já estão no banco: sinal de planilha
    # importada duas vezes. Só aviso — duas retiradas iguais no mesmo dia existem.
    ja_lancadas: int = 0
    erros: list[str] = field(default_factory=list)
    importado: bool = False


@dataclass(frozen=True)
class _Linha:
    insumo: Insumo  # o cadastrado, ou um novo que só entra na sessão ao confirmar
    tipo: str
    data: date
    quantidade: float
    validade: date | None
    motivo: str | None
    valor_unitario: float | None
    talhao: Talhao | None  # onde a retirada foi aplicada; só cadastrado


# (número da linha no arquivo, células) — ou, no lugar das células, o erro de forma da
# linha, que segue junto para a lista de erros sair na ordem do arquivo.
type _Lidas = list[tuple[int, list[Celula] | str]]


def importar(db: Session, arquivo: str | bytes, *, confirmar: bool) -> ResultadoImportacao:
    """Confere a planilha — texto = CSV, bytes = .xlsx — e devolve o resumo. Com
    `confirmar` e sem erros, grava."""
    resultado = ResultadoImportacao()
    if isinstance(arquivo, bytes):
        lido = _ler_xlsx(arquivo, resultado.erros)
    else:
        lido = _ler_csv(arquivo, resultado.erros)
    if lido is None:
        return resultado
    cabecalho, lidas, decimal_padrao = lido
    try:
        decimal = _separador_decimal(lidas, cabecalho, decimal_padrao)
    except ValueError as erro:
        resultado.erros.append(str(erro))
        return resultado

    cadastrados: dict[str, list[Insumo]] = {}
    for insumo in db.scalars(select(Insumo).order_by(Insumo.nome)):
        cadastrados.setdefault(insumo.nome.casefold(), []).append(insumo)
    novos: dict[str, Insumo] = {}
    talhoes: dict[str, list[Talhao]] = {}
    for talhao in db.scalars(select(Talhao).order_by(Talhao.nome)):
        talhoes.setdefault(talhao.nome.casefold(), []).append(talhao)
    linhas: list[_Linha] = []
    for numero, celulas in lidas:
        if isinstance(celulas, str):
            resultado.erros.append(f"Linha {numero}: {celulas}")
            continue
        campos = dict(zip(cabecalho, celulas, strict=False))
        try:
            linhas.append(_ler(campos, decimal, cadastrados, novos, talhoes))
        except ValueError as erro:
            resultado.erros.append(f"Linha {numero}: {erro}")

    if not linhas and not resultado.erros:
        resultado.erros.append("Nenhuma movimentação na planilha.")
    _resumir(resultado, linhas)
    resultado.ja_lancadas = _ja_lancadas(db, linhas)
    if resultado.erros or not confirmar:
        return resultado

    for linha in linhas:
        movimentacao = Movimentacao(
            tipo=linha.tipo,
            data=linha.data,
            quantidade=linha.quantidade,
            validade=linha.validade,
            motivo=linha.motivo,
            valor_unitario=linha.valor_unitario,
            talhao_id=linha.talhao and linha.talhao.id,
        )
        if linha.insumo.id is None:
            linha.insumo.movimentacoes.append(movimentacao)
        else:
            movimentacao.insumo_id = linha.insumo.id
            db.add(movimentacao)
    db.add_all(novos.values())
    db.commit()
    resultado.importado = True
    return resultado


def _ler_csv(conteudo: str, erros: list[str]) -> tuple[list[str], _Lidas, str] | None:
    """Cabeçalho, linhas e o separador decimal costumeiro do arquivo. None se não deu
    para ler (o motivo vai para `erros`)."""
    conteudo = conteudo.removeprefix("\N{BYTE ORDER MARK}")
    if not conteudo.strip():
        erros.append("O arquivo está vazio.")
        return None
    # O Excel em português separa por ponto e vírgula; o resto do mundo, por vírgula.
    delimitador = ";" if ";" in conteudo.split("\n", 1)[0] else ","
    leitor = csv.reader(io.StringIO(conteudo, newline=""), delimiter=delimitador)
    cabecalho = _cabecalho(next(leitor), erros)
    if cabecalho is None:
        return None
    lidas = _linhas_do_csv(leitor, cabecalho, erros)
    if lidas is None:
        return None
    # Sem pista nas células, vale o costume: quem separa colunas por ponto e vírgula (o
    # Excel em português) usa vírgula decimal.
    return cabecalho, lidas, "," if delimitador == ";" else "."


def _linhas_do_csv(leitor: _csv.Reader, cabecalho: list[str], erros: list[str]) -> _Lidas | None:
    lidas: _Lidas = []
    inicio = 2  # a linha 1 é o cabeçalho
    try:
        for celulas in leitor:
            numero, inicio = inicio, leitor.line_num + 1
            if not any(celula.strip() for celula in celulas):
                continue
            if any("\n" in celula or "\r" in celula for celula in celulas):
                # Uma aspa aberta faz o csv juntar o que vem depois numa célula só.
                lidas.append(
                    (numero, 'aspa (") sem fechar; o resto do arquivo virou uma célula só.')
                )
            elif any("\x00" in celula for celula in celulas):
                lidas.append((numero, "caractere inválido; salve a planilha de novo como CSV."))
            elif letra := _coluna_sem_titulo(cabecalho, celulas[: len(cabecalho)]):
                lidas.append((numero, _sem_titulo(letra)))
            elif any(celula.strip() for celula in celulas[len(cabecalho) :]):
                # Excel deixa ";;" no fim das linhas: célula vazia a mais não é erro.
                lidas.append(
                    (
                        numero,
                        f"{len(celulas)} colunas, mas o cabeçalho tem {len(cabecalho)}. Se a "
                        "quantidade tem vírgula decimal, separe as colunas por ponto e vírgula "
                        "ou ponha o número entre aspas.",
                    )
                )
            else:
                lidas.append((numero, [celula.strip() for celula in celulas]))
            if len(lidas) > LINHAS_MAXIMAS:
                erros.append(_MUITAS_LINHAS)
                return None
    except csv.Error as erro:
        erros.append(
            f'Linha {inicio}: não deu para ler o arquivo ({erro}). Há uma aspa (") sem fechar?'
        )
        return None
    return lidas


def _ler_xlsx(dados: bytes, erros: list[str]) -> tuple[list[str], _Lidas, str] | None:
    """Lê a aba Movimentações, ou a primeira. O número de cada linha é o da linha no Excel."""
    lidas: _Lidas = []
    try:
        with ZipFile(io.BytesIO(dados)) as compactado:
            # O tamanho declarado é obedecido: o zipfile para de descompactar nele.
            descompactado = sum(item.file_size for item in compactado.infolist())
        if descompactado > DESCOMPACTADO_MAXIMO:
            erros.append(_GRANDE_DEMAIS)
            return None
        # read_only: lê linha a linha, sem montar a planilha inteira na memória.
        livro = load_workbook(io.BytesIO(dados), read_only=True, data_only=True)
    except Exception:  # zip quebrado, .xls renomeado, XML malicioso barrado pelo defusedxml
        erros.append(_XLSX_ILEGIVEL)
        return None
    try:
        # A aba de preencher do modelo (e a do relatório, que serve de cópia de segurança);
        # numa planilha qualquer, a primeira.
        abas = {aba.title.casefold(): aba for aba in livro.worksheets}
        aba = abas.get("movimentações", livro.worksheets[0])
        # A dimensão gravada no arquivo pode mentir: menor, cortaria linhas sem aviso;
        # enorme, faria o leitor preencher milhões de células vazias.
        aba.reset_dimensions()
        valores = aba.iter_rows(values_only=True)
        primeira = next(valores, None)
        if primeira is None:
            erros.append("A planilha está vazia.")
            return None
        cabecalho = _cabecalho([_texto(_celula(valor)) for valor in primeira], erros)
        if cabecalho is None:
            return None
        celulas_lidas = 0
        for numero, linha in enumerate(valores, start=2):
            celulas_lidas += len(linha)
            if numero > LINHA_MAXIMA_DO_EXCEL or celulas_lidas > CELULAS_MAXIMAS:
                erros.append(_GRANDE_DEMAIS)
                return None
            # Contar no próprio C, antes de converter: as linhas vazias são a maioria.
            if linha.count(None) + linha.count("") == len(linha):
                continue
            celulas = [_celula(valor) for valor in linha]
            if not any(_texto(celula) for celula in celulas):
                continue
            if letra := _coluna_sem_titulo(cabecalho, celulas):
                lidas.append((numero, _sem_titulo(letra)))
            else:
                lidas.append((numero, celulas))
            if len(lidas) > LINHAS_MAXIMAS:
                erros.append(_MUITAS_LINHAS)
                return None
    except Exception:
        erros.append(_XLSX_ILEGIVEL)
        return None
    finally:
        livro.close()
    # Número digitado como texto numa planilha do Excel em português usa vírgula decimal.
    return cabecalho, lidas, ","


def _celula(valor: object) -> Celula:
    """Valor do openpyxl → célula: datetime vira data (a hora não interessa), número
    vira float, o resto vira texto."""
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, int | float) and not isinstance(valor, bool):
        return float(valor)
    return "" if valor is None else str(valor).strip()


def _texto(celula: Celula) -> str:
    if isinstance(celula, float):
        # 1020.0 → "1020": um código numérico no nome não ganha ".0".
        return str(int(celula)) if celula.is_integer() else str(celula)
    if isinstance(celula, date):
        return f"{celula:%d/%m/%Y}"
    return celula


def _coluna_sem_titulo(cabecalho: list[str], celulas: list[Celula]) -> str | None:
    """A letra da primeira coluna com valor e sem título: o valor seria descartado sem
    ninguém ver."""
    for indice, celula in enumerate(celulas):
        if (indice >= len(cabecalho) or not cabecalho[indice]) and _texto(celula):
            return get_column_letter(indice + 1)
    return None


def _sem_titulo(letra: str) -> str:
    return (
        f"há valores na coluna {letra}, que não tem título na linha 1; apague-os ou dê um "
        "título à coluna."
    )


def _cabecalho(valores: list[str], erros: list[str]) -> list[str] | None:
    cabecalho = [valor.strip().lower() for valor in valores]
    cabecalho = [SINONIMOS_DE_COLUNA.get(titulo, titulo) for titulo in cabecalho]
    if faltando := [c for c in COLUNAS_OBRIGATORIAS if c not in cabecalho]:
        erros.append(
            f"Faltam as colunas: {', '.join(faltando)}. A primeira linha precisa ter os "
            "títulos. Use a planilha modelo."
        )
        return None
    return cabecalho


def _separador_decimal(lidas: _Lidas, cabecalho: list[str], padrao: str) -> str:
    """Vírgula ou ponto, decidido pela planilha inteira: sozinha, a célula "25.500" não diz
    se é vinte e cinco e meio ou vinte e cinco mil. Só olha número escrito como texto, na
    quantidade e no valor: quem escreve "3,20" no preço escreve "1.000" na quantidade."""
    colunas = [(nome, cabecalho.index(nome)) for nome in COLUNAS_DE_NUMERO if nome in cabecalho]
    pistas: dict[str, tuple[int, str, str]] = {}
    for numero, celulas in lidas:
        if isinstance(celulas, str):
            continue  # linha com erro de forma
        for nome, coluna in colunas:
            valor = celulas[coluna] if coluna < len(celulas) else ""
            if not isinstance(valor, str):
                continue
            texto = _sem_moeda(valor)
            if (separador := _pista(texto)) is not None:
                pistas.setdefault(separador, (numero, nome, texto))
    if len(pistas) > 1:
        virgula, ponto = pistas[","], pistas["."]
        assunto = "Quantidades" if virgula[1] == ponto[1] == "quantidade" else "Números"
        raise ValueError(
            f"{assunto} com vírgula decimal ({_onde(*virgula)}) e com ponto decimal "
            f"({_onde(*ponto)}) na mesma planilha: use um formato só."
        )
    return next(iter(pistas)) if pistas else padrao


def _onde(numero: int, coluna: str, texto: str) -> str:
    # A quantidade é a coluna de sempre; as outras se dizem pelo nome.
    if coluna == "quantidade":
        return f"linha {numero}: {texto}"
    return f"linha {numero}, {coluna}: {texto}"


def _sem_moeda(texto: str) -> str:
    """O preço pode vir como "R$ 3,20": o símbolo e os espaços não fazem parte do número."""
    return texto.replace("R$", "").replace(" ", "")


def _pista(texto: str) -> str | None:
    """O separador decimal que a célula revela, se revela algum."""
    virgulas, pontos = texto.count(","), texto.count(".")
    if virgulas and pontos:  # 1.234,5 ou 1,234.5: o último é o decimal
        return "," if texto.rfind(",") > texto.rfind(".") else "."
    if virgulas + pontos > 1:  # 1.234.567: o repetido é o de milhar; o decimal é o outro
        return "," if pontos else "."
    if not virgulas + pontos:
        return None
    separador = "," if virgulas else "."
    inteiro, decimais = texto.split(separador)
    # "25.500" ou "1,000": três casas depois de um número que não é zero tanto podem ser
    # milhar quanto decimal. "0.250" e "2.5" não deixam dúvida.
    if len(decimais) == 3 and inteiro.lstrip("+-") not in ("", "0"):
        return None
    return separador


def _ler(
    campos: dict[str, Celula],
    decimal: str,
    cadastrados: dict[str, list[Insumo]],
    novos: dict[str, Insumo],
    talhoes: dict[str, list[Talhao]],
) -> _Linha:
    nome = _texto(campos.get("insumo", ""))
    if not nome:
        raise ValueError("falta o nome do insumo.")
    if len(nome) > 120:
        raise ValueError("nome do insumo com mais de 120 caracteres.")
    if not re.fullmatch(SEM_CARACTERE_DE_CONTROLE, nome):
        raise ValueError("o nome do insumo tem um caractere inválido.")
    unidade = _texto(campos.get("unidade", ""))
    if len(unidade) > 20:
        raise ValueError("unidade com mais de 20 caracteres.")
    if not re.fullmatch(SEM_CARACTERE_DE_CONTROLE, unidade):
        raise ValueError("a unidade tem um caractere inválido.")

    tipo_escrito = _texto(campos.get("tipo", ""))
    if (tipo_e_motivo := TIPOS.get(tipo_escrito.casefold())) is None:
        raise ValueError(f'tipo "{tipo_escrito}" desconhecido; use entrada, saída ou descarte.')
    tipo, motivo = tipo_e_motivo

    data = _data(campos.get("data", ""), "data")
    if data > hoje():
        raise ValueError("data no futuro.")

    quantidade = _numero(campos.get("quantidade", ""), decimal, "quantidade", casas=3)
    # `not >` em vez de `<=`: pega também o NaN, que float("nan") aceita.
    if not quantidade > 0:
        raise ValueError("quantidade deve ser maior que zero.")
    if quantidade > QUANTIDADE_MAXIMA:
        raise ValueError("quantidade grande demais.")

    validade = None
    if _texto(valor_validade := campos.get("validade", "")):
        if tipo != "entrada":
            raise ValueError("validade só se informa numa entrada.")
        validade = _data(valor_validade, "validade")

    valor_unitario = None
    if _texto(celula_valor := campos.get("valor", "")):
        # Preço é o que se pagou numa compra: descarte, falta e sobra não têm.
        if tipo != "entrada" or motivo is not None:
            raise ValueError("valor só se informa numa entrada (compra): é o preço pago.")
        # Quatro casas, como a coluna: o preço por grama de um defensivo tem centavos
        # depois da vírgula.
        valor_unitario = _numero(celula_valor, decimal, "valor", casas=4)
        if not valor_unitario > 0:
            raise ValueError("valor deve ser maior que zero.")
        if valor_unitario > VALOR_UNITARIO_MAXIMO:
            raise ValueError("valor grande demais.")
    talhao = None
    if nome_do_talhao := _texto(campos.get("talhao", "")):
        # Para onde o insumo foi aplicado: compra, descarte, sobra e falta não vão à lavoura.
        if tipo != "saida" or motivo is not None:
            raise ValueError(
                "talhão só se informa numa saída (retirada para uso): é onde o insumo foi aplicado."
            )
        talhao = _talhao(nome_do_talhao, talhoes)
    insumo = _insumo(nome, unidade, cadastrados, novos)
    return _Linha(insumo, tipo, data, quantidade, validade, motivo, valor_unitario, talhao)


def _numero(celula: Celula, decimal: str, coluna: str, *, casas: int) -> float:
    if isinstance(celula, float):
        return round(celula, casas)
    texto = _texto(celula)
    milhar = "." if decimal == "," else ","
    try:
        return round(float(_sem_moeda(texto).replace(milhar, "").replace(decimal, ".")), casas)
    except ValueError:
        raise ValueError(f'{coluna} "{texto}" não é um número.') from None


def _insumo(
    nome: str, unidade: str, cadastrados: dict[str, list[Insumo]], novos: dict[str, Insumo]
) -> Insumo:
    """Casa o nome da planilha com o cadastro, sem diferenciar maiúsculas. Insumo novo
    entra com a grafia e a unidade da primeira linha em que aparece."""
    chave = nome.casefold()
    match cadastrados.get(chave, []):
        case [insumo]:
            pass
        case []:
            if chave not in novos:
                if not unidade:
                    raise ValueError(
                        f'"{nome}" não está cadastrado; informe a unidade para cadastrá-lo.'
                    )
                novos[chave] = Insumo(nome=nome, unidade=unidade)
            insumo = novos[chave]
        case varios:
            # O cadastro diferencia maiúsculas; a planilha, não. Escolher um seria chute.
            raise ValueError(
                f'há mais de um insumo cadastrado como "{nome}" '
                f"({', '.join(i.nome for i in varios)}); renomeie um deles na tela de Insumos."
            )
    if unidade and unidade.casefold() != insumo.unidade.casefold():
        # Somar kg com saca daria um estoque sem sentido.
        raise ValueError(f'"{nome}" já está cadastrado em {insumo.unidade}, não em {unidade}.')
    return insumo


def _talhao(nome: str, talhoes: dict[str, list[Talhao]]) -> Talhao:
    """Casa o nome com o cadastro, sem diferenciar maiúsculas, como o insumo. Ao contrário
    dele, talhão novo não nasce da planilha: falta a área, sem a qual não há custo por
    hectare."""
    match talhoes.get(nome.casefold(), []):
        case [talhao]:
            return talhao
        case []:
            raise ValueError(f'talhão "{nome}" não cadastrado; cadastre na tela de Talhões.')
        case varios:
            raise ValueError(
                f'há mais de um talhão cadastrado como "{nome}" '
                f"({', '.join(t.nome for t in varios)}); renomeie um deles na tela de Talhões."
            )


def _data(celula: Celula, coluna: str) -> date:
    if isinstance(celula, date):
        return celula
    # Número numa coluna de data é o número de série do Excel, sinal de célula sem
    # formato de data: recusar é mais seguro que adivinhar.
    texto = _texto(celula)
    for formato in FORMATOS_DE_DATA:
        try:
            return datetime.strptime(texto, formato).date()  # noqa: DTZ007 — só a data
        except ValueError:
            continue
    raise ValueError(f'{coluna} "{texto}" inválida; use DD/MM/AAAA.')


def _resumir(resultado: ResultadoImportacao, linhas: list[_Linha]) -> None:
    totais: dict[int, TotalDoInsumo] = {}
    for linha in linhas:
        insumo = linha.insumo
        total = totais.setdefault(
            id(insumo), TotalDoInsumo(insumo.nome, insumo.unidade, novo=insumo.id is None)
        )
        if linha.tipo == "entrada":
            total.entradas = round(total.entradas + linha.quantidade, 3)
        else:
            total.saidas = round(total.saidas + linha.quantidade, 3)
    resultado.insumos = list(totais.values())
    resultado.linhas = len(linhas)
    resultado.entradas = sum(linha.tipo == "entrada" for linha in linhas)
    resultado.saidas = resultado.linhas - resultado.entradas
    if linhas:
        resultado.primeira_data = min(linha.data for linha in linhas)
        resultado.ultima_data = max(linha.data for linha in linhas)


def _ja_lancadas(db: Session, linhas: list[_Linha]) -> int:
    ids = {linha.insumo.id for linha in linhas if linha.insumo.id is not None}
    if not ids:
        return 0
    no_banco = Counter(
        (
            m.insumo_id,
            m.tipo,
            m.data,
            m.quantidade,
            m.validade,
            m.motivo,
            m.valor_unitario,
            m.talhao_id,
        )
        for m in db.scalars(select(Movimentacao).where(Movimentacao.insumo_id.in_(ids)))
    )
    repetidas = 0
    for linha in linhas:
        chave = (
            linha.insumo.id,
            linha.tipo,
            linha.data,
            linha.quantidade,
            linha.validade,
            linha.motivo,
            linha.valor_unitario,
            linha.talhao and linha.talhao.id,
        )
        if no_banco[chave] > 0:
            no_banco[chave] -= 1
            repetidas += 1
    return repetidas
